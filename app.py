import os
import csv
import re
import shutil
import subprocess
import tempfile
import threading
import concurrent.futures
import queue
import datetime
import platform
import time
import json
import ipaddress
import traceback
import tkinter as tk
import pandas as pd
from tkinter import ttk
import customtkinter as ctk
from tkinter import filedialog, messagebox, simpledialog
import paramiko
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter








# ============================================================
# CONFIGURATION
# ============================================================
APP_TITLE = "INFOBLOX DIG AUTOMATION By Suraj Joshi"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
COMMANDS_FILE = os.path.join(BASE_DIR, "commands.txt")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
SERVER_PROFILES_FILE = os.path.join(BASE_DIR, "server_profiles.json")
UI_PREFERENCES_FILE = os.path.join(BASE_DIR, "ui_preferences.json")








# ============================================================
# QUARTERLY APPLICATION LICENSE
# ============================================================
QUARTERLY_LICENSE_CODES = {
    "2026-Q4": "IBX-9B4Y-3FZK-TRPN",
    "2027-Q1": "IBX-MVVY-C8QJ-YY8L",
    "2027-Q2": "IBX-DTEF-YR2Z-AETM",
    "2027-Q3": "IBX-TR2H-2XVA-AL8L",
    "2027-Q4": "IBX-7HXC-MPW2-N979",
    "2028-Q1": "IBX-482L-G9GW-U6AM",
    "2028-Q2": "IBX-K2XR-YGTL-W8SC",
    "2028-Q3": "IBX-G5KH-W2N9-ESHD",
    "2028-Q4": "IBX-4QA8-HNSU-W6E3",
    "2029-Q1": "IBX-7D7J-5R8D-UHD9",
    "2029-Q2": "IBX-Q9GB-FBCK-8FVT",
    "2029-Q3": "IBX-N6PN-UHNS-8VHB",
}




LICENSE_STORAGE_DIR = os.path.join(
    os.environ.get("LOCALAPPDATA", os.path.expanduser("~")),
    "Infoblox_DIG_Automation"
)
LICENSE_STORAGE_FILE = os.path.join(LICENSE_STORAGE_DIR, "license.json")








def get_current_license_period(today=None):
    today = today or datetime.date.today()
    year = today.year
    month = today.month
    if month <= 3:
        quarter = 1
        expiry = datetime.date(year, 3, 31)
    elif month <= 6:
        quarter = 2
        expiry = datetime.date(year, 6, 30)
    elif month <= 9:
        quarter = 3
        expiry = datetime.date(year, 9, 30)
    else:
        quarter = 4
        expiry = datetime.date(year, 12, 31)
    return f"{year}-Q{quarter}", expiry








def load_local_license():
    try:
        with open(LICENSE_STORAGE_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except (OSError, json.JSONDecodeError, TypeError, ValueError):
        return {}








def save_local_license(code, quarter_key):
    os.makedirs(LICENSE_STORAGE_DIR, exist_ok=True)
    data = {
        "quarter": quarter_key,
        "code": code,
        "activated_on": datetime.date.today().isoformat(),
    }
    temp_file = LICENSE_STORAGE_FILE + ".tmp"
    with open(temp_file, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    os.replace(temp_file, LICENSE_STORAGE_FILE)








def ensure_quarterly_license(parent=None):
    quarter_key, expiry_date = get_current_license_period()
    expected_code = QUARTERLY_LICENSE_CODES.get(quarter_key)




    dialog_root = None
    if parent is None:
        dialog_root = tk.Tk()
        dialog_root.withdraw()
        parent = dialog_root




    try:
        if not expected_code:
            messagebox.showerror(
                "License Not Configured",
                (
                    f"No application license code is configured for {quarter_key}.\n\n"
                    "Please contact the application administrator."
                ),
                parent=parent,
            )
            return False




        saved = load_local_license()
        saved_code = str(saved.get("code", "")).strip().upper()
        saved_quarter = str(saved.get("quarter", "")).strip()




        if saved_quarter == quarter_key and saved_code == expected_code:
            return True




        expiry_text = expiry_date.strftime("%d-%b-%Y")
        prompt = (
            f"Quarterly license required for {quarter_key}.\n\n"
            f"Valid until: {expiry_text}\n\n"
            "Enter the application code provided by the administrator:"
        )




        while True:
            code = simpledialog.askstring(
                "Infoblox DIG Automation License",
                prompt,
                parent=parent,
            )
            if code is None:
                return False
            code = code.strip().upper()
            if code == expected_code:
                try:
                    save_local_license(code, quarter_key)
                except OSError as exc:
                    messagebox.showerror(
                        "License Save Error",
                        f"The license was valid, but it could not be saved locally.\n\n{exc}",
                        parent=parent,
                    )
                    return False
                return True
            messagebox.showerror(
                "Invalid License Code",
                (
                    "The code entered is not valid for the current quarter.\n\n"
                    "Please enter the current quarterly code."
                ),
                parent=parent,
            )
    finally:
        if dialog_root is not None:
            try:
                dialog_root.destroy()
            except Exception:
                pass








os.makedirs(OUTPUT_DIR, exist_ok=True)




# Active SSH channel used by the manual Infoblox pager button.
ACTIVE_SSH_CHANNEL = None
ACTIVE_SSH_CHANNEL_LOCK = threading.Lock()
# Additive cancellation support for single- and multi-server DIG runs.
STOP_DIG_EVENT = threading.Event()
ACTIVE_SSH_CHANNELS = set()








# ============================================================
# MODERN UI THEME CONSTANTS
# ============================================================
# Color palette (dark modern theme)
COLOR_BG            = "#0b0f14"   # app background
COLOR_SURFACE       = "#131a22"   # card surface
COLOR_SURFACE_ALT   = "#1c2530"   # elevated surface
COLOR_SURFACE_HOVER = "#232e3b"   # hover surface
COLOR_BORDER        = "#2a3542"   # subtle border
COLOR_ACCENT        = "#4f8cff"   # primary blue
COLOR_ACCENT_HOVER  = "#3a7aee"
COLOR_SUCCESS       = "#22c55e"   # green
COLOR_SUCCESS_HOVER = "#16a34a"
COLOR_WARNING       = "#f59e0b"   # amber
COLOR_DANGER        = "#ef4444"   # red
COLOR_DANGER_HOVER  = "#dc2626"
COLOR_TEXT          = "#e8eef5"   # primary text
COLOR_TEXT_MUTED    = "#8b98a8"   # secondary text
COLOR_TEXT_DIM      = "#5a6676"   # tertiary text
COLOR_PURPLE        = "#8b5cf6"
COLOR_PURPLE_HOVER  = "#7c3aed"
COLOR_CYAN          = "#06b6d4"




# Fonts
FONT_FAMILY = "Segoe UI"
FONT_MONO = "Consolas"








def font(size=12, weight="normal"):
    return ctk.CTkFont(family=FONT_FAMILY, size=size, weight=weight)








def mono_font(size=12, weight="normal"):
    return ctk.CTkFont(family=FONT_MONO, size=size, weight=weight)








# ============================================================
# UTILITY
# ============================================================
def now_string():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")








def clean_terminal_text(text):
    if not text:
        return ""
    text = re.sub(r"\x1B\][^\x07\x1B]*(?:\x07|\x1B\\)", "", text)
    text = re.sub(r"\x1B(?:[@-Z\\-_]|\[[0-?]*[ -/]*[@-~])", "", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    while "\x08" in text:
        text = re.sub(r".\x08", "", text)
    text = "".join(ch for ch in text if ch in "\r\n\t" or ord(ch) >= 32)
    return text








def normalize_command(command):
    if not command:
        return ""
    command = clean_terminal_text(command)
    command = command.strip()
    if ">" in command:
        command = command.split(">")[-1].strip()
    command = re.sub(r"\s+\+short\b", "", command, flags=re.IGNORECASE)
    command = re.sub(r"\s+", " ", command)
    return command.strip()








def command_key(command):
    return normalize_command(command).lower()








# ============================================================
# PUTTY / PPK SUPPORT
# ============================================================
def find_puttygen():
    candidates = [
        shutil.which("puttygen"),
        r"C:\Program Files\PuTTY\puttygen.exe",
        r"C:\Program Files (x86)\PuTTY\puttygen.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\Programs\PuTTY\puttygen.exe"),
    ]
    for path in candidates:
        if path and os.path.isfile(path):
            return path
    return None








def convert_ppk_to_openssh(ppk_file, passphrase=None):
    puttygen = find_puttygen()
    if not puttygen:
        raise Exception(
            "puttygen.exe was not found.\n\n"
            "Please install PuTTY or add puttygen.exe "
            "to your Windows PATH."
        )
    temp_dir = tempfile.mkdtemp(prefix="infoblox_key_")
    openssh_key = os.path.join(temp_dir, "converted_key")
    command = [puttygen, ppk_file, "-O", "private-openssh", "-o", openssh_key]
    if passphrase:
        command.extend(["--old-passphrase", passphrase])
    result = subprocess.run(
        command,
        capture_output=True,
        text=True,
        creationflags=(
            subprocess.CREATE_NO_WINDOW
            if platform.system() == "Windows"
            else 0
        )
    )
    if result.returncode != 0:
        error = result.stderr.strip() or result.stdout.strip() or "Unknown PuTTYgen error."
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise Exception("PPK conversion failed:\n\n" + error)
    if not os.path.exists(openssh_key):
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise Exception(
            "PuTTYgen completed but the converted "
            "private key was not created."
        )
    return openssh_key, temp_dir








# ============================================================
# PRIVATE KEY LOADING
# ============================================================
def load_private_key(key_file, passphrase=None):
    if not key_file:
        raise Exception("Please select a private key.")
    key_file = os.path.abspath(key_file)
    if not os.path.isfile(key_file):
        raise Exception(f"Private key not found:\n{key_file}")
    extension = os.path.splitext(key_file)[1].lower()
    temporary_directory = None
    try:
        if extension == ".ppk":
            key_file, temporary_directory = convert_ppk_to_openssh(key_file, passphrase)
        key_classes = []
        if hasattr(paramiko, "RSAKey"):
            key_classes.append(paramiko.RSAKey)
        if hasattr(paramiko, "ECDSAKey"):
            key_classes.append(paramiko.ECDSAKey)
        if hasattr(paramiko, "Ed25519Key"):
            key_classes.append(paramiko.Ed25519Key)
        if hasattr(paramiko, "DSSKey"):
            key_classes.append(paramiko.DSSKey)
        last_error = None
        for key_class in key_classes:
            try:
                key = key_class.from_private_key_file(key_file, password=passphrase)
                return key, temporary_directory
            except Exception as error:
                last_error = error
        raise Exception(
            "Unable to load the private key.\n\n"
            f"File:\n{key_file}\n\n"
            f"Last error:\n{last_error}"
        )
    except Exception:
        if temporary_directory:
            shutil.rmtree(temporary_directory, ignore_errors=True)
        raise








# ============================================================
# READ COMMANDS
# ============================================================
def ipv4_cidr_to_reverse_zone(value):
    value = (value or "").strip()
    if "/" not in value:
        return None
    try:
        network = ipaddress.ip_network(value, strict=False)
    except ValueError:
        return None
    if not isinstance(network, ipaddress.IPv4Network):
        return None
    if network.prefixlen not in (8, 16, 24, 32):
        return None
    octet_count = network.prefixlen // 8
    octets = str(network.network_address).split(".")[:octet_count]
    return ".".join(reversed(octets)) + ".in-addr.arpa"








def convert_cidr_tokens_in_dig_command(command):
    converted_parts = []
    for token in command.split():
        if token.startswith("@"):
            converted_parts.append(token)
            continue
        reverse_zone = ipv4_cidr_to_reverse_zone(token)
        converted_parts.append(reverse_zone if reverse_zone else token)
    return " ".join(converted_parts)








def read_commands(commands_file=None, trailing_option=""):
    if not commands_file:
        commands_file = COMMANDS_FILE
    trailing_option = (trailing_option or "").strip()
    if not os.path.exists(commands_file):
        raise Exception(f"Input file not found:\n\n{commands_file}")
    extension = os.path.splitext(commands_file)[1].lower()
    if extension in (".xlsx", ".xlsm", ".csv"):
        commands = []
        if extension in (".xlsx", ".xlsm"):
            try:
                workbook = load_workbook(commands_file, read_only=True, data_only=True)
            except Exception as error:
                raise Exception("Unable to read Excel file.\n\n" + str(error))
            try:
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                try:
                    headers = next(rows)
                except StopIteration:
                    headers = []
                header_map = {}
                for index, header in enumerate(headers):
                    if header is None:
                        continue
                    header_name = str(header).strip().lower()
                    header_map[header_name] = index
                if "fqdn*" not in header_map:
                    raise Exception("Excel file must contain an 'fqdn*' column.")
                if "forward_to" not in header_map:
                    raise Exception("Excel file must contain a 'forward_to' column.")
                fqdn_index = header_map["fqdn*"]
                forward_to_index = header_map["forward_to"]
                for row in rows:
                    if not row:
                        continue
                    fqdn = (
                        str(row[fqdn_index]).strip()
                        if fqdn_index < len(row) and row[fqdn_index] is not None
                        else ""
                    )
                    forward_to = (
                        str(row[forward_to_index]).strip()
                        if forward_to_index < len(row) and row[forward_to_index] is not None
                        else ""
                    )
                    if not fqdn or not forward_to:
                        continue
                    for target in forward_to.split(","):
                        target = target.strip()
                        if "/" not in target:
                            continue
                        ip = target.rsplit("/", 1)[1].strip()
                        if not ip:
                            continue
                        command = f"dig @{ip} {fqdn}"
                        if trailing_option:
                            command += f" {trailing_option}"
                        commands.append(command)
            finally:
                workbook.close()
        else:
            try:
                with open(commands_file, "r", encoding="utf-8-sig", newline="") as file:
                    reader = csv.reader(file)
                    try:
                        headers = next(reader)
                    except StopIteration:
                        headers = []
                    header_map = {}
                    for index, header in enumerate(headers):
                        header_name = str(header).strip().lower()
                        header_map[header_name] = index
                    if "fqdn*" not in header_map:
                        raise Exception("CSV file must contain an 'fqdn*' column.")
                    if "forward_to" not in header_map:
                        raise Exception("CSV file must contain a 'forward_to' column.")
                    fqdn_index = header_map["fqdn*"]
                    forward_to_index = header_map["forward_to"]
                    for row in reader:
                        if not row:
                            continue
                        fqdn = row[fqdn_index].strip() if fqdn_index < len(row) else ""
                        forward_to = row[forward_to_index].strip() if forward_to_index < len(row) else ""
                        if not fqdn or not forward_to:
                            continue
                        for target in forward_to.split(","):
                            target = target.strip()
                            if "/" not in target:
                                continue
                            ip = target.rsplit("/", 1)[1].strip()
                            if not ip:
                                continue
                            command = f"dig @{ip} {fqdn}"
                            if trailing_option:
                                command += f" {trailing_option}"
                            commands.append(command)
            except UnicodeDecodeError:
                raise Exception("Unable to read CSV file as UTF-8.")
        return commands




    commands = []
    with open(commands_file, "r", encoding="utf-8") as file:
        for line in file:
            line = line.strip()
            if not line:
                continue
            if line.lower().startswith("dig "):
                command = line
            else:
                command = "dig " + line
            command = convert_cidr_tokens_in_dig_command(command)
            command_parts = command.split()
            if (
                trailing_option
                and (
                    len(command_parts) == 2
                    or (len(command_parts) == 3 and command_parts[1].startswith("@"))
                )
            ):
                command += f" {trailing_option}"
            commands.append(command)
    return commands








# ============================================================
# SSH EXECUTION
# ============================================================
def execute_ssh(
    server,
    username,
    port,
    authentication,
    password,
    key_file,
    key_passphrase,
    commands,
    log_callback,
    keep_connection=False,
    progress_callback=None,
    command_callback=None
):
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    private_key = None
    temp_key_dir = None
    keep_open = False
    try:
        log_callback(f"Connecting to {server}:{port}...")
        if authentication == "Password":
            client.connect(
                hostname=server,
                port=int(port),
                username=username,
                password=password,
                timeout=20,
                banner_timeout=20,
                auth_timeout=20,
                look_for_keys=False,
                allow_agent=False
            )
        else:
            private_key, temp_key_dir = load_private_key(key_file, key_passphrase)
            log_callback("Private key loaded successfully.")
            client.connect(
                hostname=server,
                port=int(port),
                username=username,
                pkey=private_key,
                timeout=20,
                banner_timeout=20,
                auth_timeout=20,
                look_for_keys=False,
                allow_agent=False
            )
        log_callback(f"Connected to {server}")
        channel = client.invoke_shell(term="xterm", width=160, height=200)
        global ACTIVE_SSH_CHANNEL
        with ACTIVE_SSH_CHANNEL_LOCK:
            ACTIVE_SSH_CHANNEL = channel
            ACTIVE_SSH_CHANNELS.add(channel)
        channel.settimeout(1.0)
        time.sleep(2)
        initial = ""
        while True:
            try:
                if channel.recv_ready():
                    data = channel.recv(65535).decode("utf-8", errors="replace")
                    initial += data
                else:
                    break
            except Exception:
                break
        log_callback(f"Executing {len(commands)} DIG commands...")
        output = initial
        command_index = 0
        exit_sent = False
        pager_handled = False
        start_time = time.time()
        last_data_time = time.time()
        max_runtime = max(120, len(commands) * 15)
        current_command_index = None
        current_output_start = len(output)
        current_command_started = None




        def emit_command_event(event, command, status=None, started=None, ended=None):
            if command_callback:
                try:
                    command_callback(event, server, command, status, started, ended)
                except Exception:
                    pass




        if commands:
            current_command_index = 0
            current_command_started = time.time()
            current_output_start = len(output)
            emit_command_event("start", commands[0], None, current_command_started, None)
            channel.send(commands[0] + "\n")
            command_index = 1
        while True:
            if STOP_DIG_EVENT.is_set():
                log_callback("[STOP] DIG execution stopped by user.")
                break
            if channel.recv_ready():
                data = channel.recv(65535).decode("utf-8", errors="replace")
                output += data
                last_data_time = time.time()
                log_callback(data, raw=True)
                if (
                    not pager_handled
                    and re.search(
                        r"Enter <return> for next page or q<return> to cancel the command\.?",
                        clean_terminal_text(data),
                        flags=re.IGNORECASE
                    )
                ):
                    channel.send("q\n")
                    pager_handled = True
                    last_data_time = time.time()
                    log_callback(
                        "[AUTO] Infoblox pager detected - sent q + Enter. Continuing..."
                    )
            else:
                time.sleep(0.1)
            elapsed = time.time() - start_time
            cleaned = clean_terminal_text(output)
            prompt_returned = re.search(r"(?:^|\n)[^\n]*>\s*$", cleaned, flags=re.MULTILINE)
            if prompt_returned and time.time() - last_data_time > 0.5:
                if current_command_index is not None:
                    command_text = commands[current_command_index]
                    command_output = output[current_output_start:]
                    command_status = detect_actual_dns_status(command_output)
                    command_finished_at = time.time()
                    emit_command_event(
                        "complete", command_text, command_status,
                        current_command_started, command_finished_at
                    )
                    current_command_index = None
                if STOP_DIG_EVENT.is_set():
                    break
                if command_index < len(commands):
                    if progress_callback:
                        progress_callback(command_index, len(commands))
                    current_command_index = command_index
                    current_command_started = time.time()
                    current_output_start = len(output)
                    emit_command_event(
                        "start", commands[current_command_index], None,
                        current_command_started, None
                    )
                    channel.send(commands[command_index] + "\n")
                    command_index += 1
                    pager_handled = False
                    last_data_time = time.time()
                    continue
                if progress_callback:
                    progress_callback(len(commands), len(commands))
                if keep_connection:
                    break
                if not exit_sent:
                    channel.send("exit\n")
                    exit_sent = True
                    last_data_time = time.time()
                    continue
            if (
                exit_sent
                and re.search(r"(^|\n).*\>\s*exit\s*", output)
                and time.time() - last_data_time > 1.5
            ):
                break
            if elapsed > max_runtime:
                log_callback("SSH execution timeout reached.")
                break
        cleaned_output = clean_terminal_text(output)
        if keep_connection:
            keep_open = True
            return (cleaned_output, client, channel, temp_key_dir)
        return cleaned_output
    except Exception as exc:
        log_callback(f"SSH ERROR: {exc}")
        return ""
    finally:
        if not keep_open:
            with ACTIVE_SSH_CHANNEL_LOCK:
                if 'channel' in locals():
                    ACTIVE_SSH_CHANNELS.discard(channel)
                    if ACTIVE_SSH_CHANNEL is channel:
                        ACTIVE_SSH_CHANNEL = next(iter(ACTIVE_SSH_CHANNELS), None)
            try:
                client.close()
            except Exception:
                pass
            if temp_key_dir:
                shutil.rmtree(temp_key_dir, ignore_errors=True)








def execute_on_existing_session(channel, commands, log_callback, progress_callback=None, command_callback=None, server=""):
    if not commands:
        return ""
    global ACTIVE_SSH_CHANNEL
    with ACTIVE_SSH_CHANNEL_LOCK:
        ACTIVE_SSH_CHANNEL = channel
        ACTIVE_SSH_CHANNELS.add(channel)
    log_callback(f"Executing {len(commands)} DIG commands on existing SSH session...")
    output = ""
    command_index = 0
    pager_handled = False
    start_time = time.time()
    last_data_time = time.time()
    max_runtime = max(120, len(commands) * 15)
    current_command_index = 0
    current_command_started = time.time()
    current_output_start = len(output)
    if command_callback:
        try:
            command_callback("start", server, commands[0], None, current_command_started, None)
        except Exception:
            pass
    channel.send(commands[0] + "\n")
    command_index = 1
    while True:
        if STOP_DIG_EVENT.is_set():
            log_callback("[STOP] DIG execution stopped by user.")
            break
        if channel.recv_ready():
            data = channel.recv(65535).decode("utf-8", errors="replace")
            output += data
            last_data_time = time.time()
            log_callback(data, raw=True)
            if (
                not pager_handled
                and re.search(
                    r"Enter <return> for next page or q<return> to cancel the command\.?",
                    clean_terminal_text(data),
                    flags=re.IGNORECASE
                )
            ):
                channel.send("q\n")
                pager_handled = True
                last_data_time = time.time()
                log_callback(
                    "[AUTO] Infoblox pager detected - sent q + Enter. Continuing..."
                )
        else:
            time.sleep(0.1)
        elapsed = time.time() - start_time
        cleaned = clean_terminal_text(output)
        prompt_returned = re.search(r"(?:^|\n)[^\n]*>\s*$", cleaned, flags=re.MULTILINE)
        if prompt_returned and time.time() - last_data_time > 0.5:
            if current_command_index is not None:
                command_status = detect_actual_dns_status(output[current_output_start:])
                if command_callback:
                    try:
                        command_callback(
                            "complete", server, commands[current_command_index], command_status,
                            current_command_started, time.time()
                        )
                    except Exception:
                        pass
                current_command_index = None
            if STOP_DIG_EVENT.is_set():
                break
            if command_index < len(commands):
                if progress_callback:
                    progress_callback(command_index, len(commands))
                current_command_index = command_index
                current_command_started = time.time()
                current_output_start = len(output)
                if command_callback:
                    try:
                        command_callback(
                            "start", server, commands[current_command_index], None,
                            current_command_started, None
                        )
                    except Exception:
                        pass
                channel.send(commands[command_index] + "\n")
                command_index += 1
                pager_handled = False
                last_data_time = time.time()
                continue
            if progress_callback:
                progress_callback(len(commands), len(commands))
            break
        if elapsed > max_runtime:
            log_callback("SSH execution timeout reached.")
            break
    return clean_terminal_text(output)








# ============================================================
# PARSE DIG BLOCKS
# ============================================================
def parse_dig_blocks(text):
    text = clean_terminal_text(text or "")
    blocks = []
    current_command = None
    current_output = []
    for line in text.splitlines():
        clean_line = clean_terminal_text(line).strip()
        match = re.search(r"(?:^|\s)>\s*(dig\b.*)$", clean_line, re.IGNORECASE)
        if match:
            if current_command:
                blocks.append((current_command, "\n".join(current_output)))
            current_command = match.group(1).strip()
            current_output = []
            continue
        if re.search(r"(?:^|\s)>\s*exit\s*$", clean_line, re.IGNORECASE):
            if current_command:
                blocks.append((current_command, "\n".join(current_output)))
            current_command = None
            current_output = []
            continue
        if current_command:
            current_output.append(line)
    if current_command:
        blocks.append((current_command, "\n".join(current_output)))
    return blocks








def detect_actual_dns_status(output):
    output = clean_terminal_text(output or "")
    stripped = output.strip()
    lower = stripped.lower()
    if any(x in lower for x in (
        "connection timed out",
        "no servers could be reached",
        "communications error",
        "timed out",
    )):
        return "CONNECTION TIMED OUT"
    status_match = re.search(r"\bstatus\s*:\s*([A-Z]+)", stripped, flags=re.IGNORECASE)
    if status_match:
        return status_match.group(1).upper()
    ipv4_pattern = (
        r"\b(?:25[0-5]|2[0-4]\d|1?\d?\d)"
        r"(?:\.(?:25[0-5]|2[0-4]\d|1?\d?\d)){3}\b"
    )
    if re.search(ipv4_pattern, stripped):
        return "NOERROR"
    if re.search(r"\b[0-9a-fA-F]*:[0-9a-fA-F:]+\b", stripped):
        return "NOERROR"
    return "UNSUCCESSFUL" if not stripped else "OTHER"








def parse_step3_statuses(step3_output, unsuccessful_commands):
    text = clean_terminal_text(step3_output or "")
    statuses = []
    for line in text.splitlines():
        lower = line.lower()
        if "connection timed out" in lower or "no servers could be reached" in lower:
            statuses.append("CONNECTION TIMED OUT")
            continue
        match = re.search(r"\bstatus\s*:\s*([A-Z]+)", line, flags=re.IGNORECASE)
        if match:
            statuses.append(match.group(1).upper())
    return statuses[:len(unsuccessful_commands)]








def parse_actual_results(step_output):
    result = {}
    for command, output in parse_dig_blocks(step_output):
        key = command_key(command)
        result[key] = {
            "command": normalize_command(command),
            "status": detect_actual_dns_status(output),
            "output": output
        }
    return result








def extract_unsuccessful(step1_output):
    blocks = parse_dig_blocks(step1_output)
    unsuccessful = []
    for command, output in blocks:
        status = detect_actual_dns_status(output)
        if status == "UNSUCCESSFUL":
            unsuccessful.append(normalize_command(command))
    return unsuccessful








# ============================================================
# STEP 3 EXECUTION
# ============================================================
def run_step3(
    server,
    username,
    port,
    authentication,
    password,
    key_file,
    key_passphrase,
    unsuccessful_commands,
    log_callback,
    session=None,
    progress_callback=None,
    command_callback=None
):
    if not unsuccessful_commands:
        return ""
    commands = [cmd for cmd in unsuccessful_commands if cmd.strip()]
    if session:
        return execute_on_existing_session(
            channel=session[1],
            commands=commands,
            log_callback=log_callback,
            progress_callback=progress_callback,
            command_callback=command_callback,
            server=server
        )
    return execute_ssh(
        server=server,
        username=username,
        port=port,
        authentication=authentication,
        password=password,
        key_file=key_file,
        key_passphrase=key_passphrase,
        commands=commands,
        log_callback=log_callback,
        progress_callback=progress_callback,
        command_callback=command_callback
    )








# ============================================================
# FINAL TXT
# ============================================================
def create_final_file(
    server,
    username,
    authentication,
    commands,
    step1_output,
    unsuccessful,
    step3_output
):
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    final_file = os.path.join(OUTPUT_DIR, f"infoblox_dig_final_{timestamp}.txt")
    with open(final_file, "w", encoding="utf-8") as file:
        file.write("=" * 70 + "\n")
        file.write("INFOBLOX DIG AUTOMATION RESULT\n")
        file.write("=" * 70 + "\n\n")
        file.write(f"Server   : {server}\n")
        file.write(f"Username : {username}\n")
        file.write(f"Generated: {now_string()}\n")
        file.write(f"Auth     : {authentication}\n")
        file.write("\n")
        file.write("=" * 70 + "\n")
        file.write("STEP 1 - ALL DIG RESULTS\n")
        file.write("=" * 70 + "\n\n")
        file.write(step1_output)
        file.write("\n\n")
        file.write("=" * 70 + "\n")
        file.write("STEP 2 - UNSUCCESSFUL COMMANDS\n")
        file.write("=" * 70 + "\n\n")
        for command in unsuccessful:
            file.write(command + "\n")
        file.write("\n" + "=" * 70 + "\n")
        file.write("STEP 3 - RETRY RESULTS\n")
        file.write("=" * 70 + "\n\n")
        if step3_output:
            file.write(step3_output)
        else:
            file.write("No commands required retry.\n")
        file.write("\n\n" + "=" * 70 + "\n")
        file.write("SUMMARY\n")
        file.write("=" * 70 + "\n")
        file.write(f"Total commands       : {len(commands)}\n")
        file.write(f"Unsuccessful commands: {len(unsuccessful)}\n")
        file.write("Step 1              : Completed\n")
        file.write("Step 2              : Completed\n")
        file.write("Step 3              : Completed\n")
    return final_file








# ============================================================
# EXCEL
# ============================================================
def create_excel_result(commands, step1_output, step3_output, unsuccessful, server):
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_file = os.path.join(OUTPUT_DIR, f"infoblox_dig_report_{timestamp}.xlsx")
    step1_results = parse_actual_results(step1_output)
    step3_statuses = parse_step3_statuses(step3_output, unsuccessful)
    print("\n========== STEP 3 PARSED STATUS ==========")
    for i, command in enumerate(unsuccessful):
        status = step3_statuses[i] if i < len(step3_statuses) else "NO RESULT"
        print(f"{normalize_command(command)}  =>  {status}")
    print("==========================================\n")
    step3_by_key = {}
    for i, command in enumerate(unsuccessful):
        if i < len(step3_statuses):
            step3_by_key[command_key(command)] = step3_statuses[i]
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Server / IP", "Command", "Status"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    final_rows = []
    for original_command in commands:
        clean_command = normalize_command(original_command)
        key = command_key(clean_command)
        if key in step3_by_key:
            final_status = step3_by_key[key]
        else:
            step1 = step1_results.get(key)
            final_status = step1["status"] if step1 else "NO RESULT"
        final_rows.append((clean_command, final_status))
        ws.append([server, clean_command, final_status])
    summary = wb.create_sheet("Summary")
    summary["A1"] = "INFOBLOX DIG SUMMARY"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A3"] = "Total Commands"
    summary["B3"] = len(commands)
    summary["A4"] = "Retried Commands"
    summary["B4"] = len(unsuccessful)
    summary["A8"] = "FINAL DNS STATUS"
    summary["A8"].font = Font(bold=True, size=13)
    summary["A10"] = "Status"
    summary["B10"] = "Count"
    status_counts = {}
    for _, status in final_rows:
        status_counts[status] = status_counts.get(status, 0) + 1
    for cell in summary[10]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    row = 11
    preferred_order = [
        "NOERROR", "NXDOMAIN", "SERVFAIL", "REFUSED",
        "CONNECTION TIMED OUT", "UNSUCCESSFUL", "OTHER", "NO RESULT"
    ]
    added = set()
    for status in preferred_order:
        if status in status_counts:
            summary.cell(row=row, column=1, value=status)
            summary.cell(row=row, column=2, value=status_counts[status])
            added.add(status)
            row += 1
    for status, count in sorted(status_counts.items()):
        if status not in added:
            summary.cell(row=row, column=1, value=status)
            summary.cell(row=row, column=2, value=count)
            row += 1
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 28
    summary.column_dimensions["A"].width = 35
    summary.column_dimensions["B"].width = 15
    wb.save(excel_file)
    return excel_file








# ============================================================
# DIG EXCEL REPORT COMPARISON
# ============================================================
def read_dig_excel_statuses(file_path):
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        if "Results" in workbook.sheetnames:
            sheet = workbook["Results"]
        else:
            sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            raise ValueError(f"No rows found in Excel file: {file_path}")
        header_map = {
            str(value).strip().lower(): index
            for index, value in enumerate(headers)
            if value is not None
        }
        if "command" not in header_map or "status" not in header_map:
            raise ValueError(
                f"Excel file must contain 'Command' and 'Status' columns:\n{file_path}"
            )
        server_index = header_map.get("server / ip")
        command_index = header_map["command"]
        status_index = header_map["status"]
        results = []
        for row in rows:
            if not row or command_index >= len(row) or row[command_index] is None:
                continue
            command = normalize_command(str(row[command_index]))
            status = str(row[status_index]).strip().upper() if status_index < len(row) and row[status_index] is not None else "NO RESULT"
            server = str(row[server_index]).strip() if server_index is not None and server_index < len(row) and row[server_index] is not None else ""
            results.append({"command": command, "key": command_key(command), "status": status, "server": server})
        return results
    finally:
        workbook.close()








def create_dig_comparison_report(before_file, after_file, output_dir=None):
    before_rows = read_dig_excel_statuses(before_file)
    after_rows = read_dig_excel_statuses(after_file)
    before_by_key = {}
    after_by_key = {}
    for row in before_rows:
        before_by_key.setdefault(row["key"], []).append(row)
    for row in after_rows:
        after_by_key.setdefault(row["key"], []).append(row)




    ordered_keys = []
    for row in before_rows + after_rows:
        if row["key"] not in ordered_keys:
            ordered_keys.append(row["key"])




    report_rows = []
    for key in ordered_keys:
        old_list = before_by_key.get(key, [])
        new_list = after_by_key.get(key, [])
        count = max(len(old_list), len(new_list))
        for index in range(count):
            old = old_list[index] if index < len(old_list) else None
            new = new_list[index] if index < len(new_list) else None
            command = (new or old)["command"]
            old_status = old["status"] if old else "MISSING BEFORE"
            new_status = new["status"] if new else "MISSING AFTER"
            result = "UNCHANGED" if old and new and old_status == new_status else "CHANGED"
            server = (new or old).get("server", "")
            report_rows.append([server, command, old_status, new_status, result])




    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = output_dir or OUTPUT_DIR
    os.makedirs(output_dir, exist_ok=True)
    report_path = os.path.join(output_dir, f"infoblox_dig_comparison_{timestamp}.xlsx")
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Comparison"
    ws.append(["Server / IP", "Command", "Before Status", "After Status", "Difference"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for row in report_rows:
        ws.append(row)
        if row[4] == "CHANGED":
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor="FCE4D6")
    summary = workbook.create_sheet("Summary")
    changed = sum(1 for row in report_rows if row[4] == "CHANGED")
    unchanged = len(report_rows) - changed
    summary.append(["DIG RESULT COMPARISON SUMMARY", "Value"])
    summary.append(["Before File", os.path.basename(before_file)])
    summary.append(["After File", os.path.basename(after_file)])
    summary.append(["Total Commands Compared", len(report_rows)])
    summary.append(["Changed / Missing Commands", changed])
    summary.append(["Unchanged Commands", unchanged])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for col, width in {"A": 60, "B": 48, "C": 24, "D": 24, "E": 18}.items():
        ws.column_dimensions[col].width = width
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 60
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    workbook.save(report_path)
    return report_path, changed, unchanged, len(report_rows)








# ============================================================
# MODERN UI COMPONENTS
# ============================================================
class ModernCard(ctk.CTkFrame):
    """A rounded card with subtle border and optional title (grid-based)."""
    def __init__(self, master, title=None, **kwargs):
        super().__init__(
            master,
            fg_color=COLOR_SURFACE,
            corner_radius=16,
            border_width=1,
            border_color=COLOR_BORDER,
            **kwargs
        )
        self._title = title
        self.title_label = None
        if title:
            self.title_label = ctk.CTkLabel(
                self,
                text=title,
                font=font(11, "bold"),
                text_color=COLOR_TEXT_MUTED,
                anchor="w"
            )
            self.title_label.grid(
                row=0,
                column=0,
                columnspan=10,
                padx=18,
                pady=(14, 2),
                sticky="w"
            )








class StatusPill(ctk.CTkFrame):
    """A status indicator with colored dot and label."""
    def __init__(self, master, title, initial_text="Waiting", color=COLOR_TEXT_MUTED):
        super().__init__(
            master,
            fg_color=COLOR_SURFACE_ALT,
            corner_radius=12,
            border_width=1,
            border_color=COLOR_BORDER
        )
        self.color = color
        self.dot = ctk.CTkLabel(
            self,
            text="●",
            font=font(14),
            text_color=color
        )
        self.dot.pack(side="left", padx=(14, 6), pady=10)
        self.text_label = ctk.CTkLabel(
            self,
            text=initial_text,
            font=font(11, "bold"),
            text_color=COLOR_TEXT,
            anchor="w"
        )
        self.text_label.pack(side="left", padx=(0, 14), pady=10)
        self.title_label = ctk.CTkLabel(
            self,
            text=title,
            font=font(9),
            text_color=COLOR_TEXT_MUTED,
            anchor="e"
        )
        self.title_label.pack(side="right", padx=(0, 14), pady=10)




    def set_status(self, text, color):
        self.dot.configure(text_color=color)
        self.text_label.configure(text=text)
        self.color = color








# ============================================================
# APPLICATION
# ============================================================
class InfobloxDigApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.configure(fg_color=COLOR_BG)




        screen_w = self.root.winfo_screenwidth()
        screen_h = self.root.winfo_screenheight()
        initial_w = min(1400, max(1100, screen_w - 80))
        initial_h = min(920, max(700, screen_h - 100))
        self.root.geometry(f"{initial_w}x{initial_h}")
        self.root.minsize(960, 640)




        self.theme_mode = self._load_theme_preference()
        ctk.set_appearance_mode(self.theme_mode.lower())
        ctk.set_default_color_theme("blue")




        self.log_queue = queue.Queue()
        self.last_txt = None
        self.last_excel = None
        self._last_responsive_width = 0
        self._last_quick_width = 0
        self._resize_job = None




        # Live output window
        self.console = None
        self.live_output_window = None
        self.live_output_buffer = ""




        # Interactive SSH Terminal
        self.terminal_window = None
        self.terminal_text = None
        self.terminal_command_entry = None
        self.terminal_status_label = None
        self.terminal_connect_button = None
        self.terminal_disconnect_button = None
        self.terminal_client = None
        self.terminal_channel = None
        self.terminal_reader_thread = None
        self.terminal_connected = False
        self.terminal_output_buffer = ""
        self.terminal_output_lock = threading.Lock()
        self.terminal_temp_key_dir = None
        self.terminal_profile_var = None
        self.terminal_profile_menu = None
        self.terminal_trailing_option_entry = None
        self.terminal_profiles = []




        # Additive live DIG monitor state (does not change execution decisions).
        self.dig_monitor_window = None
        self.dig_monitor_tree = None
        self.dig_monitor_summary = None
        self.dig_monitor_progress = None
        self.dig_monitor_started_label = None
        self.dig_monitor_elapsed_label = None
        self.dig_monitor_rows = []
        self.dig_monitor_lock = threading.RLock()
        self.dig_monitor_started_at = None
        self.dig_monitor_finished_at = None
        self.dig_monitor_timer_running = False
        # Additive monitor controls: search/filter, failed-query view, ETA estimate.
        self.dig_monitor_search_var = None
        self.dig_monitor_status_var = None
        self.dig_monitor_eta_label = None
        self.dig_monitor_failed_window = None
        self.dig_monitor_failed_tree = None




        self._status_pills = {}




        # FQDN tab responsive state
        self._fqdn_last_width = 0
        self._fqdn_resize_job = None




        # Navigation state
        self._nav_buttons = {}
        self._pages = {}
        self._active_page = None




        self.build_gui()
        if self.theme_mode == "Light":
            # Widgets are initially built with the default dark palette.
            self.theme_mode = "Dark"
            self.apply_theme_palette("Light", persist=False)
        self.process_log_queue()




    def _load_theme_preference(self):
        try:
            with open(UI_PREFERENCES_FILE, "r", encoding="utf-8") as file:
                value = json.load(file).get("theme", "Dark")
            return value if value in ("Dark", "Light") else "Dark"
        except (OSError, json.JSONDecodeError, AttributeError, TypeError):
            return "Dark"




    def change_theme(self, value):
        if value not in ("Dark", "Light") or value == self.theme_mode:
            return
        self.apply_theme_palette(value, persist=True)




    def apply_theme_palette(self, mode, persist=False):
        """Recolor the existing UI while leaving DIG/SSH execution logic untouched."""
        # Normalize mapping based on the palette currently displayed.
        source_values = {
            "bg": "#F3F6FA" if self.theme_mode == "Light" else "#0b0f14",
            "surface": "#FFFFFF" if self.theme_mode == "Light" else "#131a22",
            "alt": "#E8EEF5" if self.theme_mode == "Light" else "#1c2530",
            "hover": "#DCE6F0" if self.theme_mode == "Light" else "#232e3b",
            "border": "#C7D2DF" if self.theme_mode == "Light" else "#2a3542",
            "text": "#17212F" if self.theme_mode == "Light" else "#e8eef5",
            "muted": "#526174" if self.theme_mode == "Light" else "#8b98a8",
            "dim": "#65758B" if self.theme_mode == "Light" else "#5a6676",
        }
        target_values = {
            "bg": "#0b0f14" if mode == "Dark" else "#F3F6FA",
            "surface": "#131a22" if mode == "Dark" else "#FFFFFF",
            "alt": "#1c2530" if mode == "Dark" else "#E8EEF5",
            "hover": "#232e3b" if mode == "Dark" else "#DCE6F0",
            "border": "#2a3542" if mode == "Dark" else "#C7D2DF",
            "text": "#e8eef5" if mode == "Dark" else "#17212F",
            "muted": "#8b98a8" if mode == "Dark" else "#526174",
            "dim": "#5a6676" if mode == "Dark" else "#65758B",
        }
        color_map = {source_values[key]: target_values[key] for key in source_values}
        color_properties = ("fg_color", "bg_color", "text_color", "border_color", "hover_color",
                            "placeholder_text_color", "button_color", "button_hover_color")
        def recolor(widget):
            try:
                for prop in color_properties:
                    try:
                        current = widget.cget(prop)
                    except Exception:
                        continue
                    if isinstance(current, str) and current in color_map:
                        try:
                            widget.configure(**{prop: color_map[current]})
                        except Exception:
                            pass
                    elif isinstance(current, (tuple, list)):
                        changed = tuple(color_map.get(item, item) for item in current)
                        if tuple(current) != changed:
                            try:
                                widget.configure(**{prop: changed})
                            except Exception:
                                pass
            except Exception:
                pass
            try:
                children = widget.winfo_children()
            except Exception:
                children = []
            for child in children:
                recolor(child)
        recolor(self.root)
        self.theme_mode = mode
        ctk.set_appearance_mode(mode.lower())
        # Additive theme fix: keep every CTkOptionMenu popup readable in both themes.
        # The dropdown popup colors are separate from the main button colors, so
        # changing the appearance mode alone does not reliably update them.
        dropdown_palette = {
            "dropdown_fg_color": "#FFFFFF" if mode == "Light" else "#1c2530",
            "dropdown_text_color": "#17212F" if mode == "Light" else "#e8eef5",
            "dropdown_hover_color": "#DCE6F0" if mode == "Light" else "#4f8cff",
            # Also set the selected value's text; this is separate from popup text.
            "text_color": "#17212F" if mode == "Light" else "#e8eef5",
            "text_color_disabled": "#526174" if mode == "Light" else "#8b98a8",
        }

        def style_dropdown_menus(widget):
            try:
                if isinstance(widget, ctk.CTkOptionMenu):
                    widget.configure(**dropdown_palette)
            except Exception:
                pass
            try:
                children = widget.winfo_children()
            except Exception:
                children = []
            for child in children:
                style_dropdown_menus(child)

        style_dropdown_menus(self.root)
        if getattr(self, "theme_menu", None):
            try:
                self.theme_menu.set(mode)
            except Exception:
                pass
        if persist:
            try:
                with open(UI_PREFERENCES_FILE, "w", encoding="utf-8") as file:
                    json.dump({"theme": mode}, file, indent=2)
            except OSError:
                pass




    # ========================================================
    # GUI
    # ========================================================
    def build_gui(self):
        # Root layout: sidebar + main content
        self.root.grid_rowconfigure(0, weight=1)
        self.root.grid_columnconfigure(0, weight=0)  # sidebar fixed
        self.root.grid_columnconfigure(1, weight=1)  # main content expands




        self._build_sidebar()
        self._build_main_area()




        # Build page contents
        self.build_dig_page()
        self.build_fqdn_tab()
        self.build_compare_tab()




        # Show default page
        self._show_page("dig")




        # Responsive
        self.root.bind("<Configure>", self._responsive_layout)
        self.root.after(150, self._responsive_layout)




    def _build_sidebar(self):
        """Modern left sidebar with brand + navigation."""
        sidebar = ctk.CTkFrame(
            self.root,
            width=240,
            corner_radius=0,
            fg_color=COLOR_SURFACE,
            border_width=0
        )
        sidebar.grid(row=0, column=0, sticky="nsw")
        sidebar.grid_propagate(False)
        sidebar.grid_rowconfigure(10, weight=1)  # spacer
        sidebar.grid_columnconfigure(0, weight=1)




        self.sidebar = sidebar




        # Brand block
        brand = ctk.CTkFrame(sidebar, fg_color="transparent")
        brand.grid(row=0, column=0, sticky="ew", padx=18, pady=(22, 18))




        logo_row = ctk.CTkFrame(brand, fg_color="transparent")
        logo_row.pack(fill="x")
        ctk.CTkLabel(
            logo_row,
            text="⬢",
            font=ctk.CTkFont(family=FONT_FAMILY, size=28, weight="bold"),
            text_color=COLOR_ACCENT
        ).pack(side="left", padx=(0, 10))




        title_box = ctk.CTkFrame(logo_row, fg_color="transparent")
        title_box.pack(side="left", fill="x", expand=True)
        ctk.CTkLabel(
            title_box,
            text="INFOBLOX",
            font=font(14, "bold"),
            text_color=COLOR_TEXT,
            anchor="w"
        ).pack(anchor="w")
        ctk.CTkLabel(
            title_box,
            text="DIG Automation",
            font=font(10),
            text_color=COLOR_TEXT_MUTED,
            anchor="w"
        ).pack(anchor="w")




        # Divider
        ctk.CTkFrame(
            sidebar, height=1, fg_color=COLOR_BORDER, corner_radius=0
        ).grid(row=1, column=0, sticky="ew", padx=18, pady=(0, 16))




        # Nav section label
        ctk.CTkLabel(
            sidebar,
            text="WORKSPACE",
            font=font(9, "bold"),
            text_color=COLOR_TEXT_DIM,
            anchor="w"
        ).grid(row=2, column=0, sticky="ew", padx=24, pady=(0, 8))




        # Navigation buttons
        nav_items = [
            ("dig", "▶   DIG Automation", "Run DIG & reports"),
            ("fqdn", "⌕   FQDN Search", "Search & row DIG"),
            ("compare", "⇄   Compare Results", "Before vs after"),
        ]
        for index, (key, label, subtitle) in enumerate(nav_items):
            btn = self._make_nav_button(sidebar, key, label, subtitle)
            btn.grid(row=3 + index, column=0, sticky="ew", padx=14, pady=3)
            self._nav_buttons[key] = btn




        # Bottom quick actions
        bottom = ctk.CTkFrame(sidebar, fg_color="transparent")
        bottom.grid(row=20, column=0, sticky="ew", padx=14, pady=(0, 18))




        self.live_output_button = ctk.CTkButton(
            bottom,
            text="⬒   Live Output",
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT,
            font=font(11, "bold"),
            anchor="w",
            command=self.open_live_output
        )
        self.live_output_button.pack(fill="x", pady=3)




        self.ssh_terminal_button = ctk.CTkButton(
            bottom,
            text="⌨   SSH Terminal",
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT,
            font=font(11, "bold"),
            anchor="w",
            command=self.open_ssh_terminal
        )
        self.ssh_terminal_button.pack(fill="x", pady=3)




        self.open_button = ctk.CTkButton(
            bottom,
            text="📁   Open Output",
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT,
            font=font(11, "bold"),
            anchor="w",
            command=self.open_output
        )
        self.open_button.pack(fill="x", pady=3)




        # Version footer
        ctk.CTkLabel(
            sidebar,
            text="By Suraj Joshi",
            font=font(9),
            text_color=COLOR_TEXT_DIM,
            anchor="center"
        ).grid(row=21, column=0, sticky="ew", padx=18, pady=(0, 14))




    def _make_nav_button(self, parent, key, label, subtitle):
        btn = ctk.CTkButton(
            parent,
            text=f"{label}",
            height=48,
            corner_radius=12,
            fg_color="transparent",
            hover_color=COLOR_SURFACE_ALT,
            text_color=COLOR_TEXT_MUTED,
            font=font(12, "bold"),
            anchor="w",
            command=lambda k=key: self._show_page(k)
        )
        return btn




    def _build_main_area(self):
        """Main content area with a top header bar and a page container."""
        main = ctk.CTkFrame(self.root, fg_color=COLOR_BG, corner_radius=0)
        main.grid(row=0, column=1, sticky="nsew")
        main.grid_rowconfigure(1, weight=1)
        main.grid_columnconfigure(0, weight=1)




        # Header bar
        header = ctk.CTkFrame(
            main,
            fg_color=COLOR_SURFACE,
            corner_radius=0,
            height=72,
            border_width=0
        )
        header.grid(row=0, column=0, sticky="ew")
        header.grid_propagate(False)
        header.grid_columnconfigure(0, weight=1)




        title_area = ctk.CTkFrame(header, fg_color="transparent")
        title_area.grid(row=0, column=0, sticky="w", padx=28, pady=14)




        self.page_title_label = ctk.CTkLabel(
            title_area,
            text="DIG Automation",
            font=font(18, "bold"),
            text_color=COLOR_TEXT,
            anchor="w"
        )
        self.page_title_label.pack(anchor="w")




        self.page_subtitle_label = ctk.CTkLabel(
            title_area,
            text="Run DIG commands, validate results, and export reports",
            font=font(10),
            text_color=COLOR_TEXT_MUTED,
            anchor="w"
        )
        self.page_subtitle_label.pack(anchor="w")




        # Status indicators in header (right)
        self.header_status = ctk.CTkFrame(header, fg_color="transparent")
        self.header_status.grid(row=0, column=1, sticky="e", padx=20, pady=14)
        ctk.CTkLabel(
            self.header_status, text="UI THEME", font=font(9, "bold"),
            text_color=COLOR_TEXT_MUTED
        ).pack(side="left", padx=(0, 8))
        self.theme_menu = ctk.CTkOptionMenu(
            self.header_status, values=["Dark", "Light"], width=100, height=32,
            corner_radius=9, font=font(10, "bold"),
            fg_color=COLOR_ACCENT, button_color=COLOR_ACCENT,
            button_hover_color=COLOR_ACCENT_HOVER, text_color="#FFFFFF",
            dropdown_fg_color=COLOR_SURFACE, dropdown_text_color=COLOR_TEXT,
            dropdown_hover_color=COLOR_SURFACE_ALT,
            command=self.change_theme
        )
        self.theme_menu.set(self.theme_mode)
        self.theme_menu.pack(side="left")




        # Page container
        self.page_container = ctk.CTkFrame(main, fg_color="transparent", corner_radius=0)
        self.page_container.grid(row=1, column=0, sticky="nsew")
        self.page_container.grid_rowconfigure(0, weight=1)
        self.page_container.grid_columnconfigure(0, weight=1)




    def _show_page(self, key):
        """Switch the visible page and update nav highlight."""
        if key not in self._pages:
            return
        for page_key, page in self._pages.items():
            if page_key == key:
                page.grid(row=0, column=0, sticky="nsew")
            else:
                page.grid_remove()




        self._active_page = key




        # Update nav button styles
        for nav_key, btn in self._nav_buttons.items():
            if nav_key == key:
                btn.configure(
                    fg_color=COLOR_SURFACE_ALT,
                    text_color=COLOR_TEXT
                )
            else:
                btn.configure(
                    fg_color="transparent",
                    text_color=COLOR_TEXT_MUTED
                )




        # Update header title
        titles = {
            "dig": ("DIG Automation", "Run DIG commands, validate results, and export reports"),
            "fqdn": ("FQDN Search", "Search the CSV data, review forwarders, and run DIG on selected rows"),
            "compare": ("Compare DIG Results", "Compare before/after activity Excel reports"),
        }
        if key in titles:
            title, subtitle = titles[key]
            try:
                self.page_title_label.configure(text=title)
                self.page_subtitle_label.configure(text=subtitle)
            except Exception:
                pass




    # --------------------------------------------------------
    # DIG PAGE
    # --------------------------------------------------------
    def build_dig_page(self):
        page = ctk.CTkFrame(self.page_container, fg_color="transparent", corner_radius=0)
        self._pages["dig"] = page




        scroll = ctk.CTkScrollableFrame(
            page,
            fg_color="transparent",
            corner_radius=0
        )
        scroll.pack(fill="both", expand=True)
        scroll.grid_columnconfigure(0, weight=1)




        # ====================================================
        # CONNECTION CARD
        # ====================================================
        conn_card = ModernCard(scroll, title="CONNECTION")
        conn_card.grid(row=0, column=0, sticky="ew", padx=24, pady=(20, 12))
        conn_card.grid_columnconfigure(1, weight=1)
        conn_card.grid_columnconfigure(3, weight=1)




        # Row: Server / Username
        ctk.CTkLabel(
            conn_card, text="Server / IP",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        ).grid(row=1, column=0, padx=(18, 10), pady=(10, 4), sticky="w")
        self.server_entry = self._entry(conn_card)
        self.server_entry.grid(row=1, column=1, padx=(0, 16), pady=(10, 4), sticky="ew")
        self.server_entry.insert(0, "eudefra11dns02.info.corp")




        ctk.CTkLabel(
            conn_card, text="Username",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        ).grid(row=1, column=2, padx=(0, 10), pady=(10, 4), sticky="w")
        self.username_entry = self._entry(conn_card)
        self.username_entry.grid(row=1, column=3, padx=(0, 18), pady=(10, 4), sticky="ew")
        self.username_entry.insert(0, "admin")




        # Row: Port / Auth
        ctk.CTkLabel(
            conn_card, text="SSH Port",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        ).grid(row=2, column=0, padx=(18, 10), pady=4, sticky="w")
        self.port_entry = self._entry(conn_card, width=110)
        self.port_entry.grid(row=2, column=1, padx=(0, 16), pady=4, sticky="w")
        self.port_entry.insert(0, "22")




        ctk.CTkLabel(
            conn_card, text="Authentication",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        ).grid(row=2, column=2, padx=(0, 10), pady=4, sticky="w")
        self.auth_var = ctk.StringVar(value="Password")
        self.auth_menu = ctk.CTkOptionMenu(
            conn_card,
            variable=self.auth_var,
            values=["Password", "SSH Key"],
            command=self.authentication_changed,
            width=180,
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            button_color=COLOR_ACCENT,
            button_hover_color=COLOR_ACCENT_HOVER,
            dropdown_fg_color=COLOR_SURFACE_ALT,
            dropdown_hover_color=COLOR_ACCENT,
            font=font(11)
        )
        self.auth_menu.grid(row=2, column=3, padx=(0, 18), pady=4, sticky="ew")




        # Row: Password
        self.password_label = ctk.CTkLabel(
            conn_card, text="SSH Password",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        )
        self.password_label.grid(row=3, column=0, padx=(18, 10), pady=(4, 14), sticky="w")
        self.password_entry = self._entry(conn_card, show="*")
        self.password_entry.grid(row=3, column=1, columnspan=3, padx=(0, 18), pady=(4, 14), sticky="ew")




        # Additive server-profile dropdown for the main DIG connection.
        ctk.CTkLabel(
            conn_card, text="Server Profile",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        ).grid(row=5, column=0, padx=(18, 10), pady=(4, 14), sticky="w")
        self.main_server_profile_var = ctk.StringVar(value="Select Server Profile")
        self.main_server_profile_menu = ctk.CTkOptionMenu(
            conn_card,
            variable=self.main_server_profile_var,
            values=["Select Server Profile"],
            command=self.main_server_profile_changed,
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            button_color=COLOR_ACCENT,
            button_hover_color=COLOR_ACCENT_HOVER,
            dropdown_fg_color=COLOR_SURFACE_ALT,
            dropdown_hover_color=COLOR_ACCENT,
            font=font(11)
        )
        self.main_server_profile_menu.grid(
            row=5, column=1, padx=(0, 10), pady=(4, 14), sticky="ew"
        )
        ctk.CTkButton(
            conn_card,
            text="↻ Refresh",
            width=100,
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT,
            font=font(10, "bold"),
            command=self.refresh_main_server_profiles
        ).grid(row=5, column=2, columnspan=2, padx=(0, 18), pady=(4, 14), sticky="w")
        self.refresh_main_server_profiles()

        # Key widgets (hidden initially)
        self.key_label = ctk.CTkLabel(
            conn_card, text="Private Key",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        )
        self.key_entry = self._entry(conn_card)
        self.key_browse = ctk.CTkButton(
            conn_card,
            text="Browse",
            width=100,
            height=38,
            corner_radius=10,
            fg_color=COLOR_ACCENT,
            hover_color=COLOR_ACCENT_HOVER,
            font=font(11, "bold"),
            command=self.browse_key
        )
        self.key_passphrase_label = ctk.CTkLabel(
            conn_card, text="Key Passphrase",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        )
        self.key_passphrase_entry = self._entry(conn_card, show="*")




        # ====================================================
        # INPUT CARD
        # ====================================================
        input_card = ModernCard(scroll, title="INPUT")
        input_card.grid(row=1, column=0, sticky="ew", padx=24, pady=12)
        input_card.grid_columnconfigure(1, weight=1)




        # Command File
        ctk.CTkLabel(
            input_card, text="Command File",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        ).grid(row=1, column=0, padx=(18, 10), pady=(10, 4), sticky="w")
        self.commands_file_entry = self._entry(input_card)
        self.commands_file_entry.grid(row=1, column=1, padx=(0, 10), pady=(10, 4), sticky="ew")
        self.commands_file_entry.insert(0, COMMANDS_FILE)
        self.commands_browse = ctk.CTkButton(
            input_card,
            text="Browse",
            width=100,
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT,
            font=font(11, "bold"),
            command=self.browse_commands_file
        )
        self.commands_browse.grid(row=1, column=2, padx=(0, 18), pady=(10, 4))




        # Quick DIG
        ctk.CTkLabel(
            input_card, text="Quick DIG / FQDN",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="nw"
        ).grid(row=2, column=0, padx=(18, 10), pady=4, sticky="nw")
        self.quick_commands_text = ctk.CTkTextbox(
            input_card,
            height=120,
            corner_radius=10,
            fg_color=COLOR_BG,
            border_width=1,
            border_color=COLOR_BORDER,
            text_color=COLOR_TEXT,
            font=mono_font(12),
            wrap="none"
        )
        self.quick_commands_text.grid(row=2, column=1, padx=(0, 10), pady=4, sticky="ew")




        trailing_option_frame = ctk.CTkFrame(input_card, fg_color="transparent")
        trailing_option_frame.grid(row=2, column=2, padx=(0, 18), pady=4, sticky="nw")




        ctk.CTkLabel(
            trailing_option_frame, text="Trailing Option A/PTR/SOA",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        ).pack(anchor="w", pady=(0, 4))
        self.trailing_option_entry = self._entry(trailing_option_frame, width=210)
        self.trailing_option_entry.pack(fill="x", pady=(0, 8))
        ctk.CTkLabel(
            trailing_option_frame, text="Empty = no trailing option",
            font=font(9), text_color=COLOR_TEXT_DIM, anchor="w"
        ).pack(anchor="w")




        # Output Folder
        ctk.CTkLabel(
            input_card, text="Output Folder",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        ).grid(row=3, column=0, padx=(18, 10), pady=(4, 14), sticky="w")
        self.output_entry = self._entry(input_card)
        self.output_entry.grid(row=3, column=1, padx=(0, 10), pady=(4, 14), sticky="ew")
        self.output_entry.insert(0, OUTPUT_DIR)
        ctk.CTkButton(
            input_card,
            text="Browse",
            width=100,
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT,
            font=font(11, "bold"),
            command=self.browse_output
        ).grid(row=3, column=2, padx=(0, 18), pady=(4, 14))




        # Fast DIG Mode
        fast_mode_frame = ctk.CTkFrame(input_card, fg_color="transparent")
        fast_mode_frame.grid(
            row=4, column=0, columnspan=3, sticky="ew", padx=18, pady=(0, 14)
        )
        self.fast_mode_var = ctk.BooleanVar(value=False)
        ctk.CTkCheckBox(
            fast_mode_frame,
            text="Enable Fast DIG Mode (parallel SSH workers)",
            variable=self.fast_mode_var,
            font=font(10, "bold"),
            text_color=COLOR_TEXT,
            fg_color=COLOR_ACCENT,
            hover_color=COLOR_ACCENT_HOVER,
            checkbox_width=20,
            checkbox_height=20,
            corner_radius=6
        ).pack(side="left", padx=(0, 18))
        ctk.CTkLabel(
            fast_mode_frame, text="Workers (1–10):",
            font=font(10), text_color=COLOR_TEXT_MUTED
        ).pack(side="left", padx=(0, 8))
        self.fast_workers_entry = self._entry(fast_mode_frame, width=70)
        self.fast_workers_entry.insert(0, "3")
        self.fast_workers_entry.pack(side="left")
        ctk.CTkLabel(
            fast_mode_frame,
            text="Commands are divided as evenly as possible; one SSH session per worker.",
            font=font(9), text_color=COLOR_TEXT_DIM
        ).pack(side="left", padx=(12, 0))




        # ====================================================
        # STATUS CARD
        # ====================================================
        status_card = ModernCard(scroll, title="STATUS")
        status_card.grid(row=2, column=0, sticky="ew", padx=24, pady=12)
        status_card.grid_columnconfigure(0, weight=1)
        status_card.grid_columnconfigure(1, weight=1)
        status_card.grid_columnconfigure(2, weight=1)
        status_card.grid_columnconfigure(3, weight=1)




        self.step1_pill = StatusPill(status_card, "STEP 1", "Waiting", COLOR_TEXT_MUTED)
        self.step1_pill.grid(row=1, column=0, padx=(18, 6), pady=(4, 6), sticky="ew")
        self.step2_pill = StatusPill(status_card, "STEP 2", "Waiting", COLOR_TEXT_MUTED)
        self.step2_pill.grid(row=1, column=1, padx=6, pady=(4, 6), sticky="ew")
        self.step3_pill = StatusPill(status_card, "STEP 3", "Waiting", COLOR_TEXT_MUTED)
        self.step3_pill.grid(row=1, column=2, padx=6, pady=(4, 6), sticky="ew")
        self.excel_pill = StatusPill(status_card, "EXCEL", "Waiting", COLOR_TEXT_MUTED)
        self.excel_pill.grid(row=1, column=3, padx=(6, 18), pady=(4, 6), sticky="ew")




        self.step1_status = self.step1_pill.text_label
        self.step2_status = self.step2_pill.text_label
        self.step3_status = self.step3_pill.text_label
        self.excel_status = self.excel_pill.text_label




        self.step1_progress = ctk.CTkLabel(
            status_card, text="DIGs: 0 / 0",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        )
        self.step1_progress.grid(row=2, column=0, columnspan=2, padx=(24, 6), pady=(0, 14), sticky="w")
        self.step3_progress = ctk.CTkLabel(
            status_card, text="DIGs: 0 / 0",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        )
        self.step3_progress.grid(row=2, column=2, columnspan=2, padx=(6, 24), pady=(0, 14), sticky="w")




        # ====================================================
        # ACTION BUTTONS
        # ====================================================
        action_card = ModernCard(scroll)
        action_card.grid(row=3, column=0, sticky="ew", padx=24, pady=(12, 24))
        action_card.grid_columnconfigure(0, weight=2)
        action_card.grid_columnconfigure(1, weight=1)
        action_card.grid_columnconfigure(2, weight=1)




        self.run_button = ctk.CTkButton(
            action_card,
            text="▶   RUN ALL STEPS",
            height=50,
            corner_radius=12,
            fg_color="#16A34A",
            hover_color="#15803D",
            text_color="#ffffff",
            font=font(13, "bold"),
            command=self.start_run
        )
        self.run_button.grid(row=0, column=0, padx=(18, 6), pady=18, sticky="ew")




        self.stop_dig_button = ctk.CTkButton(
            action_card,
            text="■   STOP DIG",
            height=50,
            corner_radius=12,
            fg_color=COLOR_DANGER,
            hover_color=COLOR_DANGER_HOVER,
            text_color="#ffffff",
            font=font(12, "bold"),
            command=self.stop_dig,
            state="disabled"
        )
        self.stop_dig_button.grid(row=0, column=2, padx=(6, 18), pady=18, sticky="ew")




        self.profiles_button = ctk.CTkButton(
            action_card,
            text="⚙   SERVER PROFILES / MULTI TEST",
            height=50,
            corner_radius=12,
            fg_color="#EA580C",
            hover_color="#C2410C",
            text_color="#ffffff",
            font=font(12, "bold"),
            command=self.open_server_profiles
        )
        self.profiles_button.grid(row=0, column=1, padx=6, pady=18, sticky="ew")




        self.authentication_changed("Password")




    # --------------------------------------------------------
    # Entry helper
    # --------------------------------------------------------
    def _entry(self, master, width=240, show=None):
        return ctk.CTkEntry(
            master,
            height=38,
            corner_radius=10,
            fg_color=COLOR_BG,
            border_width=1,
            border_color=COLOR_BORDER,
            text_color=COLOR_TEXT,
            placeholder_text_color=COLOR_TEXT_MUTED,
            font=font(11),
            width=width,
            show=show
        )




    def _responsive_layout(self, event=None):
        try:
            if event is not None and event.widget is not self.root:
                return
            if self._resize_job is not None:
                try:
                    self.root.after_cancel(self._resize_job)
                except Exception:
                    pass
            self._resize_job = self.root.after(120, self._apply_responsive_layout)
        except Exception:
            pass




    def _apply_responsive_layout(self):
        self._resize_job = None
        try:
            width = self.root.winfo_width()
            height = self.root.winfo_height()
            if width < 100 or height < 100:
                return
            if width == self._last_responsive_width:
                return
            quick_width = max(280, min(850, int(width * 0.42)))
            if quick_width != self._last_quick_width:
                try:
                    self.quick_commands_text.configure(width=quick_width)
                    self._last_quick_width = quick_width
                except Exception:
                    pass
            self._last_responsive_width = width
        except Exception:
            pass




    # ========================================================
    # COMPARE TAB PAGE
    # ========================================================
    def build_compare_tab(self):
        page = ctk.CTkFrame(self.page_container, fg_color="transparent", corner_radius=0)
        self._pages["compare"] = page




        outer = ctk.CTkScrollableFrame(page, fg_color="transparent", corner_radius=0)
        outer.pack(fill="both", expand=True)
        outer.grid_columnconfigure(0, weight=1)




        card = ModernCard(outer, title="COMPARE BEFORE / AFTER DIG RESULTS")
        card.grid(row=0, column=0, sticky="ew", padx=24, pady=(20, 12))
        card.grid_columnconfigure(1, weight=1)




        ctk.CTkLabel(
            card,
            text=("Select the Excel reports generated by this tool: first the BEFORE activity report, "
                  "then the AFTER activity report. Results are matched by DIG command and DNS status."),
            font=font(11), text_color=COLOR_TEXT, justify="left",
            wraplength=900, anchor="w"
        ).grid(row=1, column=0, columnspan=3, padx=18, pady=(10, 18), sticky="ew")




        self.compare_before_entry = self._entry(card)
        self.compare_before_entry.configure(placeholder_text="Select BEFORE activity Excel report")
        self.compare_before_entry.grid(row=2, column=1, padx=(10, 10), pady=8, sticky="ew")
        ctk.CTkLabel(
            card, text="Before Activity",
            font=font(10), text_color=COLOR_TEXT_MUTED
        ).grid(row=2, column=0, padx=(18, 4), pady=8, sticky="w")
        ctk.CTkButton(
            card, text="Browse...", width=110, height=36, corner_radius=10,
            fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT, font=font(10, "bold"),
            command=self.browse_compare_before
        ).grid(row=2, column=2, padx=(0, 18), pady=8)




        self.compare_after_entry = self._entry(card)
        self.compare_after_entry.configure(placeholder_text="Select AFTER activity Excel report")
        self.compare_after_entry.grid(row=3, column=1, padx=(10, 10), pady=8, sticky="ew")
        ctk.CTkLabel(
            card, text="After Activity",
            font=font(10), text_color=COLOR_TEXT_MUTED
        ).grid(row=3, column=0, padx=(18, 4), pady=8, sticky="w")
        ctk.CTkButton(
            card, text="Browse...", width=110, height=36, corner_radius=10,
            fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT, font=font(10, "bold"),
            command=self.browse_compare_after
        ).grid(row=3, column=2, padx=(0, 18), pady=8)




        self.compare_status_label = ctk.CTkLabel(
            card, text="Select both reports to begin.",
            font=font(10), text_color=COLOR_TEXT_MUTED, anchor="w"
        )
        self.compare_status_label.grid(row=4, column=0, columnspan=3, padx=18, pady=(12, 8), sticky="ew")




        ctk.CTkButton(
            card, text="COMPARE DIG RESULTS", height=46, corner_radius=12,
            fg_color=COLOR_ACCENT, hover_color=COLOR_ACCENT_HOVER,
            text_color="white", font=font(12, "bold"),
            command=self.compare_dig_reports
        ).grid(row=5, column=0, columnspan=3, padx=18, pady=(8, 18), sticky="ew")




    def browse_compare_before(self):
        path = filedialog.askopenfilename(
            title="Select BEFORE activity DIG Excel report",
            filetypes=[("Excel reports", "*.xlsx *.xlsm"), ("All files", "*.*")]
        )
        if path:
            self.compare_before_entry.delete(0, "end")
            self.compare_before_entry.insert(0, path)




    def browse_compare_after(self):
        path = filedialog.askopenfilename(
            title="Select AFTER activity DIG Excel report",
            filetypes=[("Excel reports", "*.xlsx *.xlsm"), ("All files", "*.*")]
        )
        if path:
            self.compare_after_entry.delete(0, "end")
            self.compare_after_entry.insert(0, path)




    def compare_dig_reports(self):
        before_file = self.compare_before_entry.get().strip()
        after_file = self.compare_after_entry.get().strip()
        if not before_file or not after_file:
            messagebox.showwarning("Select Reports", "Please select both BEFORE and AFTER Excel reports.", parent=self.root)
            return
        if not os.path.isfile(before_file) or not os.path.isfile(after_file):
            messagebox.showerror("File Not Found", "One or both selected Excel reports could not be found.", parent=self.root)
            return
        if os.path.abspath(before_file) == os.path.abspath(after_file):
            messagebox.showwarning("Select Two Reports", "Please select two different Excel reports.", parent=self.root)
            return
        try:
            report_path, changed, unchanged, total = create_dig_comparison_report(before_file, after_file)
            self.compare_status_label.configure(
                text=f"Compared {total} command entries: {changed} changed/missing, {unchanged} unchanged. Report: {report_path}",
                text_color=COLOR_SUCCESS if changed == 0 else COLOR_WARNING
            )
            messagebox.showinfo(
                "DIG Comparison Complete",
                f"Comparison completed successfully.\n\nTotal commands: {total}\nChanged / missing: {changed}\nUnchanged: {unchanged}\n\nReport saved to:\n{report_path}",
                parent=self.root
            )
        except Exception as exc:
            messagebox.showerror("DIG Comparison Failed", str(exc), parent=self.root)




    # ========================================================
    # FQDN SEARCH PAGE
    # ========================================================
    def build_fqdn_tab(self):
        page = ctk.CTkFrame(self.page_container, fg_color="transparent", corner_radius=0)
        self._pages["fqdn"] = page




        self.fqdn_server_var = ctk.StringVar(value="Current Connection")
        # Additive FQDN matching mode; Exact Match preserves the original search behavior.
        self.fqdn_match_mode_var = ctk.StringVar(value="Exact Match")
        self.fqdn_server_profiles = []
        self.fqdn_results_df = pd.DataFrame()
        self.fqdn_current_fqdn = ""
        self.fqdn_visible_columns = []
        self.fqdn_visible_rows = []
        self.fqdn_table = None




        outer = ctk.CTkScrollableFrame(page, fg_color="transparent", corner_radius=0)
        outer.pack(fill="both", expand=True)
        outer.grid_columnconfigure(0, weight=1)
        outer.grid_rowconfigure(0, weight=0)




        header = ctk.CTkFrame(outer, fg_color=COLOR_SURFACE, corner_radius=16, border_width=1, border_color=COLOR_BORDER)
        header.grid(row=0, column=0, sticky="ew", padx=24, pady=(20, 12))
        header.grid_columnconfigure(1, weight=1)




        ctk.CTkLabel(
            header,
            text="⌕   FQDN SEARCH + DIG",
            font=font(16, "bold"),
            text_color=COLOR_TEXT,
            anchor="w"
        ).grid(row=0, column=0, columnspan=4, padx=20, pady=(16, 2), sticky="w")




        ctk.CTkLabel(
            header,
            text="Search the existing CSV data, review forwarders, then run DIG against every IP in a selected row.",
            font=font(10),
            text_color=COLOR_TEXT_MUTED,
            anchor="w"
        ).grid(row=1, column=0, columnspan=4, padx=20, pady=(0, 14), sticky="w")




        ctk.CTkLabel(
            header, text="CSV Folder", font=font(10), text_color=COLOR_TEXT_MUTED
        ).grid(row=2, column=0, padx=(20, 10), pady=(4, 12), sticky="w")
        self.fqdn_folder_entry = self._entry(header)
        self.fqdn_folder_entry.grid(row=2, column=1, padx=(0, 10), pady=(4, 12), sticky="ew")
        ctk.CTkButton(
            header, text="Browse", width=100, height=38, corner_radius=10,
            fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT, font=font(10, "bold"),
            command=self.fqdn_select_folder
        ).grid(row=2, column=2, padx=(0, 20), pady=(4, 12))




        ctk.CTkLabel(
            header, text="FQDN", font=font(10), text_color=COLOR_TEXT_MUTED
        ).grid(row=3, column=0, padx=(20, 10), pady=(4, 12), sticky="w")
        self.fqdn_search_entry = self._entry(header)
        self.fqdn_search_entry.grid(row=3, column=1, padx=(0, 10), pady=(4, 12), sticky="ew")
        self.fqdn_search_entry.bind("<Return>", lambda event: self.fqdn_search())
        self.fqdn_match_mode_menu = ctk.CTkOptionMenu(
            header,
            variable=self.fqdn_match_mode_var,
            values=["Exact Match", "Starts With", "Ends With"],
            width=145,
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            button_color=COLOR_ACCENT,
            button_hover_color=COLOR_ACCENT_HOVER,
            dropdown_fg_color=COLOR_SURFACE_ALT,
            dropdown_hover_color=COLOR_ACCENT,
            font=font(10),
        )
        self.fqdn_match_mode_menu.grid(row=3, column=2, padx=(0, 10), pady=(4, 12))
        ctk.CTkButton(
            header, text="⌕  SEARCH", width=120, height=38, corner_radius=10,
            fg_color=COLOR_ACCENT, hover_color=COLOR_ACCENT_HOVER,
            text_color="#ffffff", font=font(10, "bold"),
            command=self.fqdn_search
        ).grid(row=3, column=3, padx=(0, 20), pady=(4, 12))




        ctk.CTkLabel(
            header, text="DIG Server", font=font(10), text_color=COLOR_TEXT_MUTED
        ).grid(row=4, column=0, padx=(20, 10), pady=(4, 14), sticky="w")
        self.fqdn_server_menu = ctk.CTkOptionMenu(
            header,
            variable=self.fqdn_server_var,
            values=["Current Connection"],
            width=280,
            height=38,
            corner_radius=10,
            fg_color=COLOR_SURFACE_ALT,
            button_color=COLOR_ACCENT,
            button_hover_color=COLOR_ACCENT_HOVER,
            dropdown_fg_color=COLOR_SURFACE_ALT,
            dropdown_hover_color=COLOR_ACCENT,
            font=font(11),
            command=self.fqdn_server_changed
        )
        self.fqdn_server_menu.grid(row=4, column=1, padx=(0, 10), pady=(4, 14), sticky="w")
        ctk.CTkButton(
            header, text="↻  REFRESH SERVERS", width=160, height=38, corner_radius=10,
            fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT, font=font(10, "bold"),
            command=self.refresh_fqdn_servers
        ).grid(row=4, column=2, padx=(0, 20), pady=(4, 14))




        results_card = ctk.CTkFrame(
            outer, fg_color=COLOR_SURFACE, corner_radius=16,
            border_width=1, border_color=COLOR_BORDER
        )
        results_card.grid(row=1, column=0, sticky="nsew", padx=24, pady=(0, 20))
        outer.grid_rowconfigure(1, weight=1)
        results_card.grid_rowconfigure(3, weight=1)
        results_card.grid_columnconfigure(0, weight=1)




        self.fqdn_result_summary = ctk.CTkLabel(
            results_card,
            text="No search performed.",
            font=font(10, "bold"),
            text_color=COLOR_TEXT_MUTED,
            anchor="w"
        )
        self.fqdn_result_summary.grid(row=1, column=0, padx=18, pady=(12, 6), sticky="ew")




        # Display controls bar
        display_bar = ctk.CTkFrame(results_card, fg_color="transparent")
        display_bar.grid(row=2, column=0, sticky="ew", padx=14, pady=(0, 8))
        display_bar.grid_columnconfigure(0, weight=1)
        display_bar.grid_columnconfigure(1, weight=0)
        display_bar.grid_columnconfigure(2, weight=0)
        display_bar.grid_columnconfigure(3, weight=0)




        self.fqdn_display_info = ctk.CTkLabel(
            display_bar,
            text="Columns: all    •    Rows: all",
            font=font(9),
            text_color=COLOR_TEXT_MUTED,
            anchor="w"
        )
        self.fqdn_display_info.grid(row=0, column=0, sticky="w", padx=(4, 8))




        self.fqdn_select_columns_button = ctk.CTkButton(
            display_bar, text="☷  SELECT COLUMNS", width=150, height=34,
            corner_radius=10, fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT, font=font(9, "bold"),
            command=self.fqdn_select_columns
        )
        self.fqdn_select_columns_button.grid(row=0, column=1, padx=4)




        self.fqdn_select_rows_button = ctk.CTkButton(
            display_bar, text="☑  SELECT ROWS", width=135, height=34,
            corner_radius=10, fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT, font=font(9, "bold"),
            command=self.fqdn_select_rows
        )
        self.fqdn_select_rows_button.grid(row=0, column=2, padx=4)




        self.fqdn_show_all_button = ctk.CTkButton(
            display_bar, text="SHOW ALL", width=95, height=34,
            corner_radius=10, fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT, font=font(9, "bold"),
            command=self.fqdn_show_all
        )
        self.fqdn_show_all_button.grid(row=0, column=3, padx=(4, 0))




        table_frame = ctk.CTkFrame(results_card, fg_color=COLOR_BG, corner_radius=10)
        table_frame.grid(row=3, column=0, sticky="nsew", padx=14, pady=(0, 12))
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)




        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure(
            "FQDN.Treeview",
            background=COLOR_BG,
            foreground=COLOR_TEXT,
            fieldbackground=COLOR_BG,
            rowheight=34,
            font=(FONT_FAMILY, 10)
        )
        style.configure(
            "FQDN.Treeview.Heading",
            background=COLOR_SURFACE_ALT,
            foreground=COLOR_TEXT,
            font=(FONT_FAMILY, 10, "bold"),
            relief="flat"
        )
        style.map(
            "FQDN.Treeview",
            background=[("selected", COLOR_ACCENT)],
            foreground=[("selected", "#ffffff")]
        )




        self.fqdn_table = ttk.Treeview(
            table_frame,
            show="headings",
            style="FQDN.Treeview",
            selectmode="extended"
        )
        self.fqdn_table.grid(row=0, column=0, sticky="nsew")
        self.fqdn_table.bind("<Button-1>", self.fqdn_table_click)




        y_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.fqdn_table.yview)
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.fqdn_table.xview)
        x_scroll.grid(row=1, column=0, sticky="ew")
        self.fqdn_table.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)




        bottom = ctk.CTkFrame(results_card, fg_color="transparent")
        bottom.grid(row=4, column=0, sticky="ew", padx=14, pady=(0, 14))
        bottom.grid_columnconfigure(0, weight=1)




        self.fqdn_export_button = ctk.CTkButton(
            bottom, text="⇩  EXPORT RESULTS", width=160, height=38,
            corner_radius=10, fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT, font=font(10, "bold"),
            command=self.fqdn_export
        )
        self.fqdn_export_button.grid(row=0, column=1, padx=(8, 0), sticky="e")




        self.refresh_fqdn_servers()




        try:
            page.bind("<Configure>", self._fqdn_tab_configure)
        except Exception:
            pass




    def _fqdn_tab_configure(self, event=None):
        try:
            if event is not None and event.widget is not self._pages.get("fqdn"):
                return
            if self._fqdn_resize_job is not None:
                try:
                    self.root.after_cancel(self._fqdn_resize_job)
                except Exception:
                    pass
            self._fqdn_resize_job = self.root.after(120, self._apply_fqdn_reflow)
        except Exception:
            pass




    def _apply_fqdn_reflow(self):
        self._fqdn_resize_job = None
        try:
            page = self._pages.get("fqdn")
            if page is None:
                return
            width = page.winfo_width()
            if width < 100:
                return
            if width == self._fqdn_last_width:
                return
            self._fqdn_last_width = width




            display_bar = self.fqdn_select_columns_button.master
            info_label = self.fqdn_display_info
            b1 = self.fqdn_select_columns_button
            b2 = self.fqdn_select_rows_button
            b3 = self.fqdn_show_all_button




            for widget in (info_label, b1, b2, b3):
                try:
                    widget.grid_forget()
                except Exception:
                    pass




            if width >= 760:
                display_bar.grid_columnconfigure(0, weight=1)
                display_bar.grid_columnconfigure(1, weight=0)
                display_bar.grid_columnconfigure(2, weight=0)
                display_bar.grid_columnconfigure(3, weight=0)
                info_label.grid(row=0, column=0, sticky="w", padx=(4, 8))
                b1.grid(row=0, column=1, padx=4)
                b2.grid(row=0, column=2, padx=4)
                b3.grid(row=0, column=3, padx=(4, 0))
                return




            display_bar.grid_columnconfigure(0, weight=1)
            display_bar.grid_columnconfigure(1, weight=0)
            display_bar.grid_columnconfigure(2, weight=0)
            display_bar.grid_columnconfigure(3, weight=0)
            info_label.grid(row=0, column=0, columnspan=4, sticky="w", padx=(4, 8))
            b1.grid(row=1, column=0, padx=4, pady=(6, 0), sticky="w")
            b2.grid(row=1, column=1, padx=4, pady=(6, 0), sticky="w")
            b3.grid(row=1, column=2, padx=(4, 0), pady=(6, 0), sticky="w")
        except Exception:
            pass




    def fqdn_update_display_info(self):
        total_cols = len(self.fqdn_results_df.columns) if not self.fqdn_results_df.empty else 0
        visible_cols = len(self.fqdn_visible_columns)
        total_rows = len(self.fqdn_results_df.index) if not self.fqdn_results_df.empty else 0
        visible_rows = len(self.fqdn_visible_rows)
        col_text = f"{visible_cols}/{total_cols}" if total_cols else "0"
        row_text = f"{visible_rows}/{total_rows}" if total_rows else "0"
        try:
            self.fqdn_display_info.configure(
                text=f"Columns: {col_text}    •    Rows: {row_text}    •    DIG is available in the last column"
            )
        except Exception:
            pass




    def fqdn_show_all(self):
        if self.fqdn_results_df.empty:
            return
        self.fqdn_visible_columns = list(self.fqdn_results_df.columns)
        self.fqdn_visible_rows = list(range(len(self.fqdn_results_df.index)))
        self.fqdn_render_table()
        self.fqdn_update_display_info()




    def fqdn_select_columns(self):
        if self.fqdn_results_df.empty:
            messagebox.showinfo("FQDN Display", "Search for an FQDN first.")
            return




        window = ctk.CTkToplevel(self.root)
        window.title("Select Columns to Display")
        window.geometry("560x620")
        window.minsize(480, 480)
        window.transient(self.root)
        window.configure(fg_color=COLOR_BG)




        ctk.CTkLabel(
            window,
            text="SELECT COLUMNS",
            font=font(15, "bold"),
            text_color=COLOR_TEXT
        ).pack(anchor="w", padx=20, pady=(18, 2))
        ctk.CTkLabel(
            window,
            text="Choose which FQDN result columns should be visible in the comparison table.",
            font=font(9),
            text_color=COLOR_TEXT_MUTED
        ).pack(anchor="w", padx=20, pady=(0, 12))




        selected = {
            column: ctk.BooleanVar(value=column in self.fqdn_visible_columns)
            for column in self.fqdn_results_df.columns
        }




        scroll = ctk.CTkScrollableFrame(window, fg_color=COLOR_SURFACE, corner_radius=12)
        scroll.pack(fill="both", expand=True, padx=18, pady=8)




        for index, column in enumerate(self.fqdn_results_df.columns):
            ctk.CTkCheckBox(
                scroll,
                text=str(column),
                variable=selected[column],
                font=font(10),
                text_color=COLOR_TEXT,
                checkbox_width=18,
                checkbox_height=18,
                corner_radius=5
            ).grid(row=index, column=0, padx=14, pady=7, sticky="w")




        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.pack(fill="x", padx=18, pady=(6, 18))




        def select_all():
            for var in selected.values():
                var.set(True)




        def clear_all():
            for var in selected.values():
                var.set(False)




        def apply():
            chosen = [column for column in self.fqdn_results_df.columns if selected[column].get()]
            if not chosen:
                messagebox.showwarning("FQDN Display", "Select at least one column.")
                return
            self.fqdn_visible_columns = chosen
            self.fqdn_render_table()
            self.fqdn_update_display_info()
            window.destroy()




        ctk.CTkButton(
            buttons, text="ALL", width=80, height=36,
            fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            corner_radius=10, command=select_all
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            buttons, text="CLEAR", width=80, height=36,
            fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            corner_radius=10, command=clear_all
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            buttons, text="APPLY", width=100, height=36,
            fg_color=COLOR_ACCENT, hover_color=COLOR_ACCENT_HOVER,
            text_color="#ffffff", font=font(10, "bold"),
            corner_radius=10, command=apply
        ).pack(side="right", padx=4)




    def fqdn_select_rows(self):
        if self.fqdn_results_df.empty:
            messagebox.showinfo("FQDN Display", "Search for an FQDN first.")
            return




        window = ctk.CTkToplevel(self.root)
        window.title("Select Rows to Display")
        window.geometry("700x620")
        window.minsize(560, 480)
        window.transient(self.root)
        window.configure(fg_color=COLOR_BG)




        ctk.CTkLabel(
            window,
            text="SELECT ROWS",
            font=font(15, "bold"),
            text_color=COLOR_TEXT
        ).pack(anchor="w", padx=20, pady=(18, 2))
        ctk.CTkLabel(
            window,
            text="Select the result rows you want to keep visible in the comparison table.",
            font=font(9),
            text_color=COLOR_TEXT_MUTED
        ).pack(anchor="w", padx=20, pady=(0, 12))




        selected = {
            index: ctk.BooleanVar(value=index in self.fqdn_visible_rows)
            for index in range(len(self.fqdn_results_df.index))
        }




        scroll = ctk.CTkScrollableFrame(window, fg_color=COLOR_SURFACE, corner_radius=12)
        scroll.pack(fill="both", expand=True, padx=18, pady=8)




        columns = list(self.fqdn_results_df.columns)
        preview_columns = columns[:4]




        for index, row in self.fqdn_results_df.iterrows():
            preview_parts = []
            for column in preview_columns:
                value = "" if pd.isna(row[column]) else str(row[column])
                if len(value) > 80:
                    value = value[:77] + "..."
                preview_parts.append(f"{column}: {value}")
            preview = " | ".join(preview_parts)




            ctk.CTkCheckBox(
                scroll,
                text=f"ROW {index + 1}   {preview}",
                variable=selected[index],
                font=font(9),
                text_color=COLOR_TEXT,
                checkbox_width=18,
                checkbox_height=18,
                corner_radius=5
            ).grid(row=index, column=0, padx=14, pady=7, sticky="w")




        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.pack(fill="x", padx=18, pady=(6, 18))




        def select_all():
            for var in selected.values():
                var.set(True)




        def clear_all():
            for var in selected.values():
                var.set(False)




        def apply():
            chosen = [index for index in range(len(self.fqdn_results_df.index)) if selected[index].get()]
            if not chosen:
                messagebox.showwarning("FQDN Display", "Select at least one row.")
                return
            self.fqdn_visible_rows = chosen
            self.fqdn_render_table()
            self.fqdn_update_display_info()
            window.destroy()




        ctk.CTkButton(
            buttons, text="ALL", width=80, height=36,
            fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            corner_radius=10, command=select_all
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            buttons, text="CLEAR", width=80, height=36,
            fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            corner_radius=10, command=clear_all
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            buttons, text="APPLY", width=100, height=36,
            fg_color=COLOR_ACCENT, hover_color=COLOR_ACCENT_HOVER,
            text_color="#ffffff", font=font(10, "bold"),
            corner_radius=10, command=apply
        ).pack(side="right", padx=4)




    def fqdn_render_table(self):
        if self.fqdn_table is None:
            return




        for item in self.fqdn_table.get_children():
            self.fqdn_table.delete(item)




        if self.fqdn_results_df.empty:
            self.fqdn_table["columns"] = ()
            return




        columns = list(self.fqdn_visible_columns or self.fqdn_results_df.columns)
        table_columns = [f"c{index}" for index in range(len(columns))] + ["__dig__"]
        self.fqdn_table["columns"] = table_columns




        for index, column in enumerate(columns):
            table_column = f"c{index}"
            self.fqdn_table.heading(table_column, text=str(column))
            sample_values = [
                "" if pd.isna(v) else str(v)
                for v in self.fqdn_results_df[column].tolist()
            ]
            max_len = max([len(str(column))] + [len(v) for v in sample_values[:100]])
            width = min(600, max(110, max_len * 7 + 24))
            if "allow" in str(column).lower():
                width = max(110, min(width, 150))
            self.fqdn_table.column(
                table_column,
                width=width,
                minwidth=90,
                anchor="w",
                stretch=False
            )




        self.fqdn_table.heading("__dig__", text="ACTION")
        self.fqdn_table.column("__dig__", width=110, minwidth=110, anchor="center", stretch=False)




        for dataframe_index in self.fqdn_visible_rows:
            if dataframe_index >= len(self.fqdn_results_df.index):
                continue
            row = self.fqdn_results_df.iloc[dataframe_index]
            values = ["" if pd.isna(row[column]) else str(row[column]) for column in columns]
            all_row_values = ["" if pd.isna(value) else str(value) for value in row.tolist()]
            ips = self.fqdn_extract_ips(all_row_values)
            action_text = f"▶ DIG ({len(ips)})" if ips else "NO IP"
            self.fqdn_table.insert(
                "",
                "end",
                iid=f"row_{dataframe_index}",
                values=values + [action_text]
            )




    def fqdn_table_click(self, event):
        if self.fqdn_table is None:
            return
        region = self.fqdn_table.identify("region", event.x, event.y)
        if region != "cell":
            return
        column_id = self.fqdn_table.identify_column(event.x)
        row_id = self.fqdn_table.identify_row(event.y)
        if not row_id or column_id != f"#{len(self.fqdn_visible_columns) + 1}":
            return
        try:
            dataframe_index = int(row_id.replace("row_", ""))
            row = self.fqdn_results_df.iloc[dataframe_index]
            row_values = ["" if pd.isna(value) else str(value) for value in row.tolist()]
            ips = self.fqdn_extract_ips(row_values)
            if ips:
                self.fqdn_run_row_dig(row_values, ips)
        except Exception as exc:
            messagebox.showerror("FQDN DIG", f"Unable to run row DIG:\n{exc}")




    def fqdn_select_folder(self):
        folder = filedialog.askdirectory(title="Select CSV Folder")
        if folder:
            self.fqdn_folder_entry.delete(0, "end")
            self.fqdn_folder_entry.insert(0, folder)




    def refresh_fqdn_servers(self):
        profiles = self.load_server_profiles()
        self.fqdn_server_profiles = profiles
        values = ["Current Connection"]
        for profile in profiles:
            name = profile.get("name", "").strip()
            server = profile.get("server", "").strip()
            label = f"{name}  ({server})" if name and server else (name or server)
            if label and label not in values:
                values.append(label)
        self.fqdn_server_menu.configure(values=values)
        if self.fqdn_server_var.get() not in values:
            self.fqdn_server_var.set("Current Connection")




    def fqdn_server_changed(self, value):
        if value == "Current Connection":
            return
        profile = self.fqdn_profile_from_selection(value)
        if not profile:
            return
        current_password = self.password_entry.get()
        self.server_entry.delete(0, "end")
        self.server_entry.insert(0, profile.get("server", ""))
        self.username_entry.delete(0, "end")
        self.username_entry.insert(0, profile.get("username", ""))
        self.port_entry.delete(0, "end")
        self.port_entry.insert(0, profile.get("port", "22"))
        self.auth_var.set(profile.get("authentication", "Password"))
        self.authentication_changed(self.auth_var.get())
        self.password_entry.delete(0, "end")
        self.password_entry.insert(0, current_password)
        self.key_entry.delete(0, "end")
        self.key_entry.insert(0, profile.get("key_file", ""))
        self.key_passphrase_entry.delete(0, "end")
        self.key_passphrase_entry.insert(0, profile.get("key_passphrase", ""))
        self.log(f"[FQDN] Selected DIG server profile: {profile.get('name', value)}")




    def fqdn_profile_from_selection(self, value):
        for profile in self.fqdn_server_profiles:
            name = profile.get("name", "").strip()
            server = profile.get("server", "").strip()
            label = f"{name}  ({server})" if name and server else (name or server)
            if label == value:
                return profile
        return None




    def fqdn_search(self):
        fqdn = self.fqdn_search_entry.get().strip()
        folder_path = self.fqdn_folder_entry.get().strip()




        if not folder_path:
            messagebox.showerror("FQDN Search", "Please select folder first")
            return
        if not fqdn:
            messagebox.showerror("FQDN Search", "Enter FQDN")
            return
        if not os.path.isdir(folder_path):
            messagebox.showerror("FQDN Search", f"Folder not found:\n{folder_path}")
            return




        required_columns = [0, 1, 2, 3, 4, 6, 8, 9, 10, 27, 30, 37, 39, 40]
        all_results = []
        for file in os.listdir(folder_path):
            if file.endswith(".csv"):
                file_path = os.path.join(folder_path, file)
                try:
                    df = pd.read_csv(file_path, dtype=str, low_memory=False)
                    if len(df.columns) < 2:
                        continue
                    # Keep the original exact-match behavior as the default.
                    # Starts/Ends With only change the FQDN matching criterion.
                    match_mode = self.fqdn_match_mode_var.get()
                    fqdn_column = df.iloc[:, 1].fillna("").astype(str)
                    if match_mode == "Starts With":
                        filtered = df[fqdn_column.str.startswith(fqdn, na=False)]
                    elif match_mode == "Ends With":
                        filtered = df[fqdn_column.str.endswith(fqdn, na=False)]
                    else:
                        filtered = df[df.iloc[:, 1] == fqdn]
                    if not filtered.empty:
                        selected = filtered.iloc[:, required_columns]
                        selected["Source_File"] = file
                        all_results.append(selected)
                except Exception as e:
                    print(f"Error reading {file}: {e}")




        if all_results:
            self.fqdn_results_df = pd.concat(all_results, ignore_index=True)
            self.fqdn_current_fqdn = fqdn
            self.fqdn_show_results()
        else:
            self.fqdn_results_df = pd.DataFrame()
            self.fqdn_current_fqdn = fqdn
            self.fqdn_clear_results()
            messagebox.showinfo("FQDN Search", "No match found")




    def fqdn_clear_results(self):
        if self.fqdn_table is not None:
            for item in self.fqdn_table.get_children():
                self.fqdn_table.delete(item)
            self.fqdn_table["columns"] = ()
        self.fqdn_visible_columns = []
        self.fqdn_visible_rows = []
        self.fqdn_result_summary.configure(text="No matching records found.")
        self.fqdn_update_display_info()




    def fqdn_extract_ips(self, row_values):
        ipv4_pattern = r"\b(?:25[0-5]|2[0-4]\d|1?\d?\d)(?:\.(?:25[0-5]|2[0-4]\d|1?\d?\d)){3}\b"
        ips = []
        for value in row_values:
            if value is None:
                continue
            text_value = str(value)
            for ip in re.findall(ipv4_pattern, text_value):
                if ip not in ips:
                    ips.append(ip)
        return ips




    def fqdn_show_results(self):
        self.fqdn_visible_columns = list(self.fqdn_results_df.columns)
        self.fqdn_visible_rows = list(range(len(self.fqdn_results_df.index)))
        self.fqdn_render_table()




        count = len(self.fqdn_results_df.index)
        self.fqdn_result_summary.configure(
            text=(
                f"FQDN: {self.fqdn_current_fqdn}    •    Matching rows: {count}    "
                f"•    Select columns/rows for comparison or click ▶ DIG in the ACTION column."
            )
        )
        self.fqdn_update_display_info()




    def fqdn_run_row_dig(self, row_values, ips):
        if not ips:
            return
        fqdn = self.fqdn_current_fqdn.strip()
        if not fqdn:
            messagebox.showerror("FQDN DIG", "No FQDN selected.")
            return




        try:
            server = self.server_entry.get().strip()
            username = self.username_entry.get().strip()
            port = self.port_entry.get().strip()
            authentication = self.auth_var.get()
            password = self.password_entry.get()
            key_file = self.key_entry.get().strip()
            key_passphrase = self.key_passphrase_entry.get()




            if not server or not username:
                raise Exception("Select a DIG server first.")
            if authentication == "Password" and not password:
                raise Exception("SSH password is required for the selected DIG server.")
            if authentication == "SSH Key" and not key_file:
                raise Exception("Please select a private key for the selected DIG server.")
            if authentication == "SSH Key" and not os.path.exists(key_file):
                raise Exception("Selected key does not exist.")
        except Exception as exc:
            messagebox.showerror("FQDN DIG", str(exc))
            return




        commands = [f"dig @{ip} {fqdn}" for ip in ips]




        # Add the same Live DIG Execution monitor used by the main DIG workflow.
        # This is UI-only: the existing FQDN commands and SSH execution remain unchanged.
        try:
            self._monitor_begin(
                [(server, command) for command in commands],
                title="Live DIG Execution"
            )
        except Exception as monitor_exc:
            self.log(f"[FQDN DIG] Unable to initialize Live DIG Execution monitor: {monitor_exc}")




        # Preserve the existing Live Output window as well.
        self.open_live_output()
        try:
            if self.live_output_window is not None and self.live_output_window.winfo_exists():
                self.live_output_window.deiconify()
                self.live_output_window.lift()
                self.live_output_window.focus_force()
                self.live_output_window.update_idletasks()
        except Exception as window_exc:
            self.log(f"[FQDN DIG] Live Output foreground warning: {window_exc}")
        self.log("\n" + "=" * 70)
        self.log("FQDN SEARCH TOOL - ROW DIG")
        self.log("=" * 70)
        self.log(f"FQDN    : {fqdn}")
        self.log(f"Server  : {server}")
        self.log(f"Commands: {len(commands)}")
        for command in commands:
            self.log(f"[FQDN DIG] {command}")




        def worker():
            try:
                output = execute_ssh(
                    server=server,
                    username=username,
                    port=port,
                    authentication=authentication,
                    password=password,
                    key_file=key_file,
                    key_passphrase=key_passphrase,
                    commands=commands,
                    log_callback=self.log,
                    keep_connection=False,
                    command_callback=self._monitor_command_event
                )
                if output:
                    self.log("[FQDN DIG] Row DIG completed.")
                else:
                    self.log("[FQDN DIG] No output returned.")
            except Exception as exc:
                self.log(f"[FQDN DIG] ERROR: {exc}")
            finally:
                try:
                    self.root.after(0, self._monitor_finish)
                except Exception as monitor_exc:
                    self.log(f"[FQDN DIG] Unable to finish Live DIG Execution monitor: {monitor_exc}")




        threading.Thread(target=worker, daemon=True).start()




    def fqdn_export(self):
        if self.fqdn_results_df.empty:
            messagebox.showerror("FQDN Search", "No data to export")
            return
        file_path = filedialog.asksaveasfilename(
            title="Export FQDN Search Results",
            defaultextension=".csv",
            filetypes=[("CSV Files", "*.csv"), ("All Files", "*.*")]
        )
        if file_path:
            self.fqdn_results_df.to_csv(file_path, index=False)
            messagebox.showinfo("FQDN Search", f"File saved: {file_path}")




    def authentication_changed(self, value):
        if value == "Password":
            self.key_label.grid_remove()
            self.key_entry.grid_remove()
            self.key_browse.grid_remove()
            self.key_passphrase_label.grid_remove()
            self.key_passphrase_entry.grid_remove()
            self.password_label.grid()
            self.password_entry.grid()
        else:
            self.password_label.grid_remove()
            self.password_entry.grid_remove()
            self.key_label.grid(row=3, column=0, padx=(18, 10), pady=4, sticky="w")
            self.key_entry.grid(row=3, column=1, padx=(0, 10), pady=4, sticky="ew")
            self.key_browse.grid(row=3, column=2, padx=(0, 18), pady=4, sticky="w")
            self.key_passphrase_label.grid(row=4, column=0, padx=(18, 10), pady=4, sticky="w")
            self.key_passphrase_entry.grid(row=4, column=1, columnspan=3, padx=(0, 18), pady=4, sticky="ew")




    # ========================================================
    # BROWSE COMMAND FILE
    # ========================================================
    def browse_commands_file(self):
        file_path = filedialog.askopenfilename(
            title="Select DIG Command File",
            filetypes=[
                ("All Files", "*.*"),
                ("Text Files", "*.txt"),
                ("Excel Files", "*.xlsx *.xlsm"),
                ("CSV Files", "*.csv")
            ]
        )
        if file_path:
            self.commands_file_entry.delete(0, "end")
            self.commands_file_entry.insert(0, file_path)
    
	
    # ========================================================
    # BROWSE KEY
    # ========================================================
    def browse_key(self):
        file_path = filedialog.askopenfilename(
            title="Select SSH Private Key",
            filetypes=[
                ("All Files", "*.*"),
                ("SSH Private Keys", "*.pem *.ppk"),
                ("PEM Files", "*.pem"),
                ("PuTTY Private Keys", "*.ppk")
            ]
        )
        if file_path:
            self.key_entry.delete(0, "end")
            self.key_entry.insert(0, file_path)
    # ========================================================
    # EFFECTIVE COMMAND INPUT
    # ========================================================
    def get_effective_commands(self):
        quick_input = self.quick_commands_text.get("1.0", "end").strip()
        if not quick_input:
            commands_file = self.commands_file_entry.get().strip()
            return read_commands(commands_file, self.trailing_option_entry.get().strip())




        commands_for_temp_file = []
        for raw_line in quick_input.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            if "\t" in raw_line:
                parts = raw_line.split("\t", 1)
                fqdn = parts[0].strip()
                forward_to = parts[1].strip()
                if fqdn and forward_to:
                    for target in forward_to.split(","):
                        target = target.strip()
                        if "/" not in target:
                            continue
                        ip = target.rsplit("/", 1)[1].strip()
                        if not ip:
                            continue
                        trailing_option = self.trailing_option_entry.get().strip()
                        command = f"dig @{ip} {fqdn}"
                        if trailing_option:
                            command += f" {trailing_option}"
                        commands_for_temp_file.append(command)
                    continue
            commands_for_temp_file.append(line)
        if not commands_for_temp_file:
            return []
        temp_path = None
        try:
            fd, temp_path = tempfile.mkstemp(prefix="infoblox_quick_commands_", suffix=".txt")
            with os.fdopen(fd, "w", encoding="utf-8") as temp_file:
                for command in commands_for_temp_file:
                    temp_file.write(command)
                    temp_file.write("\n")
            self.log("[COMMAND INPUT] Using commands/FQDN from Quick DIG input box.")
            return read_commands(temp_path, self.trailing_option_entry.get().strip())
        finally:
            if temp_path and os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except Exception:
                    pass




    # ========================================================
    # OUTPUT FOLDER
    # ========================================================
    def browse_output(self):
        folder = filedialog.askdirectory(title="Select Output Folder")
        if folder:
            self.output_entry.delete(0, "end")
            self.output_entry.insert(0, folder)




    # ========================================================
    # INTERACTIVE SSH TERMINAL
    # ========================================================
    def open_ssh_terminal(self):
        try:
            if self.terminal_window is not None and self.terminal_window.winfo_exists():
                self.terminal_window.deiconify()
                self.terminal_window.lift()
                self.terminal_window.focus_force()
                return
            with self.terminal_output_lock:
                self.terminal_output_buffer = ""
            self.terminal_window = ctk.CTkToplevel(self.root)
            self.terminal_window.title("INFOBLOX — Interactive SSH DIG Terminal")
            self.terminal_window.geometry("1100x700")
            self.terminal_window.minsize(750, 450)
            self.terminal_window.configure(fg_color=COLOR_BG)
            self.terminal_window.protocol("WM_DELETE_WINDOW", self.close_ssh_terminal)
            self.terminal_window.grid_rowconfigure(2, weight=1)
            self.terminal_window.grid_columnconfigure(0, weight=1)




            header = ctk.CTkFrame(self.terminal_window, fg_color=COLOR_SURFACE, corner_radius=12)
            header.grid(row=0, column=0, sticky="ew", padx=14, pady=(14, 7))
            header.grid_columnconfigure(0, weight=1)




            self.terminal_status_label = ctk.CTkLabel(
                header, text="●  Disconnected", font=font(12, "bold"),
                text_color=COLOR_TEXT_MUTED, anchor="w"
            )
            self.terminal_status_label.grid(row=0, column=0, padx=16, pady=12, sticky="w")




            self.terminal_connect_button = ctk.CTkButton(
                header, text="CONNECT", width=110, height=34,
                corner_radius=10, fg_color=COLOR_SUCCESS, hover_color=COLOR_SUCCESS_HOVER,
                text_color="#ffffff", font=font(11, "bold"), command=self.connect_ssh_terminal
            )
            self.terminal_connect_button.grid(row=0, column=1, padx=(6, 8), pady=10)




            self.terminal_disconnect_button = ctk.CTkButton(
                header, text="DISCONNECT", width=120, height=34,
                corner_radius=10, fg_color=COLOR_DANGER, hover_color=COLOR_DANGER_HOVER,
                text_color="#ffffff", font=font(11, "bold"),
                command=self.disconnect_ssh_terminal, state="disabled"
            )
            self.terminal_disconnect_button.grid(row=0, column=2, padx=(0, 8), pady=10)




            self.terminal_save_button = ctk.CTkButton(
                header, text="SAVE TXT", width=100, height=34,
                corner_radius=10, fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
                text_color=COLOR_TEXT, font=font(11, "bold"), command=self.save_ssh_terminal_output
            )
            self.terminal_save_button.grid(row=0, column=3, padx=(0, 16), pady=10)




            # Terminal-only connection profile selector. This does not alter
            # the main window's active connection fields.
            profile_frame = ctk.CTkFrame(self.terminal_window, fg_color=COLOR_SURFACE, corner_radius=12)
            profile_frame.grid(row=1, column=0, sticky="ew", padx=14, pady=7)
            profile_frame.grid_columnconfigure(1, weight=1)
            ctk.CTkLabel(
                profile_frame, text="Server Profile", font=font(10, "bold"),
                text_color=COLOR_TEXT_MUTED
            ).grid(row=0, column=0, padx=(14, 8), pady=10, sticky="w")




            self.terminal_profiles = self.load_server_profiles()
            self._terminal_profile_lookup = {}
            profile_values = ["Main Window Connection"]
            for profile in self.terminal_profiles:
                name = str(profile.get("name", "")).strip()
                server = str(profile.get("server", "")).strip()
                label = f"{name} ({server})" if name and server else (name or server)
                if label and label not in self._terminal_profile_lookup:
                    self._terminal_profile_lookup[label] = profile
                    profile_values.append(label)
            self.terminal_profile_var = ctk.StringVar(value="Main Window Connection")
            self.terminal_profile_menu = ctk.CTkOptionMenu(
                profile_frame, variable=self.terminal_profile_var, values=profile_values,
                height=34, corner_radius=9, fg_color=COLOR_SURFACE_ALT,
                button_color=COLOR_ACCENT, button_hover_color=COLOR_ACCENT_HOVER,
                dropdown_fg_color=COLOR_SURFACE_ALT, dropdown_hover_color=COLOR_ACCENT,
                font=font(11), command=self.terminal_profile_changed
            )
            self.terminal_profile_menu.grid(row=0, column=1, padx=(0, 14), pady=8, sticky="ew")
            ctk.CTkLabel(
                profile_frame, text="Trailing option", font=font(10),
                text_color=COLOR_TEXT_MUTED
            ).grid(row=0, column=2, padx=(8, 6), pady=10, sticky="e")
            self.terminal_trailing_option_entry = self._entry(profile_frame, width=150)
            self.terminal_trailing_option_entry.grid(row=0, column=3, padx=(0, 14), pady=8, sticky="e")
            try:
                self.terminal_trailing_option_entry.insert(0, self.trailing_option_entry.get().strip())
            except Exception:
                pass




            self.terminal_text = ctk.CTkTextbox(
                self.terminal_window, corner_radius=12, fg_color="#080b0f",
                border_width=1, border_color=COLOR_BORDER, text_color=COLOR_TEXT,
                font=mono_font(12), wrap="none"
            )
            self.terminal_text.grid(row=2, column=0, sticky="nsew", padx=14, pady=7)




            command_frame = ctk.CTkFrame(self.terminal_window, fg_color=COLOR_SURFACE, corner_radius=12)
            command_frame.grid(row=3, column=0, sticky="ew", padx=14, pady=(7, 14))
            command_frame.grid_columnconfigure(0, weight=1)
            self.terminal_command_entry = ctk.CTkEntry(
                command_frame, height=40, corner_radius=10, fg_color=COLOR_BG,
                border_width=1, border_color=COLOR_BORDER, text_color=COLOR_TEXT,
                font=mono_font(12), placeholder_text="Enter an FQDN, subnet, or DIG command only..."
            )
            self.terminal_command_entry.grid(row=0, column=0, padx=(14, 8), pady=14, sticky="ew")




            try:
                entry_widget = self.terminal_command_entry._entry
                terminal_text_widget = None
                for child in self.terminal_text.winfo_children():
                    try:
                        if child.winfo_class() == "Text":
                            terminal_text_widget = child
                            break
                    except Exception:
                        pass




                def terminal_enter(event=None):
                    return self.send_ssh_terminal_command(event)




                entry_widget.bind("<Return>", terminal_enter, add="+")
                entry_widget.bind("<KP_Enter>", terminal_enter, add="+")
                if terminal_text_widget is not None:
                    terminal_text_widget.bind("<Return>", terminal_enter, add="+")
                    terminal_text_widget.bind("<KP_Enter>", terminal_enter, add="+")
                self.terminal_window.after(100, entry_widget.focus_set)
            except Exception as exc:
                print("Enter binding failed:", exc)




            ctk.CTkButton(
                command_frame, text="RUN DIG", width=90, height=40, corner_radius=10,
                fg_color=COLOR_ACCENT, hover_color=COLOR_ACCENT_HOVER,
                font=font(11, "bold"), command=self.send_ssh_terminal_command
            ).grid(row=0, column=1, padx=(0, 14), pady=14)




            self.terminal_write(
                "Interactive SSH DIG Terminal\n"
                "Choose a saved Server Profile or use Main Window Connection.\n"
                "Only DIG queries are accepted. Enter an FQDN, a supported IPv4 subnet, or a DIG command.\n"
                "Bare FQDNs/subnets are prefixed with 'dig'; existing DIG commands are sent unchanged.\n\n"
            )
        except Exception as exc:
            self.log(f"[SSH TERMINAL] Unable to open terminal: {exc}")




    def terminal_profile_changed(self, value):
        if value and value != "Main Window Connection":
            self.terminal_write(f"\n[SSH] Selected server profile: {value}\n")




    def prepare_terminal_dig_command(self, raw_command):
        """Normalize input to a DIG-only command and reject shell/CLI commands."""
        raw_command = (raw_command or "").strip()
        if not raw_command:
            return ""
        forbidden_chars = (";", "|", "&", ">", "<", "`", "$", chr(92), "'", '"', "(", ")", "\n", "\r")
        if any(ch in raw_command for ch in forbidden_chars):
            raise ValueError("Only a single DIG query is allowed. Shell operators and command chaining are blocked.")




        if re.match(r"^dig\b", raw_command, flags=re.IGNORECASE):
            command = convert_cidr_tokens_in_dig_command(raw_command)
        else:
            if len(raw_command.split()) != 1:
                raise ValueError("Enter one FQDN/subnet, or provide a complete DIG command beginning with 'dig'.")
            command = convert_cidr_tokens_in_dig_command("dig " + raw_command)
            trailing_option = ""
            try:
                trailing_option = self.terminal_trailing_option_entry.get().strip()
            except Exception:
                pass
            if trailing_option:
                forbidden_option_chars = (";", "|", "&", ">", "<", "`", "$", chr(92), "'", '"', "(", ")", "\n", "\r")
                if any(ch in trailing_option for ch in forbidden_option_chars):
                    raise ValueError("The trailing option contains unsupported characters.")
                command += " " + trailing_option




        if not re.match(r"^dig(?:\s|$)", command, flags=re.IGNORECASE):
            raise ValueError("Only DIG commands are allowed.")
        allowed_token = re.compile(r"^[A-Za-z0-9._:/%+=@*?\-]+$")
        if any(not allowed_token.fullmatch(token) for token in command.split()):
            raise ValueError("Only DIG syntax is allowed. Shell commands, separators, and unsupported characters are blocked.")
        return command




    def terminal_write(self, text):
        if not text:
            return
        cleaned = clean_terminal_text(text)
        with self.terminal_output_lock:
            self.terminal_output_buffer += cleaned




        def update():
            try:
                if self.terminal_text is not None and self.terminal_text.winfo_exists():
                    self.terminal_text.insert("end", cleaned)
                    self.terminal_text.see("end")
            except Exception:
                pass
        try:
            self.root.after(0, update)
        except Exception:
            pass




    def connect_ssh_terminal(self):
        if self.terminal_connected:
            self.terminal_write("\n[SSH] Already connected.\n")
            return
        selected_profile = self.terminal_profile_var.get() if self.terminal_profile_var is not None else "Main Window Connection"
        if selected_profile == "Main Window Connection":
            server = self.server_entry.get().strip()
            username = self.username_entry.get().strip()
            port = self.port_entry.get().strip() or "22"
            authentication = self.auth_var.get()
            password = self.password_entry.get()
            key_file = self.key_entry.get().strip()
            key_passphrase = self.key_passphrase_entry.get()
        else:
            profile = getattr(self, "_terminal_profile_lookup", {}).get(selected_profile)
            if not profile:
                messagebox.showwarning("SSH Terminal", "Selected server profile could not be found.", parent=self.terminal_window)
                return
            server = str(profile.get("server", "")).strip()
            username = str(profile.get("username", "")).strip()
            port = str(profile.get("port", "22")).strip() or "22"
            authentication = str(profile.get("authentication", "Password"))
            password = ""
            key_file = str(profile.get("key_file", "")).strip()
            key_passphrase = str(profile.get("key_passphrase", ""))
            if authentication == "Password":
                password = simpledialog.askstring(
                    "SSH Profile Password",
                    f"Enter SSH password for profile '{selected_profile}':",
                    show="*", parent=self.terminal_window
                )
                if password is None:
                    return
        if not server or not username:
            messagebox.showwarning("SSH Terminal", "The selected connection/profile must include Server / IP and Username.", parent=self.terminal_window)
            return
        try:
            port_value = int(port)
        except ValueError:
            messagebox.showwarning("SSH Terminal", "SSH Port must be a valid number.", parent=self.terminal_window)
            return




        self.terminal_connect_button.configure(state="disabled")
        self.terminal_status_label.configure(text="●  Connecting...", text_color=COLOR_WARNING)
        self.terminal_write(f"\n[SSH] Connecting to {server}:{port_value} using {selected_profile}...\n")




        def worker():
            client = None
            temp_key_dir = None
            try:
                client = paramiko.SSHClient()
                client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
                if authentication == "Password":
                    client.connect(
                        hostname=server, port=port_value, username=username,
                        password=password, timeout=20, banner_timeout=20,
                        auth_timeout=20, look_for_keys=False, allow_agent=False
                    )
                else:
                    private_key, temp_key_dir = load_private_key(key_file, key_passphrase)
                    client.connect(
                        hostname=server, port=port_value, username=username,
                        pkey=private_key, timeout=20, banner_timeout=20,
                        auth_timeout=20, look_for_keys=False, allow_agent=False
                    )
                channel = client.invoke_shell(term="xterm", width=160, height=200)
                channel.settimeout(1.0)
                self.terminal_client = client
                self.terminal_channel = channel
                self.terminal_temp_key_dir = temp_key_dir
                self.terminal_connected = True
                self.terminal_write(f"[SSH] Connected to {server}\n")




                def connected_ui():
                    try:
                        self.terminal_status_label.configure(
                            text=f"●  Connected: {server}", text_color=COLOR_SUCCESS
                        )
                        self.terminal_connect_button.configure(state="disabled")
                        self.terminal_disconnect_button.configure(state="normal")
                        self.terminal_window.focus_force()
                        command_entry_widget = self.terminal_command_entry._entry
                        command_entry_widget.focus_force()
                        self.terminal_window.after(100, command_entry_widget.focus_force)
                        self.terminal_window.after(300, command_entry_widget.focus_force)
                    except Exception:
                        pass
                self.root.after(0, connected_ui)




                def reader():
                    try:
                        while self.terminal_connected:
                            try:
                                if channel.recv_ready():
                                    data = channel.recv(65535)
                                    if not data:
                                        break
                                    decoded_data = data.decode("utf-8", errors="replace")
                                    self.terminal_write(decoded_data)
                                    # Keep the DIG-only input restriction while handling
                                    # Infoblox's pager internally, as the main DIG runner does.
                                    if re.search(
                                        r"Enter <return> for next page or q<return> to cancel the command\.?",
                                        clean_terminal_text(decoded_data),
                                        flags=re.IGNORECASE
                                    ):
                                        try:
                                            channel.send("q\n")
                                            self.terminal_write("\n[AUTO] Infoblox pager detected; cancelled the page prompt automatically.\n")
                                        except Exception as pager_exc:
                                            self.terminal_write(f"\n[SSH] Unable to handle pager prompt: {pager_exc}\n")
                                else:
                                    time.sleep(0.05)
                            except Exception:
                                if not self.terminal_connected:
                                    break
                                time.sleep(0.1)
                    finally:
                        if self.terminal_connected:
                            self.root.after(0, self._terminal_connection_lost)




                self.terminal_reader_thread = threading.Thread(target=reader, daemon=True)
                self.terminal_reader_thread.start()
            except Exception as exc:
                if temp_key_dir:
                    shutil.rmtree(temp_key_dir, ignore_errors=True)
                try:
                    if client:
                        client.close()
                except Exception:
                    pass
                self.terminal_client = None
                self.terminal_channel = None
                self.terminal_connected = False
                self.terminal_write(f"[SSH ERROR] {exc}\n")




                def failed_ui():
                    try:
                        self.terminal_status_label.configure(
                            text="●  Connection failed", text_color=COLOR_DANGER
                        )
                        self.terminal_connect_button.configure(state="normal")
                    except Exception:
                        pass
                self.root.after(0, failed_ui)
        threading.Thread(target=worker, daemon=True).start()




    def send_ssh_terminal_command(self, event=None):
        if not self.terminal_connected or self.terminal_channel is None:
            self.terminal_write("\n[SSH] Not connected. Click CONNECT first.\n")
            return "break"
        try:
            raw_command = self.terminal_command_entry.get()
            if not raw_command.strip():
                return "break"
            try:
                command = self.prepare_terminal_dig_command(raw_command)
            except ValueError as exc:
                self.terminal_write(f"\n[BLOCKED] {exc}\n")
                return "break"
            self.terminal_command_entry.delete(0, "end")
            self.terminal_write(f"\n[LOCAL DIG] {command}\n")
            self.terminal_channel.send(command + "\n")
        except Exception as exc:
            self.terminal_write(f"\n[SSH ERROR] Failed to send DIG command: {exc}\n")
        return "break"




    def disconnect_ssh_terminal(self):
        self.terminal_connected = False
        channel = self.terminal_channel
        client = self.terminal_client
        temp_key_dir = self.terminal_temp_key_dir
        self.terminal_channel = None
        self.terminal_client = None
        self.terminal_temp_key_dir = None
        try:
            if channel:
                try:
                    channel.send("exit\n")
                except Exception:
                    pass
                try:
                    channel.close()
                except Exception:
                    pass
        finally:
            try:
                if client:
                    client.close()
            except Exception:
                pass
            if temp_key_dir:
                shutil.rmtree(temp_key_dir, ignore_errors=True)
        self.terminal_write("\n[SSH] Disconnected.\n")
        try:
            self.terminal_status_label.configure(text="●  Disconnected", text_color=COLOR_TEXT_MUTED)
            self.terminal_connect_button.configure(state="normal")
            self.terminal_disconnect_button.configure(state="disabled")
        except Exception:
            pass




    def _terminal_connection_lost(self):
        if not self.terminal_connected:
            return
        self.terminal_connected = False
        self.terminal_channel = None
        self.terminal_client = None
        try:
            self.terminal_status_label.configure(text="●  Connection lost", text_color=COLOR_DANGER)
            self.terminal_connect_button.configure(state="normal")
            self.terminal_disconnect_button.configure(state="disabled")
        except Exception:
            pass
        self.terminal_write("\n[SSH] Connection closed by remote server.\n")




    def save_ssh_terminal_output(self):
        with self.terminal_output_lock:
            output = self.terminal_output_buffer
        if not output.strip():
            messagebox.showinfo("Save Terminal Output", "There is no terminal output to save.")
            return
        default_dir = self.output_entry.get().strip() or OUTPUT_DIR
        try:
            os.makedirs(default_dir, exist_ok=True)
        except Exception:
            default_dir = OUTPUT_DIR
        filename = "infoblox_ssh_terminal_" + datetime.datetime.now().strftime("%Y%m%d_%H%M%S") + ".txt"
        file_path = filedialog.asksaveasfilename(
            title="Save SSH Terminal Output",
            initialdir=default_dir,
            initialfile=filename,
            defaultextension=".txt",
            filetypes=[("Text files", "*.txt"), ("All files", "*.*")]
        )
        if not file_path:
            return
        try:
            with open(file_path, "w", encoding="utf-8") as file:
                file.write(output)
            self.log(f"[SSH TERMINAL] Output saved: {file_path}")
            messagebox.showinfo("Save Terminal Output", f"Output saved successfully:\n\n{file_path}")
        except Exception as exc:
            messagebox.showerror("Save Terminal Output", f"Unable to save terminal output:\n\n{exc}")




    def close_ssh_terminal(self):
        try:
            bindtag = getattr(self, "_terminal_enter_bindtag", None)
            if bindtag and self.terminal_window is not None:
                self.terminal_window.unbind_class(bindtag, "<Return>")
                self.terminal_window.unbind_class(bindtag, "<KP_Enter>")
            self._terminal_enter_bindtag = None
        except Exception:
            pass
        self.disconnect_ssh_terminal()
        try:
            if self.terminal_window is not None and self.terminal_window.winfo_exists():
                self.terminal_window.destroy()
        except Exception:
            pass
        self.terminal_window = None
        self.terminal_text = None
        self.terminal_command_entry = None
        self.terminal_status_label = None
        self.terminal_connect_button = None
        self.terminal_disconnect_button = None




    # ========================================================
    # LIVE DIG EXECUTION MONITOR (ADDITIVE UI ONLY)
    # ========================================================
    @staticmethod
    def _monitor_domain_from_command(command):
        parts = (command or "").strip().split()
        if parts and parts[0].lower() == "dig":
            parts = parts[1:]
        for index, token in enumerate(parts):
            if token.startswith("@") or token.startswith("+") or token.startswith("-"):
                if token == "-x" and index + 1 < len(parts):
                    return parts[index + 1]
                continue
            return token
        return (command or "").strip()




    def _monitor_begin(self, server_commands, title="Live DIG Execution Monitor"):
        """Initialize a live status table for the current run."""
        now = time.time()
        with self.dig_monitor_lock:
            self.dig_monitor_started_at = now
            self.dig_monitor_finished_at = None
            self.dig_monitor_rows = []
            for index, (server, command) in enumerate(server_commands):
                self.dig_monitor_rows.append({
                    "id": str(index),
                    "server": str(server or ""),
                    "command": normalize_command(command),
                    "domain": self._monitor_domain_from_command(command),
                    "status": "PENDING",
                    "start": "",
                    "end": "",
                    "duration": "",
                    "started_ts": None,
                    "ended_ts": None,
                })
        self.root.after(0, lambda t=title: self._monitor_build_window(t))
        self.root.after(0, self._monitor_refresh)
        self._monitor_schedule_timer()




    def _monitor_build_window(self, title):
        try:
            if self.dig_monitor_window is not None and self.dig_monitor_window.winfo_exists():
                self.dig_monitor_window.destroy()
        except Exception:
            pass
        try:
            win = ctk.CTkToplevel(self.root)
            self.dig_monitor_window = win
            win.title(title)
            win.geometry("1380x760")
            win.minsize(960, 520)
            win.configure(fg_color=COLOR_BG)
            win.grid_columnconfigure(0, weight=1)
            win.grid_rowconfigure(3, weight=1)




            header = ctk.CTkFrame(win, fg_color=COLOR_SURFACE, corner_radius=14,
                                  border_width=1, border_color=COLOR_BORDER)
            header.grid(row=0, column=0, sticky="ew", padx=14, pady=(14, 8))
            header.grid_columnconfigure(0, weight=1)
            ctk.CTkLabel(header, text="LIVE DIG EXECUTION",
                         font=font(15, "bold"), text_color=COLOR_TEXT).grid(
                row=0, column=0, sticky="w", padx=16, pady=(12, 4))
            self.dig_monitor_summary = ctk.CTkLabel(
                header, text="Preparing commands...", font=font(10, "bold"),
                text_color=COLOR_TEXT_MUTED, anchor="w")
            self.dig_monitor_summary.grid(row=1, column=0, sticky="ew", padx=16, pady=2)




            meta = ctk.CTkFrame(header, fg_color="transparent")
            meta.grid(row=2, column=0, sticky="ew", padx=16, pady=(4, 12))
            self.dig_monitor_started_label = ctk.CTkLabel(
                meta, text="Start: —", font=font(10), text_color=COLOR_TEXT_MUTED)
            self.dig_monitor_started_label.pack(side="left", padx=(0, 18))
            self.dig_monitor_elapsed_label = ctk.CTkLabel(
                meta, text="Elapsed: 00:00:00", font=font(10, "bold"), text_color=COLOR_TEXT)
            self.dig_monitor_elapsed_label.pack(side="left", padx=(0, 18))
            self.dig_monitor_eta_label = ctk.CTkLabel(
                meta, text="Estimated remaining: calculating…", font=font(10, "bold"),
                text_color=COLOR_WARNING)
            self.dig_monitor_eta_label.pack(side="left", padx=(0, 18))
            ctk.CTkButton(
                meta, text="Stop DIG", width=100, height=30,
                fg_color=COLOR_DANGER, hover_color=COLOR_DANGER_HOVER,
                command=self._monitor_stop_dig
            ).pack(side="right", padx=(8, 0))
            ctk.CTkButton(
                meta, text="Failed Queries", width=120, height=30,
                fg_color=COLOR_WARNING, hover_color=COLOR_DANGER_HOVER,
                command=self._monitor_open_failed_view
            ).pack(side="right", padx=(8, 0))
            ctk.CTkButton(
                meta, text="Export Live CSV", width=120, height=30,
                fg_color=COLOR_ACCENT, hover_color=COLOR_ACCENT_HOVER,
                command=self._monitor_export_csv
            ).pack(side="right")




            progress_frame = ctk.CTkFrame(win, fg_color="transparent")
            progress_frame.grid(row=1, column=0, sticky="ew", padx=16, pady=(0, 6))
            self.dig_monitor_progress = ctk.CTkProgressBar(
                progress_frame, height=12, corner_radius=6,
                fg_color=COLOR_SURFACE_ALT, progress_color=COLOR_SUCCESS)
            self.dig_monitor_progress.set(0)
            self.dig_monitor_progress.pack(fill="x", expand=True)




            filters = ctk.CTkFrame(win, fg_color=COLOR_SURFACE, corner_radius=10,
                                   border_width=1, border_color=COLOR_BORDER)
            filters.grid(row=2, column=0, sticky="ew", padx=14, pady=(0, 8))
            filters.grid_columnconfigure(0, weight=1)
            self.dig_monitor_search_var = tk.StringVar(value="")
            self.dig_monitor_status_var = tk.StringVar(value="All statuses")
            ctk.CTkLabel(filters, text="Search server / domain / command:",
                         font=font(10), text_color=COLOR_TEXT_MUTED).pack(
                side="left", padx=(12, 6), pady=8)
            search_entry = ctk.CTkEntry(
                filters, textvariable=self.dig_monitor_search_var, width=300,
                placeholder_text="Type to filter results…")
            search_entry.pack(side="left", padx=(0, 14), pady=8)
            ctk.CTkLabel(filters, text="Status:", font=font(10),
                         text_color=COLOR_TEXT_MUTED).pack(side="left", padx=(0, 6))
            status_menu = ctk.CTkOptionMenu(
                filters, variable=self.dig_monitor_status_var, width=190,
                values=["All statuses", "SERVFAIL", "NXDOMAIN", "Timeouts",
                        "IN PROGRESS", "Pending", "Failed"],
                command=lambda _value: self._monitor_apply_filters())
            status_menu.pack(side="left", padx=(0, 10), pady=8)
            self.dig_monitor_search_var.trace_add("write", lambda *_args: self._monitor_apply_filters())
            ctk.CTkButton(
                filters, text="Clear Filters", width=100, height=30,
                fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
                command=self._monitor_clear_filters
            ).pack(side="left", padx=(0, 10), pady=8)




            table_frame = ctk.CTkFrame(win, fg_color=COLOR_SURFACE, corner_radius=12,
                                       border_width=1, border_color=COLOR_BORDER)
            table_frame.grid(row=3, column=0, sticky="nsew", padx=14, pady=(0, 14))
            table_frame.grid_columnconfigure(0, weight=1)
            table_frame.grid_rowconfigure(0, weight=1)
            columns = ("server", "domain", "status", "start", "end", "duration", "command")
            tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=18)
            self.dig_monitor_tree = tree
            labels = {"server": "Server", "domain": "Domain / Query", "status": "Status",
                      "start": "Start Time", "end": "End Time", "duration": "Duration",
                      "command": "DIG Command"}
            widths = {"server": 160, "domain": 180, "status": 150, "start": 145,
                      "end": 145, "duration": 95, "command": 420}
            for col in columns:
                tree.heading(col, text=labels[col])
                tree.column(col, width=widths[col], minwidth=80,
                            anchor="w" if col in ("server", "domain", "command") else "center")
            style = ttk.Style()
            try:
                style.theme_use("clam")
            except Exception:
                pass
            style.configure("DigMonitor.Treeview", background=COLOR_SURFACE_ALT,
                            foreground=COLOR_TEXT, fieldbackground=COLOR_SURFACE_ALT,
                            rowheight=28, font=(FONT_FAMILY, 9), borderwidth=0)
            style.configure("DigMonitor.Treeview.Heading", background=COLOR_BORDER,
                            foreground=COLOR_TEXT, font=(FONT_FAMILY, 9, "bold"), relief="flat")
            tree.configure(style="DigMonitor.Treeview")
            tree.tag_configure("success", foreground="#22c55e")
            tree.tag_configure("running", foreground="#f59e0b")
            tree.tag_configure("failed", foreground="#ef4444")
            tree.tag_configure("pending", foreground=COLOR_TEXT_MUTED)
            yscroll = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
            xscroll = ttk.Scrollbar(table_frame, orient="horizontal", command=tree.xview)
            tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
            tree.grid(row=0, column=0, sticky="nsew", padx=(8, 0), pady=(8, 0))
            yscroll.grid(row=0, column=1, sticky="ns", pady=(8, 0), padx=(0, 8))
            xscroll.grid(row=1, column=0, sticky="ew", padx=(8, 0), pady=(0, 8))
            with self.dig_monitor_lock:
                rows = list(self.dig_monitor_rows)
            for row in rows:
                tree.insert("", "end", iid=row["id"], values=(
                    row["server"], row["domain"], row["status"], row["start"],
                    row["end"], row["duration"], row["command"]
                ), tags=("pending",))
            win.protocol("WM_DELETE_WINDOW", win.withdraw)
            # Ensure this monitor is visible when launched from FQDN Search Action.
            try:
                win.deiconify()
                win.lift()
                win.attributes("-topmost", True)
                win.after(900, lambda w=win: w.attributes("-topmost", False)
                          if w.winfo_exists() else None)
                win.focus_force()
            except Exception:
                pass
            self._monitor_refresh()
        except Exception as exc:
            self.log(f"[MONITOR] Unable to open live DIG monitor: {exc}")




    def _monitor_stop_dig(self):
        """Stop the active run using the existing cancellation pathway."""
        if self.dig_monitor_finished_at is not None or STOP_DIG_EVENT.is_set():
            return
        if messagebox.askyesno("Stop DIG", "Stop the current DIG run and close active SSH channels?",
                               parent=self.dig_monitor_window):
            self.stop_dig()




    def _monitor_clear_filters(self):
        if self.dig_monitor_search_var is not None:
            self.dig_monitor_search_var.set("")
        if self.dig_monitor_status_var is not None:
            self.dig_monitor_status_var.set("All statuses")
        self._monitor_apply_filters()




    def _monitor_apply_filters(self):
        try:
            if self.dig_monitor_tree is None or not self.dig_monitor_tree.winfo_exists():
                return
            query = (self.dig_monitor_search_var.get() if self.dig_monitor_search_var else "").strip().lower()
            status_filter = self.dig_monitor_status_var.get() if self.dig_monitor_status_var else "All statuses"
            with self.dig_monitor_lock:
                rows = [row.copy() for row in self.dig_monitor_rows]
            for row in rows:
                status = row["status"].upper()
                searchable = " ".join((row["server"], row["domain"], row["command"], status)).lower()
                matches_query = not query or query in searchable
                if status_filter == "All statuses":
                    matches_status = True
                elif status_filter == "Timeouts":
                    matches_status = "TIMEOUT" in status or "TIMED OUT" in status
                elif status_filter == "Pending":
                    matches_status = status in ("PENDING", "NO RESULT")
                elif status_filter == "Failed":
                    matches_status = status not in ("NOERROR", "IN PROGRESS", "PENDING", "NO RESULT", "STOPPED")
                else:
                    matches_status = status == status_filter.upper()
                visible = matches_query and matches_status
                exists = self.dig_monitor_tree.exists(row["id"])
                if visible and not exists:
                    self.dig_monitor_tree.insert("", "end", iid=row["id"], values=(
                        row["server"], row["domain"], row["status"], row["start"],
                        row["end"], row["duration"], row["command"]),
                        tags=(self._monitor_status_tag(status),))
                elif visible:
                    self.dig_monitor_tree.item(row["id"], values=(
                        row["server"], row["domain"], row["status"], row["start"],
                        row["end"], row["duration"], row["command"]),
                        tags=(self._monitor_status_tag(status),))
                elif exists:
                    self.dig_monitor_tree.delete(row["id"])
        except Exception:
            pass




    @staticmethod
    def _monitor_status_tag(status):
        if status == "IN PROGRESS":
            return "running"
        if status in ("NOERROR", "NXDOMAIN"):
            return "success"
        if status in ("PENDING", "NO RESULT", "STOPPED"):
            return "pending"
        return "failed"




    def _monitor_open_failed_view(self):
        try:
            if self.dig_monitor_failed_window is not None and self.dig_monitor_failed_window.winfo_exists():
                self.dig_monitor_failed_window.deiconify()
                self.dig_monitor_failed_window.lift()
                self._monitor_refresh_failed_view()
                return
            win = ctk.CTkToplevel(self.dig_monitor_window or self.root)
            self.dig_monitor_failed_window = win
            win.title("Failed DIG Queries — Troubleshooting")
            win.geometry("1050x480")
            win.minsize(760, 350)
            win.configure(fg_color=COLOR_BG)
            win.grid_columnconfigure(0, weight=1)
            win.grid_rowconfigure(0, weight=1)
            columns = ("server", "domain", "status", "command")
            tree = ttk.Treeview(win, columns=columns, show="headings", selectmode="browse")
            self.dig_monitor_failed_tree = tree
            for col, label, width in (("server", "Server", 160), ("domain", "Domain / Query", 220),
                                      ("status", "Status", 160), ("command", "DIG Command", 470)):
                tree.heading(col, text=label)
                tree.column(col, width=width, minwidth=90, anchor="w")
            tree.grid(row=0, column=0, sticky="nsew", padx=(12, 0), pady=12)
            scroll = ttk.Scrollbar(win, orient="vertical", command=tree.yview)
            tree.configure(yscrollcommand=scroll.set)
            scroll.grid(row=0, column=1, sticky="ns", padx=(0, 12), pady=12)
            actions = ctk.CTkFrame(win, fg_color="transparent")
            actions.grid(row=1, column=0, columnspan=2, sticky="ew", padx=12, pady=(0, 12))
            ctk.CTkButton(actions, text="Copy Selected Command", command=self._monitor_copy_failed_command,
                          fg_color=COLOR_ACCENT, hover_color=COLOR_ACCENT_HOVER).pack(side="left")
            ctk.CTkButton(actions, text="Refresh", command=self._monitor_refresh_failed_view,
                          fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER).pack(side="left", padx=8)
            ctk.CTkButton(actions, text="Close", command=win.withdraw,
                          fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER).pack(side="right")
            win.protocol("WM_DELETE_WINDOW", win.withdraw)
            self._monitor_refresh_failed_view()
        except Exception as exc:
            self.log(f"[MONITOR] Unable to open failed-query view: {exc}")




    def _monitor_refresh_failed_view(self):
        try:
            tree = self.dig_monitor_failed_tree
            if tree is None or not tree.winfo_exists():
                return
            tree.delete(*tree.get_children())
            with self.dig_monitor_lock:
                rows = [row.copy() for row in self.dig_monitor_rows]
            for row in rows:
                status = row["status"].upper()
                if status in ("NOERROR", "IN PROGRESS", "PENDING", "NO RESULT", "STOPPED"):
                    continue
                tree.insert("", "end", values=(row["server"], row["domain"], status, row["command"]))
        except Exception:
            pass




    def _monitor_copy_failed_command(self):
        try:
            tree = self.dig_monitor_failed_tree
            selected = tree.selection() if tree is not None else ()
            if not selected:
                messagebox.showinfo("Copy command", "Select a failed query first.",
                                    parent=self.dig_monitor_failed_window)
                return
            values = tree.item(selected[0], "values")
            command = str(values[3]) if len(values) > 3 else ""
            if command:
                self.dig_monitor_failed_window.clipboard_clear()
                self.dig_monitor_failed_window.clipboard_append(command)
                self.dig_monitor_failed_window.update_idletasks()
                messagebox.showinfo("Copied", "DIG command copied to clipboard.",
                                    parent=self.dig_monitor_failed_window)
        except Exception as exc:
            messagebox.showerror("Copy failed", str(exc), parent=self.dig_monitor_failed_window)




    @staticmethod
    def _monitor_format_duration(seconds):
        seconds = max(0, int(seconds or 0))
        hours, remainder = divmod(seconds, 3600)
        minutes, seconds = divmod(remainder, 60)
        return f"{hours:02d}:{minutes:02d}:{seconds:02d}"




    def _monitor_command_event(self, event, server, command, status=None, started=None, ended=None):
        """Thread-safe callback from SSH execution; updates monitor only."""
        target = normalize_command(command)
        row_to_update = None
        with self.dig_monitor_lock:
            matches = [
                row for row in self.dig_monitor_rows
                if row["server"].lower() == str(server or "").lower()
                and command_key(row["command"]) == command_key(target)
            ]
            if event == "start":
                row_to_update = next((row for row in matches if row["status"] == "PENDING"), None)
                if row_to_update is None:
                    row_to_update = next((row for row in matches if row["status"] != "IN PROGRESS"), None)
                if row_to_update is not None:
                    row_to_update["status"] = "IN PROGRESS"
                    row_to_update["started_ts"] = started or time.time()
                    row_to_update["start"] = datetime.datetime.fromtimestamp(
                        row_to_update["started_ts"]).strftime("%H:%M:%S")
                    row_to_update["end"] = ""
                    row_to_update["ended_ts"] = None
                    row_to_update["duration"] = "Running..."
            elif event == "complete":
                row_to_update = next((row for row in matches if row["status"] == "IN PROGRESS"), None)
                if row_to_update is None:
                    row_to_update = next((row for row in matches if row["status"] == "PENDING"), None)
                if row_to_update is not None:
                    row_to_update["status"] = str(status or "NO RESULT").upper()
                    row_to_update["ended_ts"] = ended or time.time()
                    row_to_update["end"] = datetime.datetime.fromtimestamp(
                        row_to_update["ended_ts"]).strftime("%H:%M:%S")
                    if row_to_update["started_ts"] is None:
                        row_to_update["started_ts"] = started or row_to_update["ended_ts"]
                        row_to_update["start"] = datetime.datetime.fromtimestamp(
                            row_to_update["started_ts"]).strftime("%H:%M:%S")
                    row_to_update["duration"] = self._monitor_format_duration(
                        row_to_update["ended_ts"] - row_to_update["started_ts"])
        if row_to_update is not None:
            self.root.after(0, lambda rid=row_to_update["id"]: self._monitor_refresh_row(rid))
            self.root.after(0, self._monitor_refresh)




    def _monitor_refresh_row(self, row_id):
        try:
            if self.dig_monitor_tree is None or not self.dig_monitor_tree.winfo_exists():
                return
            with self.dig_monitor_lock:
                row = next((item.copy() for item in self.dig_monitor_rows if item["id"] == row_id), None)
            if not row:
                return
            status = row["status"]
            tag = self._monitor_status_tag(status)
            values = (row["server"], row["domain"], status, row["start"], row["end"],
                      row["duration"], row["command"])
            if self.dig_monitor_tree.exists(row_id):
                self.dig_monitor_tree.item(row_id, values=values, tags=(tag,))
        except Exception:
            pass




    def _monitor_refresh(self):
        try:
            with self.dig_monitor_lock:
                rows = [row.copy() for row in self.dig_monitor_rows]
                started_at = self.dig_monitor_started_at
                finished_at = self.dig_monitor_finished_at
            total = len(rows)
            running = sum(1 for row in rows if row["status"] == "IN PROGRESS")
            pending = sum(1 for row in rows if row["status"] == "PENDING")
            completed = sum(1 for row in rows if row["status"] not in ("IN PROGRESS", "PENDING"))
            if self.dig_monitor_summary is not None and self.dig_monitor_summary.winfo_exists():
                self.dig_monitor_summary.configure(
                    text=f"Total DIGs: {total}   •   Live running: {running}   •   Completed: {completed}   •   Pending: {pending}"
                )
            if self.dig_monitor_progress is not None and self.dig_monitor_progress.winfo_exists():
                self.dig_monitor_progress.set((completed / total) if total else 0)
            if self.dig_monitor_started_label is not None and self.dig_monitor_started_label.winfo_exists():
                start_text = datetime.datetime.fromtimestamp(started_at).strftime("%Y-%m-%d %H:%M:%S") if started_at else "—"
                self.dig_monitor_started_label.configure(text=f"Start: {start_text}")
            end_at = finished_at or time.time()
            elapsed = end_at - started_at if started_at else 0
            if self.dig_monitor_elapsed_label is not None and self.dig_monitor_elapsed_label.winfo_exists():
                suffix = "  •  Finished" if finished_at else ""
                self.dig_monitor_elapsed_label.configure(
                    text=f"Total time: {self._monitor_format_duration(elapsed)}{suffix}")
            # ETA is based on observed completion throughput; retries/SSH delays can change it.
            done_rows = [row for row in rows if row["status"] not in ("IN PROGRESS", "PENDING", "NO RESULT")]
            remaining = sum(1 for row in rows if row["status"] in ("IN PROGRESS", "PENDING"))
            if self.dig_monitor_eta_label is not None and self.dig_monitor_eta_label.winfo_exists():
                if finished_at or remaining == 0:
                    eta_text = "Estimated remaining: 00:00:00"
                elif len(done_rows) > 0 and elapsed > 0:
                    seconds_left = (elapsed / len(done_rows)) * remaining
                    eta_text = f"Estimated remaining: ~{self._monitor_format_duration(seconds_left)} (estimate)"
                else:
                    eta_text = "Estimated remaining: calculating…"
                self.dig_monitor_eta_label.configure(text=eta_text)
            self._monitor_apply_filters()
            self._monitor_refresh_failed_view()
        except Exception:
            pass




    def _monitor_schedule_timer(self):
        self.dig_monitor_timer_running = True
        def tick():
            if not self.dig_monitor_timer_running:
                return
            self._monitor_refresh()
            with self.dig_monitor_lock:
                active = self.dig_monitor_finished_at is None
            if active:
                try:
                    self.root.after(500, tick)
                except Exception:
                    pass
        try:
            self.root.after(500, tick)
        except Exception:
            pass




    def _monitor_finish(self):
        finished = time.time()
        with self.dig_monitor_lock:
            self.dig_monitor_finished_at = finished
            final_status = "STOPPED" if STOP_DIG_EVENT.is_set() else "NO RESULT"
            for row in self.dig_monitor_rows:
                if row["status"] in ("PENDING", "IN PROGRESS"):
                    row["status"] = final_status
                    row["ended_ts"] = finished
                    row["end"] = datetime.datetime.fromtimestamp(finished).strftime("%H:%M:%S")
                    if row["started_ts"] is not None:
                        row["duration"] = self._monitor_format_duration(finished - row["started_ts"])
                    else:
                        row["duration"] = "—"
        self.dig_monitor_timer_running = False
        self.root.after(0, self._monitor_refresh)




    def _monitor_export_csv(self):
        try:
            from tkinter import filedialog
            path = filedialog.asksaveasfilename(
                parent=self.dig_monitor_window, title="Export live DIG results",
                defaultextension=".csv", filetypes=[("CSV files", "*.csv")],
                initialfile="infoblox_live_dig_results.csv")
            if not path:
                return
            with self.dig_monitor_lock:
                rows = [row.copy() for row in self.dig_monitor_rows]
            with open(path, "w", newline="", encoding="utf-8-sig") as file:
                writer = csv.writer(file)
                writer.writerow(["Server", "Domain / Query", "Status", "Start Time", "End Time", "Duration", "DIG Command"])
                for row in rows:
                    writer.writerow([row["server"], row["domain"], row["status"], row["start"],
                                     row["end"], row["duration"], row["command"]])
            messagebox.showinfo("Export complete", f"Live DIG results exported to:\n{path}",
                                parent=self.dig_monitor_window)
        except Exception as exc:
            messagebox.showerror("Export failed", str(exc), parent=self.root)




    # ========================================================
    # LIVE OUTPUT WINDOW
    # ========================================================
    def open_live_output(self):
        try:
            if self.live_output_window is not None and self.live_output_window.winfo_exists():
                self.live_output_window.deiconify()
                self.live_output_window.lift()
                self.live_output_window.focus_force()
                return




            self.live_output_window = ctk.CTkToplevel(self.root)
            self.live_output_window.title("INFOBLOX DIG AUTOMATION — Live Output")
            self.live_output_window.geometry("1100x650")
            self.live_output_window.minsize(700, 400)
            self.live_output_window.configure(fg_color=COLOR_BG)
            self.live_output_window.grid_rowconfigure(1, weight=1)
            self.live_output_window.grid_columnconfigure(0, weight=1)




            header = ctk.CTkFrame(self.live_output_window, fg_color=COLOR_SURFACE, corner_radius=12)
            header.grid(row=0, column=0, sticky="ew", padx=14, pady=(14, 7))
            header.grid_columnconfigure(0, weight=1)




            ctk.CTkLabel(
                header, text="⬒   LIVE OUTPUT",
                font=font(14, "bold"), text_color=COLOR_TEXT, anchor="w"
            ).grid(row=0, column=0, padx=16, pady=12, sticky="w")




            ctk.CTkButton(
                header, text="Clear", width=90, height=34, corner_radius=10,
                fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
                text_color=COLOR_TEXT, font=font(11, "bold"),
                command=self.clear_live_output
            ).grid(row=0, column=1, padx=14, pady=10)




            self.console = ctk.CTkTextbox(
                self.live_output_window,
                font=mono_font(12),
                corner_radius=12,
                fg_color="#0a0e14",
                border_width=1,
                border_color=COLOR_BORDER,
                text_color="#c9d1d9",
                wrap="none"
            )
            self.console.grid(row=1, column=0, sticky="nsew", padx=14, pady=(0, 14))




            if self.live_output_buffer:
                self.console.insert("end", self.live_output_buffer)
                self.console.see("end")




            self.live_output_window.protocol("WM_DELETE_WINDOW", self.close_live_output)
        except Exception as exc:
            # Keep the failure visible instead of silently suppressing it.
            error_details = traceback.format_exc()
            try:
                if self.live_output_window is not None and self.live_output_window.winfo_exists():
                    self.live_output_window.destroy()
            except Exception:
                pass
            self.live_output_window = None
            self.console = None
            self.log(f"[LIVE OUTPUT] CustomTkinter window creation failed: {exc}\n{error_details}")
            # Fallback to standard Tk widgets so FQDN-triggered DIG still has
            # a visible live output window if the themed widget fails.
            try:
                self.live_output_window = tk.Toplevel(self.root)
                self.live_output_window.title("INFOBLOX DIG AUTOMATION — Live Output")
                self.live_output_window.geometry("1100x650")
                self.live_output_window.minsize(700, 400)
                self.live_output_window.configure(bg=COLOR_BG)
                header = tk.Frame(self.live_output_window, bg=COLOR_SURFACE)
                header.pack(fill="x", padx=12, pady=(12, 6))
                tk.Label(header, text="LIVE DIG OUTPUT", bg=COLOR_SURFACE,
                         fg=COLOR_TEXT, font=(FONT_FAMILY, 12, "bold")).pack(side="left", padx=12, pady=10)
                tk.Button(header, text="Clear", command=self.clear_live_output).pack(side="right", padx=12, pady=8)
                self.console = tk.Text(self.live_output_window, bg="#0a0e14", fg="#c9d1d9",
                                       insertbackground=COLOR_TEXT, font=(FONT_MONO, 10), wrap="none")
                self.console.pack(fill="both", expand=True, padx=12, pady=(0, 12))
                if self.live_output_buffer:
                    self.console.insert("end", self.live_output_buffer)
                    self.console.see("end")
                self.live_output_window.protocol("WM_DELETE_WINDOW", self.close_live_output)
                self.live_output_window.deiconify()
                self.live_output_window.lift()
                self.live_output_window.focus_force()
            except Exception as fallback_exc:
                self.log(f"[LIVE OUTPUT] Tk fallback window also failed: {fallback_exc}\n{traceback.format_exc()}")
                self.live_output_window = None
                self.console = None




    def open_fqdn_live_output(self):
        """Open/raise the live output window for DIG launched from FQDN Search."""
        def show_window():
            try:
                self.open_live_output()
                window = self.live_output_window
                if window is not None and window.winfo_exists():
                    window.deiconify()
                    window.lift()
                    try:
                        window.attributes("-topmost", True)
                        window.after(700, lambda: window.attributes("-topmost", False)
                                     if window.winfo_exists() else None)
                    except Exception:
                        pass
                    window.focus_force()
                else:
                    self.log("[FQDN DIG] Live output window could not be opened.")
            except Exception as exc:
                self.log(f"[FQDN DIG] Unable to open live output window: {exc}")




        try:
            self.root.after(0, show_window)
        except Exception as exc:
            self.log(f"[FQDN DIG] Unable to schedule live output window: {exc}")




    def close_live_output(self):
        try:
            if self.live_output_window is not None and self.live_output_window.winfo_exists():
                self.live_output_window.destroy()
        except Exception:
            pass
        finally:
            self.live_output_window = None
            self.console = None




    def clear_live_output(self):
        self.live_output_buffer = ""
        if self.console is not None:
            try:
                self.console.delete("1.0", "end")
            except Exception:
                pass




    # ========================================================
    # LOGGING
    # ========================================================
    def log(self, message, raw=False):
        self.log_queue.put(("raw" if raw else "text", message))




    def process_log_queue(self):
        processed = 0
        max_messages = 100
        try:
            while processed < max_messages:
                message_type, message = self.log_queue.get_nowait()
                if message_type == "raw":
                    display_message = message
                else:
                    display_message = "\n" + message + "\n"
                self.live_output_buffer += display_message
                if self.console is not None:
                    try:
                        self.console.insert("end", display_message)
                    except Exception:
                        self.console = None
                processed += 1
            if processed and self.console is not None:
                try:
                    self.console.see("end")
                except Exception:
                    self.console = None
        except queue.Empty:
            pass
        self.root.after(50, self.process_log_queue)




    # ========================================================
    # MANUAL INFOBLOX PAGER CONTROL
    # ========================================================
    def send_pager_q(self):
        global ACTIVE_SSH_CHANNEL
        with ACTIVE_SSH_CHANNEL_LOCK:
            channel = ACTIVE_SSH_CHANNEL
        if channel is None:
            self.log("No active SSH session. Start RUN ALL STEPS first.")
            return
        try:
            channel.send("q\n")
            self.log("[MANUAL] Sent q + Enter to Infoblox pager. Continuing...")
        except Exception as exc:
            self.log(f"[MANUAL] Failed to send q: {exc}")




    # ========================================================
    # START / STOP
    # ========================================================
    def stop_dig(self):
        STOP_DIG_EVENT.set()
        with ACTIVE_SSH_CHANNEL_LOCK:
            channels = list(ACTIVE_SSH_CHANNELS)
        for channel in channels:
            try:
                channel.close()
            except Exception:
                pass
        self.log("[STOP] Stop DIG requested. Active and pending DIG commands are being stopped.")
        try:
            self.stop_dig_button.configure(state="disabled")
        except Exception:
            pass




    def start_run(self):
        STOP_DIG_EVENT.clear()
        self.stop_dig_button.configure(state="normal")




        global OUTPUT_DIR
        selected_output = self.output_entry.get().strip()
        if selected_output:
            OUTPUT_DIR = os.path.abspath(selected_output)
        os.makedirs(OUTPUT_DIR, exist_ok=True)




        self.open_live_output()
        self.run_button.configure(state="disabled")
        self.clear_live_output()
        self.last_txt = None
        self.last_excel = None
        self.step1_pill.set_status("Running", COLOR_ACCENT)
        self.step2_pill.set_status("Waiting", COLOR_TEXT_MUTED)
        self.step3_pill.set_status("Waiting", COLOR_TEXT_MUTED)
        self.excel_pill.set_status("Waiting", COLOR_TEXT_MUTED)




        thread = threading.Thread(target=self.run_all, daemon=True)
        thread.start()




    def update_step_progress(self, step, completed, total):
        text_value = f"DIGs: {completed} / {total}"
        if step == 1:
            self.step1_progress.configure(text=text_value)
        elif step == 3:
            self.step3_progress.configure(text=text_value)




    # ========================================================
    # RUN ALL
    # ========================================================
    def run_all(self):
        session = None
        try:
            server = self.server_entry.get().strip()
            username = self.username_entry.get().strip()
            port = self.port_entry.get().strip()
            authentication = self.auth_var.get()
            password = self.password_entry.get()
            key_file = self.key_entry.get().strip()
            key_passphrase = self.key_passphrase_entry.get()




            if not server:
                raise Exception("Server / IP is required.")
            if not username:
                raise Exception("Username is required.")
            if authentication == "Password":
                if not password:
                    raise Exception("SSH password is required.")
            else:
                if not key_file:
                    raise Exception("Please select a .ppk or .pem key.")
                if not os.path.exists(key_file):
                    raise Exception("Selected key does not exist.")




            commands = self.get_effective_commands()
            if not commands:
                raise Exception("No DIG commands found in the selected input.")




            self._monitor_begin([(server, command) for command in commands])




            if getattr(self, "fast_mode_var", None) is not None and self.fast_mode_var.get():
                try:
                    worker_count = int(self.fast_workers_entry.get().strip())
                except (ValueError, AttributeError):
                    raise Exception("Fast DIG worker count must be a number between 1 and 10.")
                if worker_count < 1 or worker_count > 10:
                    raise Exception("Fast DIG worker count must be between 1 and 10.")
                self.run_fast_mode(
                    server=server, username=username, port=port,
                    authentication=authentication, password=password,
                    key_file=key_file, key_passphrase=key_passphrase,
                    commands=commands, worker_count=worker_count
                )
                return




            self.log("=" * 70)
            self.log("INFOBLOX DIG AUTOMATION")
            self.log("=" * 70)
            self.log(f"Server   : {server}")
            self.log(f"Username : {username}")
            self.log(f"Port     : {port}")
            self.log(f"Auth     : {authentication}")
            self.log(f"Commands : {len(commands)}")
            if authentication == "SSH Key":
                extension = os.path.splitext(key_file)[1].upper()
                self.log(f"Key Type : {extension}")




            self.root.after(0, lambda total=len(commands): self.update_step_progress(1, 0, total))
            self.root.after(0, lambda: self.update_step_progress(3, 0, 0))
            self.log("\n[STEP 1] Running Infoblox DIG...")




            step1_result = execute_ssh(
                server=server, username=username, port=port,
                authentication=authentication, password=password,
                key_file=key_file, key_passphrase=key_passphrase,
                commands=commands, log_callback=self.log,
                keep_connection=True,
                progress_callback=lambda completed, total: self.root.after(
                    0, lambda c=completed, t=total: self.update_step_progress(1, c, t)
                ),
                command_callback=self._monitor_command_event
            )




            step1_output, session_client, session_channel, session_key_dir = step1_result
            session = (session_client, session_channel, session_key_dir)
            self.root.after(0, lambda: self.step1_pill.set_status("Completed", COLOR_SUCCESS))




            self.log("\n[STEP 2] Extracting unsuccessful commands...")
            unsuccessful = extract_unsuccessful(step1_output)
            self.log(f"Unsuccessful commands: {len(unsuccessful)}")
            for command in unsuccessful:
                self.log("  " + command)
            self.root.after(0, lambda: self.step2_pill.set_status("Completed", COLOR_SUCCESS))




            self.log("\n[STEP 3] Retrying unsuccessful commands...")
            self.root.after(0, lambda total=len(unsuccessful): self.update_step_progress(3, 0, total))




            if unsuccessful:
                step3_output = run_step3(
                    server=server, username=username, port=port,
                    authentication=authentication, password=password,
                    key_file=key_file, key_passphrase=key_passphrase,
                    unsuccessful_commands=unsuccessful,
                    log_callback=self.log, session=session,
                    progress_callback=lambda completed, total: self.root.after(
                        0, lambda c=completed, t=total: self.update_step_progress(3, c, t)
                    ),
                    command_callback=self._monitor_command_event
                )
            else:
                step3_output = ""
                self.log("No commands require retry.")
            self.root.after(0, lambda: self.step3_pill.set_status("Completed", COLOR_SUCCESS))




            self.log("\n[OUTPUT] Creating final TXT...")
            self.last_txt = create_final_file(
                server=server, username=username, authentication=authentication,
                commands=commands, step1_output=step1_output,
                unsuccessful=unsuccessful, step3_output=step3_output
            )
            self.log(f"TXT saved: {self.last_txt}")




            self.log("\n[EXCEL] Creating report...")
            self.last_excel = create_excel_result(
                commands=commands, step1_output=step1_output,
                step3_output=step3_output, unsuccessful=unsuccessful, server=server
            )
            self.log(f"Excel saved: {self.last_excel}")
            self.root.after(0, lambda: self.excel_pill.set_status("Created", COLOR_SUCCESS))




            # Additive: only remember a new server after DIG has returned a DNS response.
            if not STOP_DIG_EVENT.is_set():
                self._auto_save_profile_after_successful_dig(
                    server, (step1_output or "") + "\\n" + (step3_output or "")
                )

            self.log("\\n" + "=" * 70)
            self.log("COMPLETED SUCCESSFULLY")
            self.log("=" * 70)




            self.root.after(
                0,
                lambda: messagebox.showinfo(
                    "Completed",
                    "Infoblox DIG automation completed.\n\n"
                    f"TXT:\n{self.last_txt}\n\n"
                    f"Excel:\n{self.last_excel}"
                )
            )
        except Exception as error:
            self.log("\n" + "=" * 70)
            self.log("ERROR")
            self.log("=" * 70)
            self.log(str(error))
            self.root.after(
                0,
                lambda: messagebox.showerror("Infoblox DIG Error", str(error))
            )
        finally:
            if session:
                try:
                    session[1].send("exit\n")
                except Exception:
                    pass
                try:
                    session[1].close()
                except Exception:
                    pass
                try:
                    session[0].close()
                except Exception:
                    pass
                if session[2]:
                    shutil.rmtree(session[2], ignore_errors=True)




            global ACTIVE_SSH_CHANNEL
            with ACTIVE_SSH_CHANNEL_LOCK:
                ACTIVE_SSH_CHANNEL = None




            self.root.after(0, self._monitor_finish)
            self.root.after(0, lambda: self.run_button.configure(state="normal"))
            self.root.after(0, lambda: self.stop_dig_button.configure(state="disabled"))




    # ========================================================
    # FAST DIG MODE
    # ========================================================
    def run_fast_mode(
        self, server, username, port, authentication,
        password, key_file, key_passphrase, commands, worker_count=3
    ):
        if not commands:
            raise Exception("No DIG commands are available for Fast DIG Mode.")




        actual_workers = min(int(worker_count), len(commands))
        chunks = []
        for index in range(actual_workers):
            start = index * len(commands) // actual_workers
            end = (index + 1) * len(commands) // actual_workers
            chunks.append((index, commands[start:end]))




        self.log("=" * 70)
        self.log("INFOBLOX DIG AUTOMATION - FAST DIG MODE")
        self.log("=" * 70)
        self.log(f"Server   : {server}")
        self.log(f"Username : {username}")
        self.log(f"Port     : {port}")
        self.log(f"Auth     : {authentication}")
        self.log(f"Total DIG commands: {len(commands)}")
        self.log(f"Workers requested: {worker_count}; workers started: {actual_workers}")
        for index, chunk in chunks:
            self.log(f"[FAST] Worker {index + 1}: {len(chunk)} commands")




        self.root.after(0, lambda total=len(commands): self.update_step_progress(1, 0, total))
        self.root.after(0, lambda: self.update_step_progress(3, 0, 0))
        self.root.after(0, lambda: self.step1_pill.set_status("Running (Fast)", COLOR_ACCENT))
        self.root.after(0, lambda: self.step2_pill.set_status("Waiting", COLOR_TEXT_MUTED))
        self.root.after(0, lambda: self.step3_pill.set_status("Waiting", COLOR_TEXT_MUTED))
        self.root.after(0, lambda: self.excel_pill.set_status("Waiting", COLOR_TEXT_MUTED))




        progress_lock = threading.Lock()
        worker_progress = {index: 0 for index, _ in chunks}
        worker_retry_progress = {index: 0 for index, _ in chunks}
        retry_totals = {index: 0 for index, _ in chunks}




        def progress_callback_for(index, phase, total_for_worker):
            def callback(completed, total):
                with progress_lock:
                    if phase == 1:
                        worker_progress[index] = min(max(0, completed), total_for_worker)
                        combined_done = sum(worker_progress.values())
                        combined_total = len(commands)
                    else:
                        worker_retry_progress[index] = min(max(0, completed), total_for_worker)
                        combined_done = sum(worker_retry_progress.values())
                        combined_total = sum(retry_totals.values())
                step_number = 1 if phase == 1 else 3
                try:
                    self.root.after(
                        0,
                        lambda c=combined_done, t=combined_total, st=step_number:
                            self.update_step_progress(st, c, t)
                    )
                except Exception:
                    pass
            return callback




        def timeout_output(chunk):
            return "".join(
                f"\n> {normalize_command(command)}\nconnection timed out\n"
                for command in chunk
            )




        def worker_job(index, chunk):
            worker_label = f"[FAST W{index + 1}]"
            self.log(f"{worker_label} Connecting; assigned {len(chunk)} commands.")
            session = None
            step1_output = ""
            unsuccessful = []
            step3_output = ""
            try:
                if index > 0:
                    stagger_seconds = min(index * 2, 8)
                    self.log(f"{worker_label} Waiting {stagger_seconds}s before SSH connection to reduce handshake bursts.")
                    for _ in range(stagger_seconds * 10):
                        if STOP_DIG_EVENT.is_set():
                            break
                        time.sleep(0.1)




                step1_result = ""
                max_connect_attempts = 3
                for connect_attempt in range(1, max_connect_attempts + 1):
                    if STOP_DIG_EVENT.is_set():
                        break
                    step1_result = execute_ssh(
                        server=server, username=username, port=port,
                        authentication=authentication, password=password,
                        key_file=key_file, key_passphrase=key_passphrase,
                        commands=chunk,
                        log_callback=lambda message, raw=False: self.log(
                            f"{worker_label} {message}" if not raw else message, raw=raw
                        ),
                        keep_connection=True,
                        progress_callback=progress_callback_for(index, 1, len(chunk)),
                        command_callback=self._monitor_command_event
                    )
                    if isinstance(step1_result, tuple) and len(step1_result) == 4:
                        break
                    if connect_attempt < max_connect_attempts and not STOP_DIG_EVENT.is_set():
                        wait_seconds = connect_attempt * 5
                        self.log(
                            f"{worker_label} SSH session did not establish (attempt "
                            f"{connect_attempt}/{max_connect_attempts}). Retrying in {wait_seconds}s. "
                            "If this repeats, reduce Fast DIG workers to 1 or 2."
                        )
                        for _ in range(wait_seconds * 10):
                            if STOP_DIG_EVENT.is_set():
                                break
                            time.sleep(0.1)




                if not isinstance(step1_result, tuple) or len(step1_result) != 4:
                    self.log(f"{worker_label} SSH connection failed after {max_connect_attempts} attempts; marking assigned commands as timed out.")
                    for failed_command in chunk:
                        self._monitor_command_event(
                            "complete", server, failed_command, "CONNECTION TIMED OUT",
                            time.time(), time.time()
                        )
                    return index, timeout_output(chunk), [], "", True




                step1_output, client, channel, key_dir = step1_result
                session = (client, channel, key_dir)
                if STOP_DIG_EVENT.is_set():
                    self.log(f"{worker_label} Stop requested; skipping retries.")
                else:
                    unsuccessful = extract_unsuccessful(step1_output)
                    with progress_lock:
                        retry_totals[index] = len(unsuccessful)
                    self.log(f"{worker_label} Step 1 finished; unsuccessful commands: {len(unsuccessful)}")
                    if unsuccessful and not STOP_DIG_EVENT.is_set():
                        step3_output = run_step3(
                            server=server, username=username, port=port,
                            authentication=authentication, password=password,
                            key_file=key_file, key_passphrase=key_passphrase,
                            unsuccessful_commands=unsuccessful,
                            log_callback=lambda message, raw=False: self.log(
                                f"{worker_label} {message}" if not raw else message, raw=raw
                            ),
                            session=session,
                            progress_callback=progress_callback_for(index, 3, len(unsuccessful)),
                            command_callback=self._monitor_command_event
                        )
                return index, step1_output, unsuccessful, step3_output, False
            except Exception as exc:
                self.log(f"{worker_label} ERROR: {exc}")
                return index, timeout_output(chunk), [], "", True
            finally:
                if session:
                    try:
                        if not STOP_DIG_EVENT.is_set():
                            session[1].send("exit\n")
                    except Exception:
                        pass
                    try:
                        session[1].close()
                    except Exception:
                        pass
                    try:
                        session[0].close()
                    except Exception:
                        pass
                    if session[2]:
                        shutil.rmtree(session[2], ignore_errors=True)




        worker_results = {}
        with concurrent.futures.ThreadPoolExecutor(max_workers=actual_workers) as executor:
            future_map = {
                executor.submit(worker_job, index, chunk): index
                for index, chunk in chunks
            }
            for future in concurrent.futures.as_completed(future_map):
                index = future_map[future]
                try:
                    worker_results[index] = future.result()
                except Exception as exc:
                    chunk = dict(chunks)[index]
                    self.log(f"[FAST W{index + 1}] Unhandled worker error: {exc}")
                    worker_results[index] = (index, timeout_output(chunk), [], "", True)




        ordered_results = [worker_results[index] for index, _ in chunks]
        combined_step1 = "\n".join(item[1] for item in ordered_results if item[1])
        combined_unsuccessful = [command for item in ordered_results for command in item[2]]
        combined_step3 = "\n".join(item[3] for item in ordered_results if item[3])
        failed_workers = sum(1 for item in ordered_results if item[4])




        self.root.after(
            0,
            lambda: self.step1_pill.set_status(
                "Stopped" if STOP_DIG_EVENT.is_set() else "Completed",
                COLOR_WARNING if STOP_DIG_EVENT.is_set() else COLOR_SUCCESS
            )
        )
        self.root.after(0, lambda: self.step2_pill.set_status("Completed", COLOR_SUCCESS))
        self.log(f"[FAST] All workers finished. Failed worker sessions: {failed_workers}")
        self.log(f"[FAST] Total unsuccessful commands to retry: {len(combined_unsuccessful)}")
        self.root.after(0, lambda total=len(combined_unsuccessful): self.update_step_progress(3, total, total))
        self.root.after(0, lambda: self.step3_pill.set_status(
            "Stopped" if STOP_DIG_EVENT.is_set() else "Completed",
            COLOR_WARNING if STOP_DIG_EVENT.is_set() else COLOR_SUCCESS
        ))




        self.log("[FAST] Creating one combined TXT report...")
        self.last_txt = create_final_file(
            server=server, username=username, authentication=authentication,
            commands=commands, step1_output=combined_step1,
            unsuccessful=combined_unsuccessful, step3_output=combined_step3
        )
        self.log(f"[FAST] Combined TXT saved: {self.last_txt}")




        self.log("[FAST] Creating one combined Excel report...")
        self.last_excel = create_excel_result(
            commands=commands, step1_output=combined_step1,
            step3_output=combined_step3,
            unsuccessful=combined_unsuccessful, server=server
        )
        self.log(f"[FAST] Combined Excel saved: {self.last_excel}")
        self.root.after(0, lambda: self.excel_pill.set_status("Created", COLOR_SUCCESS))




        if not STOP_DIG_EVENT.is_set():
            # Additive: only remember a new server when at least one DIG got a DNS response.
            self._auto_save_profile_after_successful_dig(
                server, (combined_step1 or "") + "\\n" + (combined_step3 or "")
            )

        if STOP_DIG_EVENT.is_set():
            self.log("[FAST] Stopped by user. Combined report contains results collected before stopping.")
            title = "Fast DIG Stopped"
            message = (
                "Fast DIG execution was stopped.\n\n"
                f"Combined TXT:\n{self.last_txt}\n\n"
                f"Combined Excel:\n{self.last_excel}"
            )
        else:
            self.log("[FAST] FAST DIG MODE COMPLETED")
            title = "Fast DIG Completed"
            message = (
                f"Fast DIG completed using {actual_workers} workers.\n"
                f"Commands processed: {len(commands)}\n"
                f"Failed worker sessions: {failed_workers}\n\n"
                f"Combined TXT:\n{self.last_txt}\n\n"
                f"Combined Excel:\n{self.last_excel}"
            )
        self.root.after(0, lambda t=title, m=message: messagebox.showinfo(t, m, parent=self.root))




    # ========================================================
    # SERVER PROFILES / MULTI-INFOBLOX TEST
    # ========================================================
    def load_server_profiles(self):
        if not os.path.exists(SERVER_PROFILES_FILE):
            return []
        try:
            with open(SERVER_PROFILES_FILE, "r", encoding="utf-8") as file:
                data = json.load(file)
            if not isinstance(data, list):
                return []
            sanitized_profiles = []
            password_found = False
            for profile in data:
                if not isinstance(profile, dict):
                    continue
                profile_copy = dict(profile)
                if "password" in profile_copy:
                    profile_copy.pop("password", None)
                    password_found = True
                sanitized_profiles.append(profile_copy)
            if password_found:
                with open(SERVER_PROFILES_FILE, "w", encoding="utf-8") as file:
                    json.dump(sanitized_profiles, file, indent=4)
            return sanitized_profiles
        except Exception as exc:
            self.log(f"[PROFILES] Unable to read profiles: {exc}")
            return []




    def save_server_profiles(self, profiles):
        sanitized_profiles = []
        for profile in profiles:
            if not isinstance(profile, dict):
                continue
            profile_copy = dict(profile)
            profile_copy.pop("password", None)
            sanitized_profiles.append(profile_copy)
        with open(SERVER_PROFILES_FILE, "w", encoding="utf-8") as file:
            json.dump(sanitized_profiles, file, indent=4)




    # Additive helpers for selecting a saved profile from the main DIG page.
    def refresh_main_server_profiles(self):
        try:
            profiles = self.load_server_profiles()
            self.main_server_profiles = profiles
            values = ["Select Server Profile"]
            for profile in profiles:
                name = str(profile.get("name", "")).strip()
                server = str(profile.get("server", "")).strip()
                label = f"{name}  ({server})" if name and server else (name or server)
                if label and label not in values:
                    values.append(label)
            self.main_server_profile_menu.configure(values=values)
            if self.main_server_profile_var.get() not in values:
                self.main_server_profile_var.set("Select Server Profile")
        except Exception as exc:
            self.log(f"[PROFILES] Unable to refresh main DIG server profiles: {exc}")


    def main_server_profile_changed(self, value):
        if value == "Select Server Profile":
            return
        profile = None
        for item in getattr(self, "main_server_profiles", []):
            name = str(item.get("name", "")).strip()
            server = str(item.get("server", "")).strip()
            label = f"{name}  ({server})" if name and server else (name or server)
            if label == value:
                profile = item
                break
        if not profile:
            return
        self.server_entry.delete(0, "end")
        self.server_entry.insert(0, profile.get("server", ""))
        self.username_entry.delete(0, "end")
        self.username_entry.insert(0, profile.get("username", ""))
        self.port_entry.delete(0, "end")
        self.port_entry.insert(0, profile.get("port", "22"))
        self.auth_var.set(profile.get("authentication", "Password"))
        self.authentication_changed(self.auth_var.get())
        self.password_entry.delete(0, "end")
        self.password_entry.insert(0, profile.get("password", ""))
        self.key_entry.delete(0, "end")
        self.key_entry.insert(0, profile.get("key_file", ""))
        self.key_passphrase_entry.delete(0, "end")
        self.key_passphrase_entry.insert(0, profile.get("key_passphrase", ""))
        self.log(f"[PROFILES] Selected main DIG server profile: {profile.get('name', value)}")


    def _find_main_profile_by_server(self, server, profiles=None):
        server_key = str(server or "").strip().casefold()
        if not server_key:
            return None
        for profile in (profiles if profiles is not None else self.load_server_profiles()):
            if str(profile.get("server", "")).strip().casefold() == server_key:
                return profile
        return None


    def _store_main_profile_for_server(self, name, show_message=True):
        """Add or update a main-page connection profile without changing DIG execution."""
        server = self.server_entry.get().strip()
        username = self.username_entry.get().strip()
        if not server:
            messagebox.showwarning(
                "Save Server Profile", "Please enter a Server / IP first.", parent=self.root
            )
            return False
        if not username:
            messagebox.showwarning(
                "Save Server Profile", "Please enter a Username first.", parent=self.root
            )
            return False

        profiles = self.load_server_profiles()
        existing = self._find_main_profile_by_server(server, profiles)
        profile_name = str(name or server).strip()
        if existing:
            # Keep the established profile name when this server is already known.
            profile_name = str(existing.get("name") or profile_name).strip()
            profiles = [p for p in profiles if p is not existing]
        else:
            # Avoid profile-name collisions with another server.
            if any(str(p.get("name", "")).strip().casefold() == profile_name.casefold()
                   for p in profiles):
                profile_name = f"{profile_name} ({server})"

        profiles.append(self.get_current_profile_data(profile_name))
        try:
            self.save_server_profiles(profiles)
        except Exception as exc:
            messagebox.showerror(
                "Save Server Profile", f"Unable to save server profile:\\n{exc}",
                parent=self.root
            )
            return False

        self.refresh_main_server_profiles()
        # Select the newly saved profile in the main dropdown.
        for profile in self.main_server_profiles:
            if str(profile.get("server", "")).strip().casefold() == server.casefold():
                name_value = str(profile.get("name", "")).strip()
                label = f"{name_value}  ({server})" if name_value and server else (name_value or server)
                self.main_server_profile_var.set(label)
                break
        self.log(f"[PROFILES] {'Updated' if existing else 'Saved'} main DIG server profile: {profile_name}")
        if show_message:
            messagebox.showinfo(
                "Server Profile",
                f"{'Updated' if existing else 'Saved'} server profile: {profile_name}",
                parent=self.root
            )
        return True


    def _auto_save_profile_after_successful_dig(self, server, dig_output):
        """Add a new server profile only after DIG returns a DNS response."""
        cleaned_output = clean_terminal_text(dig_output or "")
        # A DNS header status proves the DIG command received a DNS response.
        # Connection timeouts or SSH failures alone must never create a profile.
        if not re.search(r"\bstatus\s*:\s*(?:NOERROR|NXDOMAIN)\b", cleaned_output, flags=re.IGNORECASE):
            self.log(f"[PROFILES] No successful DNS response from {server}; profile not added.")
            return

        def save_on_ui_thread():
            try:
                if self._find_main_profile_by_server(server):
                    self.log(f"[PROFILES] Server already exists in profiles; no duplicate added: {server}")
                    return
                # Use the server/IP as the default profile name; no dialog or extra button.
                saved = self._store_main_profile_for_server(server, show_message=False)
                if saved:
                    self.log(f"[PROFILES] DIG response received; automatically added server profile: {server}")
            except Exception as exc:
                self.log(f"[PROFILES] Automatic profile save skipped for {server}: {exc}")

        try:
            self.root.after(0, save_on_ui_thread)
        except Exception as exc:
            self.log(f"[PROFILES] Could not schedule automatic profile save for {server}: {exc}")


    def save_current_main_server_profile(self):
        """Button action: save the current connection using a user-provided profile name."""
        server = self.server_entry.get().strip()
        if not server:
            messagebox.showwarning(
                "Save Server Profile", "Please enter a Server / IP first.", parent=self.root
            )
            return
        existing = self._find_main_profile_by_server(server)
        default_name = str(existing.get("name", "")).strip() if existing else server
        dialog = ctk.CTkInputDialog(
            text=f"Enter a profile name for {server}:",
            title="Save Current Server Profile"
        )
        entered_name = dialog.get_input()
        if entered_name is None:
            return
        entered_name = entered_name.strip()
        if not entered_name:
            messagebox.showwarning(
                "Save Server Profile", "Profile name cannot be empty.", parent=self.root
            )
            return
        # For an already-saved server, the existing profile is updated rather than duplicated.
        self._store_main_profile_for_server(entered_name or default_name, show_message=True)


    def auto_save_new_main_server_profile(self):
        """Automatically add a server to the profile list the first time it is used."""
        server = self.server_entry.get().strip()
        if not server or self._find_main_profile_by_server(server):
            return
        # No dialog in the run-start path: use Server/IP as the default profile name.
        self._store_main_profile_for_server(server, show_message=False)
        self.log(f"[PROFILES] New server detected; auto-added profile for {server}")


    def get_current_profile_data(self, name):
        return {
            "name": name,
            "server": self.server_entry.get().strip(),
            "username": self.username_entry.get().strip(),
            "port": self.port_entry.get().strip(),
            "authentication": self.auth_var.get(),
            "password": self.password_entry.get(),
            "key_file": self.key_entry.get().strip(),
            "key_passphrase": self.key_passphrase_entry.get()
        }




    def open_server_profiles(self):
        window = ctk.CTkToplevel(self.root)
        window.title("Server Profiles / Multi-Infoblox Test")
        window.geometry("1000x680")
        window.minsize(880, 580)
        window.transient(self.root)
        window.configure(fg_color=COLOR_BG)




        header = ctk.CTkFrame(window, fg_color=COLOR_SURFACE, corner_radius=16, border_width=1, border_color=COLOR_BORDER)
        header.pack(fill="x", padx=18, pady=(18, 10))
        ctk.CTkLabel(
            header,
            text="⚙   SERVER PROFILES / MULTI-INFOBLOX TEST",
            font=font(16, "bold"),
            text_color=COLOR_TEXT
        ).pack(anchor="w", padx=20, pady=(16, 2))
        ctk.CTkLabel(
            header,
            text="Select two or more profiles to run the same DIG command file against multiple Infoblox servers.",
            font=font(10),
            text_color=COLOR_TEXT_MUTED
        ).pack(anchor="w", padx=20, pady=(0, 16))




        list_frame = ctk.CTkFrame(window, fg_color=COLOR_SURFACE, corner_radius=16, border_width=1, border_color=COLOR_BORDER)
        list_frame.pack(fill="both", expand=True, padx=18, pady=8)




        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure(
            "Modern.Treeview",
            background=COLOR_SURFACE_ALT,
            foreground=COLOR_TEXT,
            fieldbackground=COLOR_SURFACE_ALT,
            rowheight=30,
            borderwidth=0,
            font=(FONT_FAMILY, 10)
        )
        style.configure(
            "Modern.Treeview.Heading",
            background=COLOR_BORDER,
            foreground=COLOR_TEXT,
            relief="flat",
            font=(FONT_FAMILY, 10, "bold")
        )
        style.map(
            "Modern.Treeview",
            background=[("selected", COLOR_ACCENT)],
            foreground=[("selected", "#ffffff")]
        )




        columns = ("select", "profile", "server", "username", "auth")
        tree = ttk.Treeview(list_frame, columns=columns, show="headings", height=12, style="Modern.Treeview")
        tree.heading("select", text="Select")
        tree.heading("profile", text="Profile")
        tree.heading("server", text="Server / IP")
        tree.heading("username", text="Username")
        tree.heading("auth", text="Authentication")
        tree.column("select", width=70, anchor="center")
        tree.column("profile", width=180)
        tree.column("server", width=280)
        tree.column("username", width=120)
        tree.column("auth", width=130)




        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        tree.pack(side="left", fill="both", expand=True, padx=(14, 0), pady=14)
        scrollbar.pack(side="right", fill="y", pady=14, padx=(0, 14))




        profiles = self.load_server_profiles()
        selected_names = set()




        def refresh_tree():
            for item in tree.get_children():
                tree.delete(item)
            for profile in profiles:
                tree.insert(
                    "", "end",
                    iid=profile.get("name", profile.get("server", "")),
                    values=(
                        "☐",
                        profile.get("name", ""),
                        profile.get("server", ""),
                        profile.get("username", ""),
                        profile.get("authentication", "")
                    )
                )




        def toggle_selected(event=None):
            item = tree.identify_row(event.y) if event else None
            if not item:
                return
            if item in selected_names:
                selected_names.remove(item)
            else:
                selected_names.add(item)
            values = list(tree.item(item, "values"))
            values[0] = "☑" if item in selected_names else "☐"
            tree.item(item, values=values)




        tree.bind("<ButtonRelease-1>", toggle_selected)
        refresh_tree()




        button_frame = ctk.CTkFrame(window, fg_color="transparent")
        button_frame.pack(fill="x", padx=18, pady=(6, 18))




        def save_current_profile():
            dialog = ctk.CTkInputDialog(text="Enter a profile name:", title="Save Server Profile")
            name = dialog.get_input()
            if not name:
                return
            name = name.strip()
            if not name:
                return
            nonlocal profiles
            profiles = [p for p in profiles if p.get("name", "") != name]
            profiles.append(self.get_current_profile_data(name))
            self.save_server_profiles(profiles)
            selected_names.add(name)
            refresh_tree()
            for item in tree.get_children():
                if item == name:
                    values = list(tree.item(item, "values"))
                    values[0] = "☑"
                    tree.item(item, values=values)
            self.log(f"[PROFILES] Saved profile: {name}")




        def load_selected_profile():
            if len(selected_names) != 1:
                messagebox.showwarning("Load Profile", "Select exactly one server profile to load.")
                return
            name = next(iter(selected_names))
            profile = next((p for p in profiles if p.get("name") == name), None)
            if not profile:
                return
            self.server_entry.delete(0, "end")
            self.server_entry.insert(0, profile.get("server", ""))
            self.username_entry.delete(0, "end")
            self.username_entry.insert(0, profile.get("username", ""))
            self.port_entry.delete(0, "end")
            self.port_entry.insert(0, profile.get("port", "22"))
            self.auth_var.set(profile.get("authentication", "Password"))
            self.authentication_changed(self.auth_var.get())
            self.password_entry.delete(0, "end")
            self.password_entry.insert(0, profile.get("password", ""))
            self.key_entry.delete(0, "end")
            self.key_entry.insert(0, profile.get("key_file", ""))
            self.key_passphrase_entry.delete(0, "end")
            self.key_passphrase_entry.insert(0, profile.get("key_passphrase", ""))
            self.log(f"[PROFILES] Loaded profile: {name}")




        def delete_selected_profiles():
            nonlocal profiles
            if not selected_names:
                messagebox.showwarning("Delete Profile", "Select at least one profile.")
                return
            profiles = [p for p in profiles if p.get("name", "") not in selected_names]
            self.save_server_profiles(profiles)
            selected_names.clear()
            refresh_tree()
            self.log("[PROFILES] Selected profiles deleted.")




        def select_all_profiles():
            selected_names.clear()
            for item in tree.get_children():
                selected_names.add(item)
                values = list(tree.item(item, "values"))
                values[0] = "☑"
                tree.item(item, values=values)




        def clear_selection():
            selected_names.clear()
            for item in tree.get_children():
                values = list(tree.item(item, "values"))
                values[0] = "☐"
                tree.item(item, values=values)




        btn_style = dict(
            height=38, corner_radius=10,
            fg_color=COLOR_SURFACE_ALT, hover_color=COLOR_SURFACE_HOVER,
            text_color=COLOR_TEXT, font=font(10, "bold")
        )




        ctk.CTkButton(button_frame, text="SAVE CURRENT", command=save_current_profile, **btn_style).pack(side="left", padx=4)
        ctk.CTkButton(button_frame, text="LOAD SELECTED", command=load_selected_profile, **btn_style).pack(side="left", padx=4)
        ctk.CTkButton(button_frame, text="DELETE SELECTED", command=delete_selected_profiles, **btn_style).pack(side="left", padx=4)
        ctk.CTkButton(button_frame, text="SELECT ALL", command=select_all_profiles, **btn_style).pack(side="left", padx=4)
        ctk.CTkButton(button_frame, text="CLEAR", command=clear_selection, **btn_style).pack(side="left", padx=4)




        def run_selected():
            selected_profiles = [p for p in profiles if p.get("name", "") in selected_names]
            if not selected_profiles:
                messagebox.showwarning("Multi-Infoblox Test", "Select at least one server profile.")
                return




            password_profiles = [
                p for p in selected_profiles
                if p.get("authentication", "Password") == "Password"
            ]
            if password_profiles:
                common_password = simpledialog.askstring(
                    "Multi-Infoblox Test",
                    "Enter the common SSH password for the selected Password-authentication servers:",
                    parent=window, show="*"
                )
                if common_password is None:
                    return
                if not common_password:
                    messagebox.showwarning(
                        "Multi-Infoblox Test",
                        "SSH password is required for the selected Password-authentication servers."
                    )
                    return
                for profile in password_profiles:
                    profile["password"] = common_password




            try:
                commands = self.get_effective_commands()
            except Exception as exc:
                messagebox.showerror("Multi-Infoblox Test", str(exc))
                return
            if not commands:
                messagebox.showwarning("Multi-Infoblox Test", "No DIG commands found.")
                return
            self.open_live_output()
            window.destroy()
            self.run_multi_server(selected_profiles, commands)




        ctk.CTkButton(
            button_frame,
            text="▶   RUN SELECTED SERVERS",
            height=42, corner_radius=10,
            fg_color=COLOR_SUCCESS, hover_color=COLOR_SUCCESS_HOVER,
            text_color="#ffffff", font=font(11, "bold"),
            command=run_selected
        ).pack(side="right", padx=4)




    def run_multi_server(self, profiles, commands):
        STOP_DIG_EVENT.clear()
        self.stop_dig_button.configure(state="normal")




        global OUTPUT_DIR
        selected_output = self.output_entry.get().strip()
        if selected_output:
            OUTPUT_DIR = os.path.abspath(selected_output)
        os.makedirs(OUTPUT_DIR, exist_ok=True)




        self._monitor_begin(
            [(profile.get("server", "") or profile.get("name", ""), command)
             for profile in profiles for command in commands],
            title="Multi-Infoblox Live DIG Monitor"
        )




        self.run_button.configure(state="disabled")
        self.profiles_button.configure(state="disabled")




        summary_window = ctk.CTkToplevel(self.root)
        summary_window.title("Multi-Infoblox Test Results")
        summary_window.geometry("980x580")
        summary_window.minsize(840, 500)
        summary_window.transient(self.root)
        summary_window.configure(fg_color=COLOR_BG)




        header = ctk.CTkFrame(summary_window, fg_color=COLOR_SURFACE, corner_radius=16, border_width=1, border_color=COLOR_BORDER)
        header.pack(fill="x", padx=18, pady=(18, 10))




        ctk.CTkLabel(
            header,
            text="MULTI-INFOBLOX TEST RESULTS",
            font=font(16, "bold"),
            text_color=COLOR_TEXT
        ).pack(anchor="w", padx=20, pady=(16, 2))




        ctk.CTkLabel(
            header,
            text=f"Commands: {len(commands)}    Servers: {len(profiles)}",
            font=font(10),
            text_color=COLOR_TEXT_MUTED
        ).pack(anchor="w", padx=20, pady=(0, 16))




        table_frame = ctk.CTkFrame(summary_window, fg_color=COLOR_SURFACE, corner_radius=16, border_width=1, border_color=COLOR_BORDER)
        table_frame.pack(fill="both", expand=True, padx=18, pady=8)




        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure(
            "Multi.Treeview",
            background=COLOR_SURFACE_ALT,
            foreground=COLOR_TEXT,
            fieldbackground=COLOR_SURFACE_ALT,
            rowheight=30,
            borderwidth=0,
            font=(FONT_FAMILY, 10)
        )
        style.configure(
            "Multi.Treeview.Heading",
            background=COLOR_BORDER,
            foreground=COLOR_TEXT,
            relief="flat",
            font=(FONT_FAMILY, 10, "bold")
        )




        columns = ("server", "success", "failed", "timeout", "total")
        result_tree = ttk.Treeview(table_frame, columns=columns, show="headings", style="Multi.Treeview")




        headings = {
            "server": "Server", "success": "Success", "failed": "Failed",
            "timeout": "Timeout", "total": "Total"
        }
        widths = {"server": 420, "success": 110, "failed": 110, "timeout": 110, "total": 110}




        for col in columns:
            result_tree.heading(col, text=headings[col])
            result_tree.column(col, width=widths[col], anchor="center" if col != "server" else "w")




        result_tree.pack(fill="both", expand=True, padx=14, pady=14)




        status_label = ctk.CTkLabel(
            summary_window,
            text="Running in parallel...",
            font=font(11, "bold"),
            text_color=COLOR_ACCENT
        )
        status_label.pack(pady=(0, 16))




        dashboard_window = ctk.CTkToplevel(self.root)
        dashboard_window.title("Multi-Infoblox Test — Live Dashboard")
        dashboard_window.geometry("1220x780")
        dashboard_window.minsize(920, 540)
        dashboard_window.configure(fg_color=COLOR_BG)
        dashboard_window.grid_rowconfigure(1, weight=1)
        dashboard_window.grid_columnconfigure(0, weight=1)




        dashboard_header = ctk.CTkFrame(dashboard_window, fg_color=COLOR_SURFACE, corner_radius=16, border_width=1, border_color=COLOR_BORDER)
        dashboard_header.grid(row=0, column=0, sticky="ew", padx=16, pady=(16, 10))




        ctk.CTkLabel(
            dashboard_header,
            text="MULTI-INFOBLOX LIVE EXECUTION",
            font=font(16, "bold"),
            text_color=COLOR_TEXT
        ).pack(anchor="w", padx=20, pady=(16, 2))




        ctk.CTkLabel(
            dashboard_header,
            text=(
                f"Running {len(profiles)} servers simultaneously  •  "
                f"Total DIG commands per server: {len(commands)}"
            ),
            font=font(10),
            text_color=COLOR_TEXT_MUTED
        ).pack(anchor="w", padx=20, pady=(0, 16))




        dashboard_scroll = ctk.CTkScrollableFrame(dashboard_window, fg_color="transparent", corner_radius=0)
        dashboard_scroll.grid(row=1, column=0, sticky="nsew", padx=16, pady=(0, 16))
        dashboard_scroll.grid_columnconfigure(0, weight=1)
        dashboard_scroll.grid_columnconfigure(1, weight=1)




        dashboard_widgets = {}




        for index, profile in enumerate(profiles):
            server_name = profile.get("server", "") or profile.get("name", "") or f"Server {index + 1}"




            card = ctk.CTkFrame(
                dashboard_scroll,
                fg_color=COLOR_SURFACE,
                corner_radius=16,
                border_width=1,
                border_color=COLOR_BORDER
            )
            card.grid(row=index // 2, column=index % 2, sticky="nsew", padx=6, pady=6)
            card.grid_columnconfigure(0, weight=1)
            card.grid_rowconfigure(3, weight=1)




            ctk.CTkLabel(
                card, text=f"🖥  {server_name}",
                font=font(12, "bold"), text_color=COLOR_TEXT, anchor="w"
            ).grid(row=0, column=0, padx=16, pady=(14, 2), sticky="ew")




            server_status = ctk.CTkLabel(
                card, text="Waiting...",
                font=font(10, "bold"), text_color=COLOR_ACCENT, anchor="w"
            )
            server_status.grid(row=1, column=0, padx=16, pady=(0, 4), sticky="ew")




            progress_text = ctk.CTkLabel(
                card, text=f"Commands executed: 0 / {len(commands)}",
                font=font(9), text_color=COLOR_TEXT_MUTED, anchor="w"
            )
            progress_text.grid(row=2, column=0, padx=16, pady=(0, 4), sticky="ew")




            progress_bar = ctk.CTkProgressBar(
                card, height=10, corner_radius=6,
                fg_color=COLOR_SURFACE_ALT, progress_color=COLOR_ACCENT
            )
            progress_bar.set(0)
            progress_bar.grid(row=3, column=0, padx=16, pady=(2, 8), sticky="ew")




            output_box = ctk.CTkTextbox(
                card, height=230, corner_radius=10,
                fg_color="#080b0f", border_width=1, border_color=COLOR_BORDER,
                text_color=COLOR_TEXT, font=mono_font(10), wrap="none"
            )
            output_box.grid(row=4, column=0, padx=16, pady=(0, 16), sticky="nsew")




            dashboard_widgets[index] = {
                "server": server_name,
                "status": server_status,
                "progress_text": progress_text,
                "progress_bar": progress_bar,
                "output": output_box
            }




            output_box.insert("end", f"[{server_name}] Waiting to start...\n")




        def dashboard_alive():
            try:
                return dashboard_window is not None and dashboard_window.winfo_exists()
            except Exception:
                return False




        def append_dashboard_output(server_index, message):
            def update():
                try:
                    if not dashboard_alive():
                        return
                    widget = dashboard_widgets[server_index]["output"]
                    widget.insert("end", message)
                    widget.see("end")
                except Exception:
                    pass
            try:
                self.root.after(0, update)
            except Exception:
                pass




        def update_dashboard_status(server_index, text_value, color=COLOR_ACCENT):
            def update():
                try:
                    if not dashboard_alive():
                        return
                    dashboard_widgets[server_index]["status"].configure(
                        text=text_value, text_color=color
                    )
                except Exception:
                    pass
            try:
                self.root.after(0, update)
            except Exception:
                pass




        def update_dashboard_progress(server_index, completed, total, phase="Step 1"):
            def update():
                try:
                    if not dashboard_alive():
                        return
                    widget_set = dashboard_widgets[server_index]
                    ratio = min(1.0, max(0.0, completed / total)) if total else 0
                    widget_set["progress_bar"].set(ratio)
                    widget_set["progress_text"].configure(
                        text=f"{phase} commands executed: {completed} / {total}"
                    )
                except Exception:
                    pass
            try:
                self.root.after(0, update)
            except Exception:
                pass




        def server_log_callback(server_index, message, raw=False):
            self.log(message, raw=raw)
            append_dashboard_output(server_index, message)




        def add_row(server, success, failed, timeout, total):
            self.root.after(
                0,
                lambda: result_tree.insert(
                    "", "end", values=(server, success, failed, timeout, total)
                )
            )




        comparison_results = {}
        comparison_lock = threading.Lock()
        completed_lock = threading.Lock()
        report_lock = threading.Lock()
        completed_servers = 0




        def store_comparison(server, command, status):
            with comparison_lock:
                comparison_results.setdefault(normalize_command(command), {})[server] = status




        def make_server_reports(
            server, username, authentication, commands_for_report,
            step1_output, unsuccessful, step3_output
        ):
            with report_lock:
                try:
                    txt_report = create_final_file(
                        server=server, username=username, authentication=authentication,
                        commands=commands_for_report, step1_output=step1_output,
                        unsuccessful=unsuccessful, step3_output=step3_output
                    )
                    excel_report = create_excel_result(
                        commands=commands_for_report, step1_output=step1_output,
                        step3_output=step3_output, unsuccessful=unsuccessful, server=server
                    )




                    safe_server = re.sub(r"[^A-Za-z0-9._-]+", "_", server).strip("_") or "server"




                    def unique_server_path(path):
                        base, ext = os.path.splitext(path)
                        candidate = f"{base}_{safe_server}{ext}"
                        counter = 2
                        while os.path.exists(candidate):
                            candidate = f"{base}_{safe_server}_{counter}{ext}"
                            counter += 1
                        return candidate




                    txt_target = unique_server_path(txt_report)
                    excel_target = unique_server_path(excel_report)




                    if os.path.abspath(txt_report) != os.path.abspath(txt_target):
                        os.replace(txt_report, txt_target)
                        txt_report = txt_target
                    if os.path.abspath(excel_report) != os.path.abspath(excel_target):
                        os.replace(excel_report, excel_target)
                        excel_report = excel_target




                    self.log(f"[MULTI TEST] TXT saved: {txt_report}")
                    self.log(f"[MULTI TEST] Excel saved: {excel_report}")
                except Exception as report_error:
                    self.log(f"[MULTI TEST] REPORT ERROR for {server}: {report_error}")




        def server_worker(server_index, profile):
            nonlocal completed_servers
            server = profile.get("server", "")
            username = profile.get("username", "")
            port = profile.get("port", "22")
            authentication = profile.get("authentication", "Password")
            password = profile.get("password", "")
            key_file = profile.get("key_file", "")
            key_passphrase = profile.get("key_passphrase", "")
            session = None




            update_dashboard_status(server_index, "Connecting...", COLOR_WARNING)
            append_dashboard_output(server_index, "\n" + "=" * 60 + "\n")
            append_dashboard_output(server_index, f"[MULTI TEST] {server}\n")
            append_dashboard_output(server_index, "=" * 60 + "\n")




            self.log("\n" + "=" * 70)
            self.log(f"[MULTI TEST] {server}")
            self.log("=" * 70)




            try:
                update_dashboard_status(server_index, "Running Step 1...", COLOR_ACCENT)




                def step1_progress(done, total):
                    update_dashboard_progress(server_index, done, total, "Step 1")
                    update_dashboard_status(server_index, f"Running Step 1... {done}/{total}", COLOR_ACCENT)




                step1 = execute_ssh(
                    server=server, username=username, port=port,
                    authentication=authentication, password=password,
                    key_file=key_file, key_passphrase=key_passphrase,
                    commands=commands,
                    log_callback=lambda message, raw=False, idx=server_index: (
                        server_log_callback(idx, message, raw=raw)
                    ),
                    keep_connection=True,
                    progress_callback=step1_progress,
                    command_callback=self._monitor_command_event
                )




                if not step1:
                    for command in commands:
                        self._monitor_command_event(
                            "complete", server, command, "CONNECTION TIMED OUT",
                            time.time(), time.time()
                        )
                        store_comparison(server, command, "CONNECTION TIMED OUT")
                    update_dashboard_progress(server_index, len(commands), len(commands), "Step 1")
                    update_dashboard_status(
                        server_index,
                        f"Connection timed out — {len(commands)}/{len(commands)}",
                        COLOR_DANGER
                    )
                    add_row(server, 0, 0, len(commands), len(commands))
                    return




                step1_output, client, channel, key_dir = step1
                session = (client, channel, key_dir)




                unsuccessful = extract_unsuccessful(step1_output)
                step3_output = ""




                update_dashboard_status(
                    server_index,
                    f"Step 1 complete — {len(commands)}/{len(commands)}",
                    COLOR_SUCCESS
                )
                update_dashboard_progress(server_index, len(commands), len(commands), "Step 1")




                if STOP_DIG_EVENT.is_set():
                    update_dashboard_status(server_index, "Stopped by user", COLOR_DANGER)
                    return




                if unsuccessful:
                    retry_total = len(unsuccessful)
                    update_dashboard_status(
                        server_index,
                        f"Retrying unsuccessful commands: 0/{retry_total}",
                        COLOR_WARNING
                    )




                    def step3_progress(done, total):
                        update_dashboard_status(
                            server_index,
                            f"Retrying unsuccessful commands: {done}/{total}",
                            COLOR_WARNING
                        )
                        append_dashboard_output(
                            server_index,
                            f"\n[STEP 3] Retry progress: {done}/{total}\n"
                        )




                    step3_output = run_step3(
                        server=server, username=username, port=port,
                        authentication=authentication, password=password,
                        key_file=key_file, key_passphrase=key_passphrase,
                        unsuccessful_commands=unsuccessful,
                        log_callback=lambda message, raw=False, idx=server_index: (
                            server_log_callback(idx, message, raw=raw)
                        ),
                        session=session,
                        progress_callback=step3_progress,
                        command_callback=self._monitor_command_event
                    )




                make_server_reports(
                    server=server, username=username, authentication=authentication,
                    commands_for_report=commands, step1_output=step1_output,
                    unsuccessful=unsuccessful, step3_output=step3_output
                )




                step1_results = parse_actual_results(step1_output)
                step3_statuses = parse_step3_statuses(step3_output, unsuccessful)




                step3_by_key = {}
                for i, command in enumerate(unsuccessful):
                    if i < len(step3_statuses):
                        step3_by_key[command_key(command)] = step3_statuses[i]




                success = 0
                failed = 0
                timeout = 0




                for command in commands:
                    key = command_key(command)
                    if key in step3_by_key:
                        result_status = step3_by_key[key]
                    else:
                        item = step1_results.get(key)
                        result_status = item["status"] if item else "NO RESULT"




                    if result_status == "CONNECTION TIMED OUT":
                        timeout += 1
                    elif result_status in ("NOERROR", "NXDOMAIN"):
                        success += 1
                    else:
                        failed += 1




                    store_comparison(server, command, result_status)




                add_row(server, success, failed, timeout, len(commands))




                update_dashboard_progress(server_index, len(commands), len(commands), "Step 1")
                update_dashboard_status(
                    server_index,
                    (
                        f"Completed — {len(commands)}/{len(commands)} commands  •  "
                        f"Success: {success}  •  Failed: {failed}  •  Timeout: {timeout}"
                    ),
                    COLOR_SUCCESS
                )




            except Exception as exc:
                self.log(f"[MULTI TEST] {server} ERROR: {exc}")
                append_dashboard_output(server_index, f"\n[MULTI TEST ERROR] {exc}\n")
                for command in commands:
                    store_comparison(server, command, "ERROR")
                add_row(server, 0, len(commands), 0, len(commands))
                update_dashboard_status(server_index, "ERROR", COLOR_DANGER)
            finally:
                if session:
                    try:
                        session[1].send("exit\n")
                    except Exception:
                        pass
                    try:
                        session[1].close()
                    except Exception:
                        pass
                    try:
                        session[0].close()
                    except Exception:
                        pass
                    if session[2]:
                        shutil.rmtree(session[2], ignore_errors=True)




                global ACTIVE_SSH_CHANNEL
                with ACTIVE_SSH_CHANNEL_LOCK:
                    ACTIVE_SSH_CHANNEL = None




                with completed_lock:
                    completed_servers += 1
                    current_completed = completed_servers




                self.root.after(
                    0,
                    lambda c=current_completed, total=len(profiles): (
                        status_label.configure(text=f"Completed servers: {c} / {total}")
                    )
                )




        def create_command_comparison():
            try:
                timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
                comparison_file = os.path.join(
                    OUTPUT_DIR,
                    f"infoblox_multi_test_command_comparison_{timestamp}.xlsx"
                )




                wb = Workbook()
                ws = wb.active
                ws.title = "Command Comparison"




                server_names = [profile.get("server", "") for profile in profiles]
                headers = ["DIG Command"] + server_names + ["Result"]
                ws.append(headers)




                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F4E78")
                    cell.alignment = Alignment(horizontal="center")




                for command in commands:
                    clean_command = normalize_command(command)
                    server_results = comparison_results.get(clean_command, {})
                    results_for_command = [
                        server_results.get(server, "NO RESULT")
                        for server in server_names
                    ]
                    same_on_all = (
                        len(set(results_for_command)) == 1
                        and results_for_command[0] != "NO RESULT"
                    )
                    comparison_status = (
                        "✅ Same on all" if same_on_all else "❌ Not Same on all"
                    )
                    ws.append([clean_command] + results_for_command + [comparison_status])




                ws.column_dimensions["A"].width = 55
                for column_index in range(2, len(server_names) + 2):
                    ws.column_dimensions[get_column_letter(column_index)].width = 30
                result_column_index = len(server_names) + 2
                ws.column_dimensions[get_column_letter(result_column_index)].width = 22




                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                wb.save(comparison_file)




                self.log(f"[MULTI TEST] Command Comparison Excel saved: {comparison_file}")




                if os.name == "nt":
                    os.startfile(comparison_file)
                elif platform.system() == "Darwin":
                    subprocess.Popen(["open", comparison_file])
                else:
                    subprocess.Popen(["xdg-open", comparison_file])




                status_label.configure(
                    text=(
                        f"Completed: {completed_servers} / {len(profiles)} servers - "
                        "Command Comparison created"
                    )
                )
            except Exception as comparison_error:
                self.log(f"[MULTI TEST] COMMAND COMPARISON ERROR: {comparison_error}")




        def run_all_servers_parallel():
            threads = []
            for index, profile in enumerate(profiles):
                thread = threading.Thread(target=server_worker, args=(index, profile), daemon=True)
                threads.append(thread)
                thread.start()




            for thread in threads:
                thread.join()




            self.root.after(0, self._monitor_finish)
            self.root.after(0, create_command_comparison)
            self.root.after(0, lambda: self.run_button.configure(state="normal"))
            self.root.after(0, lambda: self.profiles_button.configure(state="normal"))
            self.root.after(0, lambda: self.stop_dig_button.configure(state="disabled"))




        threading.Thread(target=run_all_servers_parallel, daemon=True).start()




    # ========================================================
    # OPEN OUTPUT
    # ========================================================
    def open_output(self):
        folder = self.output_entry.get().strip() or OUTPUT_DIR
        if not os.path.exists(folder):
            os.makedirs(folder, exist_ok=True)
        if os.name == "nt":
            os.startfile(folder)
        elif platform.system() == "Darwin":
            subprocess.Popen(["open", folder])
        else:
            subprocess.Popen(["xdg-open", folder])








# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    try:
        if not ensure_quarterly_license():
            raise SystemExit(0)




        root = ctk.CTk()
        app = InfobloxDigApp(root)
        root.mainloop()




    except SystemExit:
        raise
    except Exception as exc:
        error_text = traceback.format_exc()
        try:
            log_file = os.path.join(BASE_DIR, "startup_error.log")
            with open(log_file, "w", encoding="utf-8") as f:
                f.write(error_text)
        except Exception:
            pass




        try:
            error_root = tk.Tk()
            error_root.withdraw()
            messagebox.showerror(
                "Application Startup Error",
                f"The application could not start.\n\n{exc}\n\n"
                "A detailed log was written to startup_error.log.",
                parent=error_root,
            )
            error_root.destroy()
        except Exception:
            print(error_text)






